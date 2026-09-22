from dotenv import load_dotenv
from pathlib import Path
import os

ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / '.env')

from fastapi import FastAPI, APIRouter, HTTPException, Request, Response, Depends, UploadFile, File, Form, Query, Header
from starlette.middleware.cors import CORSMiddleware
from starlette.responses import RedirectResponse, Response as StarletteResponse, StreamingResponse
from motor.motor_asyncio import AsyncIOMotorClient
from pymongo import ReturnDocument
from pydantic import BaseModel, Field, EmailStr
from typing import List, Optional, Dict, Any
from datetime import datetime, timezone, timedelta
import logging
import uuid
import jwt
import bcrypt
import requests
import json
import re
import tempfile
import io
from PIL import Image as PILImage
from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Font, PatternFill
import secrets
import asyncio
import ipaddress
import random
import httpx
import base64
import hashlib
from html import escape
from html.parser import HTMLParser
from urllib.parse import urlparse
from cryptography.fernet import Fernet, InvalidToken
from emergentintegrations.llm.chat import (
    FileContentWithMimeType,
    LlmChat,
    StreamDone,
    TextDelta,
    UserMessage,
)
from rag_service import (
    GEMINI_MODEL,
    OUT_OF_SCOPE_MESSAGE,
    extract_reference_text,
    reference_tokens,
    split_reference_text,
)

# ---------------------------------------------------------------------------
# Config
# ---------------------------------------------------------------------------
mongo_url = os.environ['MONGO_URL']
client = AsyncIOMotorClient(mongo_url)
db = client[os.environ['DB_NAME']]

JWT_SECRET = os.environ['JWT_SECRET']
JWT_ALGORITHM = "HS256"
ADMIN_EMAIL = os.environ.get('ADMIN_EMAIL', 'supri@baznasbazisdki.id')
ADMIN_PASSWORD = os.environ.get('ADMIN_PASSWORD', 'MdjSuper2026!')
DEMO_ADMIN_EMAIL = os.environ['DEMO_ADMIN_EMAIL']
DEMO_ADMIN_PASSWORD = os.environ['DEMO_ADMIN_PASSWORD']
DEMO_STUDENT_EMAIL = os.environ['DEMO_STUDENT_EMAIL']
DEMO_STUDENT_PASSWORD = os.environ['DEMO_STUDENT_PASSWORD']

STORAGE_BASE = (os.environ.get("INTEGRATION_PROXY_URL") or "").strip() or "https://integrations.emergentagent.com"
STORAGE_URL = STORAGE_BASE.rstrip("/") + "/objstore/api/v1/storage"
EMERGENT_KEY = os.environ.get("EMERGENT_LLM_KEY")
GEMINI_API_KEY = os.environ["GEMINI_API_KEY"]
APP_NAME = "mdj-scholarship"

# Email (Emergent-managed Resend)
EMAIL_BASE_URL = "https://integrations.emergentagent.com"
EMAIL_KEY = os.environ.get("EMERGENT_EMAIL_KEY")
EMAIL_FROM_NAME = os.environ.get("EMAIL_FROM_NAME", "MDJ Scholarship")
EMAIL_REPLY_TO = os.environ.get("EMAIL_REPLY_TO")
APP_BASE_URL = os.environ["APP_BASE_URL"].rstrip("/")

# Twilio SMS OTP (optional — phone verification disabled until configured)
TWILIO_ACCOUNT_SID = (os.environ.get("TWILIO_ACCOUNT_SID") or "").strip()
TWILIO_AUTH_TOKEN = (os.environ.get("TWILIO_AUTH_TOKEN") or "").strip()
TWILIO_VERIFY_SERVICE_SID = (os.environ.get("TWILIO_VERIFY_SERVICE_SID") or "").strip()
TWILIO_ENABLED = bool(TWILIO_ACCOUNT_SID and TWILIO_AUTH_TOKEN and TWILIO_VERIFY_SERVICE_SID)

EMERGENT_AUTH_URL = "https://demobackend.emergentagent.com/auth/v1/env/oauth/session-data"

MIME_TYPES = {
    "jpg": "image/jpeg", "jpeg": "image/jpeg", "png": "image/png",
    "gif": "image/gif", "webp": "image/webp", "pdf": "application/pdf",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}

logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(name)s - %(levelname)s - %(message)s')
logger = logging.getLogger(__name__)

app = FastAPI()
api_router = APIRouter(prefix="/api")

# ---------------------------------------------------------------------------
# Object storage
# ---------------------------------------------------------------------------
storage_key = None


def init_storage(force: bool = False):
    global storage_key
    if storage_key and not force:
        return storage_key
    resp = requests.post(f"{STORAGE_URL}/init", json={"emergent_key": EMERGENT_KEY}, timeout=30)
    resp.raise_for_status()
    storage_key = resp.json()["storage_key"]
    return storage_key


def compress_image(data: bytes, ext: str, content_type: str, max_w: int = 2400, quality: int = 86):
    """Downscale/re-encode raster images to JPEG; non-images (PDF) pass through."""
    if ext not in ("jpg", "jpeg", "png", "webp", "gif", "bmp", "tiff"):
        return data, ext, content_type
    try:
        img = PILImage.open(io.BytesIO(data))
        img = img.convert("RGB")
        if img.width > max_w:
            img = img.resize((max_w, int(img.height * max_w / img.width)), PILImage.LANCZOS)
        buf = io.BytesIO()
        img.save(buf, format="JPEG", quality=quality, optimize=True, progressive=True)
        out = buf.getvalue()
        if len(out) < len(data):
            return out, "jpg", "image/jpeg"
        return data, ext, content_type
    except Exception as e:
        logger.warning(f"Image compression skipped: {e}")
        return data, ext, content_type


def put_object(path: str, data: bytes, content_type: str) -> dict:
    key = init_storage()
    resp = requests.put(
        f"{STORAGE_URL}/objects/{path}",
        headers={"X-Storage-Key": key, "Content-Type": content_type},
        data=data, timeout=120,
    )
    if resp.status_code == 404:
        key = init_storage(force=True)
        resp = requests.put(
            f"{STORAGE_URL}/objects/{path}",
            headers={"X-Storage-Key": key, "Content-Type": content_type},
            data=data, timeout=120,
        )
    resp.raise_for_status()
    return resp.json()


def get_object(path: str):
    key = init_storage()
    resp = requests.get(f"{STORAGE_URL}/objects/{path}", headers={"X-Storage-Key": key}, timeout=60)
    if resp.status_code == 404:
        key = init_storage(force=True)
        resp = requests.get(f"{STORAGE_URL}/objects/{path}", headers={"X-Storage-Key": key}, timeout=60)
    resp.raise_for_status()
    return resp.content, resp.headers.get("Content-Type", "application/octet-stream")


# ---------------------------------------------------------------------------
# Email (Emergent-managed Resend) — guardrail gate + async send
# ---------------------------------------------------------------------------
_SHORTENERS = ("bit.ly", "tinyurl.com", "t.co", "is.gd", "cutt.ly", "goo.gl", "rebrand.ly")
_CRED_ASK = ("reply with your password", "reply with the code", "send your password", "cvv",
             "send us your password", "enter your password below", "confirm your card number",
             "your full card number", "seed phrase", "recovery phrase", "verify your card",
             "social security number", "confirm your bank details")
_HOSTISH = re.compile(r"\b(?:https?://)?((?:[a-z0-9-]+\.)+[a-z]{2,})", re.I)


def _host_ok(host: str) -> bool:
    if not host or "xn--" in host:
        return False
    try:
        ipaddress.ip_address(host)
        return False
    except ValueError:
        pass
    return not any(host == s or host.endswith("." + s) for s in _SHORTENERS)


def _same_site(shown: str, real: str) -> bool:
    return shown == real or real.endswith("." + shown) or shown.endswith("." + real)


class _EmailScan(HTMLParser):
    def __init__(self):
        super().__init__()
        self.tags, self.urls, self.anchors = set(), [], []
        self._href, self._text = None, []

    def handle_starttag(self, tag, attrs):
        self.tags.add(tag.lower())
        self.urls += [v for k, v in attrs if k.lower() in ("href", "src") and v]
        if tag.lower() == "a":
            self._href = dict((k.lower(), v) for k, v in attrs).get("href")
            self._text = []

    def handle_data(self, data):
        if self._href is not None:
            self._text.append(data)

    def handle_endtag(self, tag):
        if tag.lower() == "a" and self._href is not None:
            self.anchors.append((self._href, "".join(self._text)))
            self._href, self._text = None, []


def _assert_safe_email(subject: str, html: str) -> None:
    scan = _EmailScan()
    scan.feed(html)
    if scan.tags & {"form", "input", "textarea", "select"}:
        raise ValueError("No forms or input fields in email (G2)")
    body = f"{subject}\n{html}".lower()
    for p in _CRED_ASK:
        if p in body:
            raise ValueError(f"Email asks the recipient for credentials: {p!r} (G2)")
    for url in scan.urls:
        low = url.strip().lower()
        if low.startswith(("mailto:", "tel:", "cid:", "#")):
            continue
        if not low.startswith("https://"):
            raise ValueError(f"Email links/assets must be absolute https: {url!r} (G3)")
        host = urlparse(low).hostname or ""
        if not _host_ok(host) or urlparse(low).username is not None:
            raise ValueError(f"Shortened, numeric-host or credential-bearing URL: {url!r} (G3)")
    for href, text in scan.anchors:
        real = urlparse(href.strip().lower()).hostname or ""
        if not real:
            continue
        for m in _HOSTISH.finditer(text):
            if not _same_site(m.group(1).lower(), real):
                raise ValueError(f"Anchor text {m.group(1)!r} != real link host {real!r} (G3)")


async def send_email(*, to: str, subject: str, html: str) -> Optional[str]:
    _assert_safe_email(subject, html)
    payload = {"to": [to], "subject": subject, "html": html, "from_name": EMAIL_FROM_NAME}
    if EMAIL_REPLY_TO:
        payload["contact_email"] = EMAIL_REPLY_TO
    try:
        async with httpx.AsyncClient(timeout=30) as http_client:
            resp = await http_client.post(
                f"{EMAIL_BASE_URL}/api/v1/email/send",
                headers={"X-Email-Key": EMAIL_KEY},
                json=payload,
            )
        resp.raise_for_status()
        return resp.json().get("id")
    except httpx.HTTPStatusError as e:
        logger.error(f"Email send failed: {e.response.status_code} {e.response.text}")
        raise HTTPException(status_code=502, detail="Gagal mengirim email verifikasi")
    except Exception as e:
        logger.error(f"Email send error: {str(e)}")
        raise HTTPException(status_code=500, detail="Gagal mengirim email verifikasi")


# ---------------------------------------------------------------------------
# Auth helpers
# ---------------------------------------------------------------------------
def hash_password(password: str) -> str:
    return bcrypt.hashpw(password.encode("utf-8"), bcrypt.gensalt()).decode("utf-8")


def verify_password(plain: str, hashed: str) -> bool:
    try:
        return bcrypt.checkpw(plain.encode("utf-8"), hashed.encode("utf-8"))
    except Exception:
        return False


def create_access_token(user_id: str, email: str, auth_version: int = 0) -> str:
    payload = {
        "sub": user_id,
        "email": email,
        "av": auth_version,
        "exp": datetime.now(timezone.utc) + timedelta(days=7),
        "type": "access",
    }
    return jwt.encode(payload, JWT_SECRET, algorithm=JWT_ALGORITHM)


def set_auth_cookie(response: Response, token: str):
    response.set_cookie(key="access_token", value=token, httponly=True, secure=True,
                        samesite="none", max_age=604800, path="/")


async def resolve_user_from_token(token: str) -> Optional[dict]:
    # Try JWT first
    try:
        payload = jwt.decode(token, JWT_SECRET, algorithms=[JWT_ALGORITHM])
        if payload.get("type") == "access":
            user = await db.users.find_one({"user_id": payload["sub"]}, {"_id": 0})
            if user and int(payload.get("av", 0)) != int(user.get("auth_version", 0)):
                return None
            return user
    except jwt.InvalidTokenError:
        pass
    # Fallback: treat as google session token
    session = await db.user_sessions.find_one({"session_token": token}, {"_id": 0})
    if session:
        expires_at = session["expires_at"]
        if isinstance(expires_at, str):
            expires_at = datetime.fromisoformat(expires_at)
        if expires_at.tzinfo is None:
            expires_at = expires_at.replace(tzinfo=timezone.utc)
        if expires_at > datetime.now(timezone.utc):
            return await db.users.find_one({"user_id": session["user_id"]}, {"_id": 0})
    return None


async def get_current_user(request: Request) -> dict:
    token = request.cookies.get("access_token") or request.cookies.get("session_token")
    if not token:
        auth_header = request.headers.get("Authorization", "")
        if auth_header.startswith("Bearer "):
            token = auth_header[7:]
    if not token:
        raise HTTPException(status_code=401, detail="Not authenticated")
    user = await resolve_user_from_token(token)
    if not user:
        raise HTTPException(status_code=401, detail="Invalid or expired session")
    if not user.get("is_active", True):
        raise HTTPException(status_code=403, detail="Account is deactivated")
    if user.get("is_archived"):
        raise HTTPException(status_code=403, detail="Akun pendaftar telah diarsipkan")
    user.pop("password_hash", None)
    return user


def require_roles(*roles):
    async def checker(user: dict = Depends(get_current_user)) -> dict:
        if user.get("role") not in roles:
            raise HTTPException(status_code=403, detail="Access denied")
        return user
    return checker


# ---------------------------------------------------------------------------
# Models
# ---------------------------------------------------------------------------
class RegisterInput(BaseModel):
    name: str
    email: EmailStr
    password: str
    nik: Optional[str] = None
    phone: Optional[str] = None


class LoginInput(BaseModel):
    email: str
    password: str


class GoogleSessionInput(BaseModel):
    session_id: str


class ProfileInput(BaseModel):
    data: Dict[str, Any]


class ChangePasswordInput(BaseModel):
    current_password: str
    new_password: str


class EmailChangeInput(BaseModel):
    new_email: EmailStr
    password: str


class EmailVerifyInput(BaseModel):
    token: str


class ForgotPasswordInput(BaseModel):
    email: EmailStr


class ResetPasswordInput(BaseModel):
    token: str
    new_password: str


class PhoneOtpSendInput(BaseModel):
    phone: str


class PhoneOtpVerifyInput(BaseModel):
    phone: str
    code: str


class RegistrationInput(BaseModel):
    category: str
    action: str = "draft"  # draft | submit
    pakta_integritas_agreed: bool = False


class StatusUpdateInput(BaseModel):
    status: str
    note: Optional[str] = None


class DocumentEditingPermissionInput(BaseModel):
    allowed: bool


class SelectionAnnouncementPublishInput(BaseModel):
    category: str = "all"


class CreateAdminInput(BaseModel):
    name: str
    email: EmailStr
    password: str
    role: str = "admin"
    region: Optional[str] = None


class UserToggleInput(BaseModel):
    is_active: bool


class SiteContentInput(BaseModel):
    content: Dict[str, Any]


class CampusInput(BaseModel):
    name: str
    code: Optional[str] = None
    bank_account_number: Optional[str] = None
    bank_account_holder_name: Optional[str] = None


class VerificationApprovalInput(BaseModel):
    recommendation_ids: List[str] = []
    approve_all: bool = False


class BeneficiaryBankInput(BaseModel):
    bank_name: str = ""
    account_number: str = ""


class DisbursementInput(BaseModel):
    status: str
    amount: Optional[float] = None
    disbursed_at: Optional[str] = None
    reference: str = ""
    notes: str = ""


class BeneficiaryReviewResolveInput(BaseModel):
    user_id: Optional[str] = None
    campus_id: Optional[str] = None


class ActiveLetterDecisionInput(BaseModel):
    recommendation_ids: List[str] = []
    action: str
    apply_all: bool = False
    workflow: str
    source_ids: List[str] = []


class StudentDataPurgeInput(BaseModel):
    confirmation: str = Field(min_length=1, max_length=100)
    include_orphaned_registrations: bool = False


class AdminLiveChatMessageInput(BaseModel):
    session_id: str = Field(default="admin-team", min_length=1, max_length=100)
    message: str = Field(min_length=1, max_length=2000)


class StudentAiChatInput(BaseModel):
    question: str = Field(min_length=1, max_length=1200)


STUDENT_CHAT_ALLOWED_EXTENSIONS = {"jpg", "jpeg", "png", "webp", "pdf"}
STUDENT_CHAT_MAX_FILE_SIZE = 10 * 1024 * 1024
ADMIN_CHAT_MIME_TYPES = {
    "pdf": "application/pdf", "doc": "application/msword",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "xls": "application/vnd.ms-excel",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "jpg": "image/jpeg", "jpeg": "image/jpeg", "png": "image/png",
    "webp": "image/webp", "mp3": "audio/mpeg", "wav": "audio/wav",
    "ogg": "audio/ogg", "webm": "audio/webm", "m4a": "audio/mp4",
}


REG_STATUSES = ["draft", "submitted", "perlu_perbaikan", "verifikasi", "lolos_administrasi",
                "wawancara", "verifikasi_faktual", "lolos", "ditolak"]
SELECTION_RESULT_PASSED_STATUSES = (
    "lolos_administrasi",
    "wawancara",
    "verifikasi_faktual",
    "lolos",
)
SELECTION_RESULT_FAILED_STATUS = "ditolak"
SELECTION_RESULT_STATUSES = (*SELECTION_RESULT_PASSED_STATUSES, SELECTION_RESULT_FAILED_STATUS)
REGIONAL_ADMIN_ROLE = "admin_wilayah"
MANAGEMENT_ROLES = ("admin", "super_admin", REGIONAL_ADMIN_ROLE)
DKI_REGIONS = [
    "Jakarta Pusat",
    "Jakarta Utara",
    "Jakarta Barat",
    "Jakarta Selatan",
    "Jakarta Timur",
    "Kepulauan Seribu",
]
DISBURSEMENT_STATUSES = [
    "belum_diproses",
    "menunggu_bukti",
    "perlu_tinjau",
    "terverifikasi",
    "disetujui",
    "dicairkan",
    "ditunda",
]
TRANSFER_AMOUNT_PER_STUDENT = 3_000_000
UNIQUE_TRANSFER_TOLERANCE = 101
PM_MANAGERS = ("admin", "super_admin")
PM_ALLOWED_EXTENSIONS = {"pdf", "jpg", "jpeg", "png"}


def clean_user(u: dict) -> dict:
    u = dict(u)
    u.pop("password_hash", None)
    u.pop("_id", None)
    return u


def regional_admin_region(user: dict) -> Optional[str]:
    if user.get("role") != REGIONAL_ADMIN_ROLE:
        return None
    region = str(user.get("region") or "").strip()
    if region not in DKI_REGIONS:
        raise HTTPException(status_code=403, detail="Wilayah Admin Wilayah tidak valid.")
    return region


def profile_region(profile_data: dict) -> str:
    return str(profile_data.get("kota") or profile_data.get("provinsi") or "").strip()


async def participant_scope_query(user: dict) -> dict:
    region = regional_admin_region(user)
    if not region:
        return {}
    profiles = await db.profiles.find(
        {"$or": [{"data.kota": region}, {"data.provinsi": region}]},
        {"_id": 0, "user_id": 1},
    ).to_list(5000)
    user_ids = [profile["user_id"] for profile in profiles if profile.get("user_id")]
    return {"user_id": {"$in": user_ids}}


async def require_participant_scope(user: dict, user_id: str) -> None:
    region = regional_admin_region(user)
    if not region:
        return
    profile = await db.profiles.find_one({"user_id": user_id}, {"_id": 0, "data": 1})
    if not profile or profile_region(profile.get("data", {})) != region:
        raise HTTPException(status_code=404, detail="Pendaftar tidak ditemukan")


def can_manage_admin_account(actor: dict, target: dict) -> bool:
    if actor.get("role") == "super_admin":
        return True
    return (
        actor.get("role") == "admin"
        and target.get("role") == REGIONAL_ADMIN_ROLE
    )


def bank_cipher() -> Fernet:
    key = hashlib.sha256(JWT_SECRET.encode("utf-8")).digest()
    return Fernet(base64.urlsafe_b64encode(key))


def normalize_nim(value: Any) -> str:
    return re.sub(r"[^A-Za-z0-9]", "", str(value or "")).upper()


def normalize_account_number(value: Any) -> str:
    return re.sub(r"\D", "", str(value or ""))


def parse_transfer_amount(value: Any) -> Optional[int]:
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return int(round(value))
    text = str(value or "").strip()
    if not text:
        return None
    text = re.sub(r"[.,]\d{2}$", "", text)
    digits = re.sub(r"\D", "", text)
    return int(digits) if digits else None


def mask_account_number(value: str) -> str:
    if not value:
        return "Belum dicatat"
    if len(value) <= 4:
        return "•" * len(value)
    return f"{'•' * (len(value) - 4)}{value[-4:]}"


def encrypt_account_number(value: str) -> str:
    return bank_cipher().encrypt(value.encode("utf-8")).decode("utf-8")


def decrypt_account_number(value: str) -> str:
    if not value:
        return ""
    try:
        return bank_cipher().decrypt(value.encode("utf-8")).decode("utf-8")
    except InvalidToken:
        return ""


async def audit_beneficiary_action(
    actor: dict,
    action: str,
    user_id: Optional[str] = None,
    detail: Optional[dict] = None,
) -> None:
    await db.beneficiary_audit_events.insert_one({
        "id": str(uuid.uuid4()),
        "actor_id": actor["user_id"],
        "actor_name": actor.get("name", "Admin"),
        "action": action,
        "user_id": user_id,
        "detail": detail or {},
        "created_at": datetime.now(timezone.utc).isoformat(),
    })


# ---------------------------------------------------------------------------
# Auth endpoints
# ---------------------------------------------------------------------------
@api_router.post("/auth/register")
async def register(payload: RegisterInput, response: Response):
    email = payload.email.lower().strip()
    if await db.users.find_one({"email": email}):
        raise HTTPException(status_code=400, detail="Email sudah terdaftar")
    user_id = f"user_{uuid.uuid4().hex[:12]}"
    doc = {
        "user_id": user_id, "email": email, "password_hash": hash_password(payload.password),
        "name": payload.name, "role": "student", "auth_provider": "password",
        "nik": payload.nik, "phone": payload.phone, "picture": None,
        "is_active": True, "auth_version": 0, "created_at": datetime.now(timezone.utc).isoformat(),
    }
    await db.users.insert_one(doc)
    token = create_access_token(user_id, email, doc["auth_version"])
    set_auth_cookie(response, token)
    return {"user": clean_user(doc), "token": token}


@api_router.post("/auth/login")
async def login(payload: LoginInput, response: Response):
    identifier = payload.email.lower().strip()
    user = await db.users.find_one(
        {"$or": [{"email": identifier}, {"nik": identifier}]},
    )
    if not user or not user.get("password_hash") or not verify_password(payload.password, user["password_hash"]):
        raise HTTPException(status_code=401, detail="Email atau kata sandi salah")
    if not user.get("is_active", True):
        raise HTTPException(status_code=403, detail="Akun dinonaktifkan")
    if user.get("is_archived"):
        raise HTTPException(status_code=403, detail="Akun pendaftar telah diarsipkan")
    token = create_access_token(user["user_id"], user["email"], int(user.get("auth_version", 0)))
    set_auth_cookie(response, token)
    return {"user": clean_user(user), "token": token}


@api_router.post("/auth/google/session")
async def google_session(payload: GoogleSessionInput, response: Response):
    resp = requests.get(EMERGENT_AUTH_URL, headers={"X-Session-ID": payload.session_id}, timeout=30)
    if resp.status_code != 200:
        raise HTTPException(status_code=401, detail="Google authentication gagal")
    data = resp.json()
    email = data["email"].lower().strip()
    user = await db.users.find_one({"email": email})
    if not user:
        user_id = f"user_{uuid.uuid4().hex[:12]}"
        user = {
            "user_id": user_id, "email": email, "password_hash": None, "name": data.get("name", email),
            "role": "student", "auth_provider": "google", "nik": None, "phone": None,
            "picture": data.get("picture"), "is_active": True,
            "created_at": datetime.now(timezone.utc).isoformat(),
        }
        await db.users.insert_one(user)
    elif user.get("is_archived"):
        raise HTTPException(status_code=403, detail="Akun pendaftar telah diarsipkan")
    session_token = data["session_token"]
    await db.user_sessions.insert_one({
        "user_id": user["user_id"], "session_token": session_token,
        "expires_at": (datetime.now(timezone.utc) + timedelta(days=7)).isoformat(),
        "created_at": datetime.now(timezone.utc).isoformat(),
    })
    response.set_cookie(key="session_token", value=session_token, httponly=True, secure=True,
                        samesite="none", max_age=604800, path="/")
    return {"user": clean_user(user), "token": session_token}


@api_router.post("/auth/logout")
async def logout(request: Request, response: Response):
    token = request.cookies.get("session_token")
    if token:
        await db.user_sessions.delete_many({"session_token": token})
    response.delete_cookie("access_token", path="/")
    response.delete_cookie("session_token", path="/")
    return {"message": "Logged out"}


@api_router.get("/auth/me")
async def me(user: dict = Depends(get_current_user)):
    return clean_user(user)


def password_reset_token_hash(token: str) -> str:
    return hashlib.sha256(f"{token}:{JWT_SECRET}".encode("utf-8")).hexdigest()


def password_reset_rate_key(email: str, client_ip: str) -> str:
    return hashlib.sha256(f"{email}:{client_ip}:{JWT_SECRET}".encode("utf-8")).hexdigest()


async def password_reset_request_allowed(email: str, client_ip: str) -> bool:
    now = datetime.now(timezone.utc)
    key = password_reset_rate_key(email, client_ip)
    record = await db.password_reset_rate_limits.find_one({"key": key}, {"_id": 0})
    if record:
        expires_at = record.get("expires_at")
        if isinstance(expires_at, str):
            expires_at = datetime.fromisoformat(expires_at)
        if expires_at and expires_at.tzinfo is None:
            expires_at = expires_at.replace(tzinfo=timezone.utc)
        if expires_at and expires_at > now and record.get("count", 0) >= 3:
            return False
    await db.password_reset_rate_limits.update_one(
        {"key": key},
        {
            "$setOnInsert": {
                "key": key,
                "expires_at": now + timedelta(minutes=15),
                "created_at": now,
            },
            "$inc": {"count": 1},
        },
        upsert=True,
    )
    return True


@api_router.post("/auth/forgot-password")
async def forgot_password(payload: ForgotPasswordInput, request: Request):
    email = payload.email.lower().strip()
    client_ip = request.client.host if request.client else "unknown"
    generic_message = (
        "Jika alamat email terdaftar, tautan untuk mengatur ulang kata sandi akan dikirim ke email tersebut."
    )
    if not await password_reset_request_allowed(email, client_ip):
        return {"message": generic_message}
    user = await db.users.find_one({"email": email}, {"_id": 0})
    if not user or not user.get("is_active", True) or not user.get("password_hash"):
        return {"message": generic_message}
    token = secrets.token_urlsafe(32)
    token_hash = password_reset_token_hash(token)
    now = datetime.now(timezone.utc)
    await db.password_reset_tokens.update_many(
        {"user_id": user["user_id"], "used": False},
        {"$set": {"used": True, "invalidated_at": now}},
    )
    await db.password_reset_tokens.insert_one({
        "user_id": user["user_id"],
        "token_hash": token_hash,
        "used": False,
        "expires_at": now + timedelta(hours=1),
        "created_at": now,
        "requested_ip_hash": hashlib.sha256(
            f"{client_ip}:{JWT_SECRET}".encode("utf-8")
        ).hexdigest(),
    })
    link = f"{APP_BASE_URL}/reset-password?token={token}"
    html = (
        '<table role="presentation" width="100%"><tr><td '
        'style="padding:24px;font-family:Arial,sans-serif;color:#1F2937">'
        f'<p>Halo {escape(user.get("name", ""))},</p>'
        '<p>Kami menerima permintaan untuk mengatur ulang kata sandi akun MDJ Scholarship Anda.</p>'
        '<p>Klik tombol berikut untuk membuat kata sandi baru. Tautan ini berlaku selama 1 jam dan hanya '
        'dapat digunakan satu kali.</p>'
        f'<p><a href="{escape(link)}" style="display:inline-block;background:#27AE60;color:#ffffff;'
        'padding:12px 20px;border-radius:10px;text-decoration:none;font-weight:bold">'
        'Atur Ulang Kata Sandi</a></p>'
        '<p style="font-size:12px;color:#666">Jika Anda tidak meminta pengaturan ulang kata sandi, '
        'abaikan email ini. Kami tidak pernah meminta kata sandi melalui email.</p>'
        '</td></tr></table>'
    )
    await send_email(to=email, subject="Atur ulang kata sandi MDJ Scholarship", html=html)
    return {"message": generic_message}


@api_router.post("/auth/reset-password")
async def reset_password(payload: ResetPasswordInput):
    if len(payload.new_password) < 8:
        raise HTTPException(status_code=400, detail="Kata sandi baru minimal 8 karakter.")
    token_hash = password_reset_token_hash(payload.token)
    now = datetime.now(timezone.utc)
    token_record = await db.password_reset_tokens.find_one(
        {
            "token_hash": token_hash,
            "used": False,
            "expires_at": {"$gt": now},
        },
        {"_id": 0},
    )
    if not token_record:
        raise HTTPException(status_code=400, detail="Tautan reset tidak valid atau sudah kedaluwarsa.")
    user = await db.users.find_one({"user_id": token_record["user_id"]}, {"_id": 0})
    if not user or not user.get("is_active", True) or not user.get("password_hash"):
        raise HTTPException(status_code=400, detail="Tautan reset tidak dapat digunakan untuk akun ini.")
    if verify_password(payload.new_password, user["password_hash"]):
        raise HTTPException(status_code=400, detail="Kata sandi baru tidak boleh sama dengan yang lama.")
    consumed = await db.password_reset_tokens.update_one(
        {
            "token_hash": token_hash,
            "used": False,
            "expires_at": {"$gt": now},
        },
        {"$set": {"used": True, "used_at": now}},
    )
    if not consumed.matched_count:
        raise HTTPException(status_code=400, detail="Tautan reset tidak valid atau sudah kedaluwarsa.")
    next_auth_version = int(user.get("auth_version", 0)) + 1
    await db.users.update_one(
        {"user_id": user["user_id"]},
        {
            "$set": {
                "password_hash": hash_password(payload.new_password),
                "auth_version": next_auth_version,
                "password_reset_at": now,
            }
        },
    )
    await db.password_reset_tokens.update_many(
        {"user_id": user["user_id"], "used": False},
        {"$set": {"used": True, "invalidated_at": now}},
    )
    await db.user_sessions.delete_many({"user_id": user["user_id"]})
    return {"message": "Kata sandi berhasil diperbarui. Silakan masuk dengan kata sandi baru Anda."}


# ---------------------------------------------------------------------------
# Account settings — password, email change (verified link), phone OTP
# ---------------------------------------------------------------------------
@api_router.post("/auth/change-password")
async def change_password(payload: ChangePasswordInput, user: dict = Depends(get_current_user)):
    doc = await db.users.find_one({"user_id": user["user_id"]})
    if not doc or not doc.get("password_hash"):
        raise HTTPException(status_code=400, detail="Akun ini masuk lewat Google dan tidak memiliki kata sandi.")
    if not verify_password(payload.current_password, doc["password_hash"]):
        raise HTTPException(status_code=400, detail="Kata sandi saat ini salah.")
    if len(payload.new_password) < 8:
        raise HTTPException(status_code=400, detail="Kata sandi baru minimal 8 karakter.")
    if verify_password(payload.new_password, doc["password_hash"]):
        raise HTTPException(status_code=400, detail="Kata sandi baru tidak boleh sama dengan yang lama.")
    await db.users.update_one({"user_id": user["user_id"]},
                              {"$set": {"password_hash": hash_password(payload.new_password)}})
    return {"message": "Kata sandi berhasil diperbarui."}


@api_router.post("/auth/email/change-request")
async def request_email_change(payload: EmailChangeInput, user: dict = Depends(get_current_user)):
    doc = await db.users.find_one({"user_id": user["user_id"]})
    if not doc or not doc.get("password_hash"):
        raise HTTPException(status_code=400, detail="Akun Google tidak dapat mengubah email di sini.")
    if not verify_password(payload.password, doc["password_hash"]):
        raise HTTPException(status_code=400, detail="Kata sandi salah.")
    new_email = payload.new_email.lower().strip()
    if new_email == doc["email"]:
        raise HTTPException(status_code=400, detail="Email baru sama dengan email saat ini.")
    if await db.users.find_one({"email": new_email}):
        raise HTTPException(status_code=400, detail="Email tersebut sudah digunakan akun lain.")
    token = secrets.token_urlsafe(32)
    await db.email_change_tokens.update_many({"user_id": user["user_id"], "used": False},
                                             {"$set": {"used": True}})
    await db.email_change_tokens.insert_one({
        "user_id": user["user_id"], "new_email": new_email, "token": token, "used": False,
        "expires_at": (datetime.now(timezone.utc) + timedelta(hours=1)).isoformat(),
        "created_at": datetime.now(timezone.utc).isoformat(),
    })
    link = f"{APP_BASE_URL}/verify-email?token={token}"
    html = (
        '<table role="presentation" width="100%"><tr><td style="padding:24px;font-family:Arial,sans-serif;color:#1F2937">'
        f'<p>Halo {escape(doc.get("name", ""))},</p>'
        f'<p>Kami menerima permintaan untuk mengubah alamat email akun MDJ Scholarship Anda menjadi '
        f'<strong>{escape(new_email)}</strong>.</p>'
        f'<p>Klik tautan berikut untuk mengonfirmasi perubahan (berlaku 1 jam):</p>'
        f'<p><a href="{escape(link)}" style="display:inline-block;background:#27AE60;color:#ffffff;'
        f'padding:12px 20px;border-radius:10px;text-decoration:none;font-weight:bold">Verifikasi Email Saya</a></p>'
        f'<p style="font-size:12px;color:#888">Jika Anda tidak meminta perubahan ini, abaikan email ini. '
        f'Sent by {escape(EMAIL_FROM_NAME)}. Kami tidak pernah meminta kata sandi Anda melalui email.</p>'
        '</td></tr></table>'
    )
    await send_email(to=new_email, subject="Verifikasi perubahan email MDJ Scholarship", html=html)
    return {"message": f"Tautan verifikasi telah dikirim ke {new_email}. Silakan cek kotak masuk Anda."}


@api_router.post("/auth/email/verify")
async def verify_email_change(payload: EmailVerifyInput):
    rec = await db.email_change_tokens.find_one({"token": payload.token, "used": False})
    if not rec:
        raise HTTPException(status_code=400, detail="Tautan verifikasi tidak valid atau sudah digunakan.")
    expires_at = rec["expires_at"]
    if isinstance(expires_at, str):
        expires_at = datetime.fromisoformat(expires_at)
    if expires_at < datetime.now(timezone.utc):
        raise HTTPException(status_code=400, detail="Tautan verifikasi sudah kedaluwarsa.")
    new_email = rec["new_email"]
    if await db.users.find_one({"email": new_email, "user_id": {"$ne": rec["user_id"]}}):
        raise HTTPException(status_code=400, detail="Email tersebut sudah digunakan akun lain.")
    await db.users.update_one({"user_id": rec["user_id"]},
                              {"$set": {"email": new_email, "email_verified": True}})
    await db.email_change_tokens.update_one({"token": payload.token}, {"$set": {"used": True}})
    return {"message": "Email berhasil diperbarui.", "email": new_email}


@api_router.post("/auth/phone/send-otp")
async def send_phone_otp(payload: PhoneOtpSendInput, user: dict = Depends(get_current_user)):
    if not TWILIO_ENABLED:
        raise HTTPException(status_code=503, detail="Verifikasi SMS belum dikonfigurasi. Hubungi administrator.")
    phone = payload.phone.strip()
    if not re.fullmatch(r"\+[1-9]\d{7,14}", phone):
        raise HTTPException(status_code=400, detail="Nomor harus format internasional, contoh: +6281234567890.")
    try:
        from twilio.rest import Client
        client = Client(TWILIO_ACCOUNT_SID, TWILIO_AUTH_TOKEN)
        await asyncio.to_thread(
            lambda: client.verify.v2.services(TWILIO_VERIFY_SERVICE_SID)
            .verifications.create(to=phone, channel="sms")
        )
    except Exception as e:
        logger.error(f"Twilio send OTP failed: {e}")
        raise HTTPException(status_code=502, detail="Gagal mengirim kode OTP. Periksa nomor Anda.")
    return {"message": f"Kode OTP telah dikirim ke {phone}."}


@api_router.post("/auth/phone/verify-otp")
async def verify_phone_otp(payload: PhoneOtpVerifyInput, user: dict = Depends(get_current_user)):
    if not TWILIO_ENABLED:
        raise HTTPException(status_code=503, detail="Verifikasi SMS belum dikonfigurasi. Hubungi administrator.")
    phone = payload.phone.strip()
    try:
        from twilio.rest import Client
        client = Client(TWILIO_ACCOUNT_SID, TWILIO_AUTH_TOKEN)
        check = await asyncio.to_thread(
            lambda: client.verify.v2.services(TWILIO_VERIFY_SERVICE_SID)
            .verification_checks.create(to=phone, code=payload.code)
        )
    except Exception as e:
        logger.error(f"Twilio verify OTP failed: {e}")
        raise HTTPException(status_code=502, detail="Gagal memverifikasi kode OTP.")
    if check.status != "approved":
        raise HTTPException(status_code=400, detail="Kode OTP salah atau kedaluwarsa.")
    await db.users.update_one({"user_id": user["user_id"]},
                              {"$set": {"phone": phone, "phone_verified": True}})
    return {"message": "Nomor telepon berhasil diverifikasi.", "phone": phone}


# ---------------------------------------------------------------------------
# Student profile
# ---------------------------------------------------------------------------
@api_router.get("/profile")
async def get_profile(user: dict = Depends(get_current_user)):
    profile = await db.profiles.find_one({"user_id": user["user_id"]}, {"_id": 0})
    if not profile:
        profile = {"user_id": user["user_id"], "data": {
            "namaLengkap": user.get("name", ""), "email": user.get("email", ""),
            "nik": user.get("nik", ""), "noTelp": user.get("phone", ""),
        }}
    return profile


@api_router.put("/profile")
async def update_profile(payload: ProfileInput, user: dict = Depends(get_current_user)):
    student_campus = None
    updated_user = None
    if user.get("role") == "student":
        student_campus = await register_student_campus(payload.data.get("institusi"))
        display_name = str(payload.data.get("namaLengkap", "")).strip()
        if display_name:
            await db.users.update_one(
                {"user_id": user["user_id"]},
                {"$set": {"name": display_name}},
            )
            updated_user = await db.users.find_one(
                {"user_id": user["user_id"]},
                {"_id": 0},
            )
    await db.profiles.update_one(
        {"user_id": user["user_id"]},
        {"$set": {"data": payload.data, "updated_at": datetime.now(timezone.utc).isoformat()}},
        upsert=True,
    )
    return {
        "message": "Profil disimpan",
        "data": payload.data,
        "campus": student_campus,
        "user": clean_user(updated_user) if updated_user else None,
    }


# ---------------------------------------------------------------------------
# Registration (student)
# ---------------------------------------------------------------------------
REGISTRATION_CLOSED_MESSAGE = (
    "Masa pendaftaran belum dibuka. Silakan pantau pengumuman MDJ Scholarship secara berkala."
)
PAKTA_INTEGRITAS_VERSION = "2026-09"
PAKTA_INTEGRITAS_DOCUMENTS = [
    "KTP DKI Jakarta",
    "Kartu Keluarga (KK)",
    "Pas Foto 3x4",
    "Kartu Tanda Mahasiswa (KTM)",
    "KRS / KHS / Transkrip Nilai",
    "Surat Keterangan Mahasiswa Aktif",
    "SKTM / Surat Rekomendasi",
    "Surat Persetujuan Orang Tua",
    "Surat Keterangan Tidak Menerima Beasiswa Lain",
]
PAKTA_INTEGRITAS_EDUCATION_FIELDS = {
    "jenjang": "Jenjang pendidikan",
    "institusi": "Perguruan tinggi",
    "jurusan": "Program studi",
    "nim": "NIM",
    "semester": "Semester",
    "ipk": "IPK",
    "biayaPendidikanSemester": "Nominal biaya pendidikan per semester",
}
PAKTA_INTEGRITAS_REQUIRED_MESSAGE = (
    "Anda wajib menyetujui Pakta Integritas sebelum mengirim pendaftaran."
)
DOCUMENTS_LOCKED_MESSAGE = (
    "Dokumen pendaftaran sudah dikirim dan terkunci. Panitia perlu mengizinkan edit kembali."
)
DOCUMENT_REVISION_DEFAULT_NOTE = (
    "Mohon lengkapi dan perbaiki data serta dokumen pendaftaran Anda dengan benar."
)
STATUS_LABELS = {
    "draft": "Draft",
    "submitted": "Terkirim",
    "perlu_perbaikan": "Perlu Perbaikan Berkas",
    "verifikasi": "Verifikasi Administrasi",
    "lolos_administrasi": "Lolos Administrasi",
    "wawancara": "Wawancara Assessment",
    "verifikasi_faktual": "Verifikasi Faktual",
    "lolos": "Lolos / Penerima Manfaat",
    "ditolak": "Tidak Lolos",
}


async def create_notification(
    user_id: str,
    notification_type: str,
    title: str,
    message: str,
    metadata: Optional[dict] = None,
) -> None:
    notification = {
        "id": str(uuid.uuid4()),
        "user_id": user_id,
        "type": notification_type,
        "title": title,
        "message": message,
        "is_read": False,
        "created_at": datetime.now(timezone.utc).isoformat(),
    }
    if metadata:
        notification.update(metadata)
    await db.notifications.insert_one(notification)


async def selection_announcement_summary(category: str) -> dict:
    selected_category = str(category or "all").strip()
    categories = sorted(filter(None, await db.registrations.distinct("category")))
    if selected_category != "all" and selected_category not in categories:
        raise HTTPException(status_code=400, detail="Kategori pendaftaran tidak valid.")

    category_query = {}
    if selected_category != "all":
        category_query["category"] = selected_category

    unpublished_query = {
        **category_query,
        "status": {"$in": SELECTION_RESULT_STATUSES},
        "$or": [
            {"is_announcement_published": {"$exists": False}},
            {"is_announcement_published": False},
        ],
    }
    recipients = await db.registrations.find(
        unpublished_query,
        {"_id": 0, "user_id": 1, "status": 1},
    ).to_list(5000)
    passed_count = sum(
        1 for registration in recipients
        if registration.get("status") in SELECTION_RESULT_PASSED_STATUSES
    )
    failed_count = sum(
        1 for registration in recipients
        if registration.get("status") == SELECTION_RESULT_FAILED_STATUS
    )
    pending_count = await db.registrations.count_documents({
        **category_query,
        "status": {"$nin": SELECTION_RESULT_STATUSES},
    })
    latest_publication = await db.selection_announcements.find_one(
        {"category": selected_category, "status_publikasi": "Published"},
        {"_id": 0},
        sort=[("published_at", -1)],
    )
    return {
        "category": selected_category,
        "categories": categories,
        "passed_count": passed_count,
        "failed_count": failed_count,
        "pending_count": pending_count,
        "recipient_count": len(recipients),
        "latest_publication": latest_publication,
    }


def announcement_signature(announcement: dict) -> str:
    fields = ("date", "category", "title", "summary")
    return "|".join(str(announcement.get(field, "")).strip() for field in fields)


async def notify_new_announcements(announcements: List[dict]) -> None:
    if not announcements:
        return
    registrations = await db.registrations.find(
        {"is_archived": {"$ne": True}},
        {"_id": 0, "user_id": 1},
    ).to_list(5000)
    recipients = {registration.get("user_id") for registration in registrations}
    notifications = []
    for announcement in announcements:
        signature = announcement_signature(announcement)
        for user_id in recipients:
            if not user_id:
                continue
            notifications.append({
                "id": str(uuid.uuid4()),
                "user_id": user_id,
                "type": "announcement",
                "source_id": signature,
                "title": announcement.get("title") or "Pengumuman MDJ Scholarship",
                "message": announcement.get("summary") or "Ada pengumuman baru dari Admin MDJ.",
                "is_read": False,
                "is_popup_seen": False,
                "created_at": datetime.now(timezone.utc).isoformat(),
            })
    if notifications:
        await db.notifications.insert_many(notifications)


async def create_cpm_id(registered_at: Optional[str] = None) -> str:
    year = str(datetime.now(timezone.utc).year)
    if registered_at:
        year = registered_at[:4]
    counter = await db.counters.find_one_and_update(
        {"key": f"cpm-{year}"},
        {"$inc": {"sequence": 1}},
        upsert=True,
        return_document=ReturnDocument.AFTER,
    )
    sequence = max(int(counter["sequence"]) - 1, 0)
    return f"CPM-MDJ/{year}/{sequence:06d}"


async def ensure_cpm_id(registration: Optional[dict]) -> Optional[dict]:
    if not registration or registration.get("cpm_id"):
        return registration

    cpm_id = await create_cpm_id(registration.get("created_at"))
    updated = await db.registrations.find_one_and_update(
        {
            "user_id": registration["user_id"],
            "$or": [{"cpm_id": {"$exists": False}}, {"cpm_id": None}],
        },
        {"$set": {"cpm_id": cpm_id}},
        return_document=ReturnDocument.AFTER,
        projection={"_id": 0},
    )
    if updated:
        return updated
    return await db.registrations.find_one(
        {"user_id": registration["user_id"]},
        {"_id": 0},
    )


async def is_registration_period_open() -> bool:
    content = await db.site_content.find_one({"key": "main"}, {"_id": 0, "settings": 1})
    settings = (content or {}).get("settings") or {}
    return settings.get("registration_open") is True


async def registration_completion_gaps(user_id: str) -> Dict[str, List[str]]:
    profile = await db.profiles.find_one({"user_id": user_id}, {"_id": 0, "data": 1})
    profile_data = (profile or {}).get("data") or {}
    missing_education = [
        label
        for field, label in PAKTA_INTEGRITAS_EDUCATION_FIELDS.items()
        if not str(profile_data.get(field, "")).strip()
    ]
    documents = await db.documents.find(
        {"user_id": user_id, "is_deleted": False},
        {"_id": 0, "doc_type": 1},
    ).to_list(100)
    document_types = {document.get("doc_type") for document in documents}
    missing_documents = [
        document
        for document in PAKTA_INTEGRITAS_DOCUMENTS
        if document not in document_types
    ]
    return {
        "missing_education": missing_education,
        "missing_documents": missing_documents,
    }


async def ensure_registration_documents_editable(user: dict) -> None:
    registration = await db.registrations.find_one(
        {"user_id": user["user_id"]},
        {"_id": 0, "status": 1, "documents_editing_allowed": 1},
    )
    if not registration or registration.get("status") == "draft":
        return
    if registration.get("documents_editing_allowed") is True:
        return
    raise HTTPException(status_code=403, detail=DOCUMENTS_LOCKED_MESSAGE)


@api_router.get("/registration")
async def get_registration(user: dict = Depends(get_current_user)):
    reg = await db.registrations.find_one({"user_id": user["user_id"]}, {"_id": 0})
    return await ensure_cpm_id(reg) or {}


@api_router.post("/registration")
async def submit_registration(payload: RegistrationInput, user: dict = Depends(get_current_user)):
    existing = await db.registrations.find_one({"user_id": user["user_id"]})
    is_resubmission = bool(existing and existing.get("revision_requested"))
    if not await is_registration_period_open() and not is_resubmission:
        raise HTTPException(status_code=403, detail=REGISTRATION_CLOSED_MESSAGE)
    if payload.action == "submit":
        if not is_resubmission and not payload.pakta_integritas_agreed:
            raise HTTPException(status_code=400, detail=PAKTA_INTEGRITAS_REQUIRED_MESSAGE)
        gaps = await registration_completion_gaps(user["user_id"])
        if gaps["missing_education"] or gaps["missing_documents"]:
            raise HTTPException(
                status_code=400,
                detail={
                    "message": "Lengkapi data pendidikan dan dokumen sebelum mengirim pendaftaran.",
                    **gaps,
                },
            )
    status = "submitted" if payload.action == "submit" else "draft"
    now = datetime.now(timezone.utc).isoformat()
    if existing:
        await ensure_cpm_id(existing)
        update = {"category": payload.category, "updated_at": now}
        if payload.action == "submit" and (existing.get("status") == "draft" or is_resubmission):
            submission_note = (
                "Berkas perbaikan dikirim ulang kepada panitia"
                if is_resubmission
                else "Pakta Integritas disetujui dan pendaftaran dikirim"
            )
            update.update(
                {
                    "status": "submitted",
                    "submitted_at": now,
                    "documents_editing_allowed": False,
                    "documents_locked_at": now,
                    "revision_requested": False,
                    "revision_note": None,
                    "revision_resubmitted_at": now if is_resubmission else None,
                }
            )
            if not is_resubmission:
                update.update(
                    {
                        "pakta_integritas_agreed": True,
                        "pakta_integritas_agreed_at": now,
                        "pakta_integritas_version": PAKTA_INTEGRITAS_VERSION,
                    }
                )
            await db.registrations.update_one(
                {"user_id": user["user_id"]},
                {
                    "$set": update,
                    "$push": {
                        "history": {
                            "status": "submitted",
                            "note": submission_note,
                            "at": now,
                        }
                    },
                },
            )
            if is_resubmission and existing.get("revision_requested_by_user_id"):
                await create_notification(
                    existing["revision_requested_by_user_id"],
                    "document_revision_resubmitted",
                    "Berkas perbaikan telah dikirim ulang",
                    "Mahasiswa telah mengirim ulang berkas perbaikan untuk verifikasi Anda.",
                    {"source_id": existing["user_id"]},
                )
        else:
            await db.registrations.update_one({"user_id": user["user_id"]}, {"$set": update})
        reg = await db.registrations.find_one({"user_id": user["user_id"]}, {"_id": 0})
        return reg
    reg = {
        "id": str(uuid.uuid4()), "user_id": user["user_id"], "name": user.get("name"),
        "email": user.get("email"), "category": payload.category, "status": status,
        "cpm_id": await create_cpm_id(now),
        "history": [{"status": status, "note": "Pendaftaran dibuat", "at": now}],
        "pakta_integritas_agreed": payload.action == "submit",
        "pakta_integritas_agreed_at": now if payload.action == "submit" else None,
        "pakta_integritas_version": PAKTA_INTEGRITAS_VERSION if payload.action == "submit" else None,
        "documents_editing_allowed": payload.action != "submit",
        "documents_locked_at": now if payload.action == "submit" else None,
        "created_at": now, "submitted_at": now if status == "submitted" else None, "updated_at": now,
    }
    await db.registrations.insert_one(reg)
    reg.pop("_id", None)
    return reg


@api_router.get("/notifications")
async def list_notifications(user: dict = Depends(get_current_user)):
    notifications = await db.notifications.find(
        {"user_id": user["user_id"]},
        {"_id": 0},
    ).sort("created_at", -1).to_list(30)
    unread_count = sum(1 for notification in notifications if not notification.get("is_read"))
    return {"notifications": notifications, "unread_count": unread_count}


@api_router.post("/notifications/{notification_id}/read")
async def mark_notification_read(
    notification_id: str,
    user: dict = Depends(get_current_user),
):
    result = await db.notifications.update_one(
        {"id": notification_id, "user_id": user["user_id"]},
        {"$set": {"is_read": True}},
    )
    if not result.matched_count:
        raise HTTPException(status_code=404, detail="Notifikasi tidak ditemukan.")
    return {"message": "Notifikasi ditandai sudah dibaca."}


@api_router.post("/notifications/read-all")
async def mark_all_notifications_read(user: dict = Depends(get_current_user)):
    await db.notifications.update_many(
        {"user_id": user["user_id"], "is_read": False},
        {"$set": {"is_read": True}},
    )
    return {"message": "Semua notifikasi ditandai sudah dibaca."}


@api_router.get("/student/site-announcements/pending")
async def get_pending_site_announcement(user: dict = Depends(require_roles("student"))):
    content = await db.site_content.find_one({"key": "main"}, {"_id": 0, "announcements": 1})
    announcements = (content or {}).get("announcements") or []
    for announcement in announcements:
        if not isinstance(announcement, dict):
            continue
        signature = announcement_signature(announcement)
        title = announcement.get("title") or "Pengumuman MDJ Scholarship"
        message = announcement.get("summary") or "Ada pengumuman baru dari Admin MDJ."
        existing = await db.notifications.find_one(
            {
                "user_id": user["user_id"],
                "type": "announcement",
                "$or": [
                    {"source_id": signature},
                    {"title": title, "message": message},
                ],
            },
            {"_id": 0, "id": 1},
        )
        if existing:
            continue
        await db.notifications.insert_one(
            {
                "id": str(uuid.uuid4()),
                "user_id": user["user_id"],
                "type": "announcement",
                "source_id": signature,
                "title": title,
                "message": message,
                "is_read": False,
                "is_popup_seen": False,
                "created_at": datetime.now(timezone.utc).isoformat(),
            }
        )
    notification = await db.notifications.find_one(
        {
            "user_id": user["user_id"],
            "type": "announcement",
            "$or": [
                {"is_popup_seen": {"$exists": False}},
                {"is_popup_seen": False},
            ],
        },
        {"_id": 0},
        sort=[("created_at", -1)],
    )
    return {"announcement": notification}


@api_router.post("/student/site-announcements/{notification_id}/seen")
async def mark_site_announcement_seen(
    notification_id: str,
    user: dict = Depends(require_roles("student")),
):
    now = datetime.now(timezone.utc).isoformat()
    result = await db.notifications.update_one(
        {
            "id": notification_id,
            "user_id": user["user_id"],
            "type": "announcement",
        },
        {"$set": {"is_read": True, "is_popup_seen": True, "popup_seen_at": now}},
    )
    if not result.matched_count:
        raise HTTPException(status_code=404, detail="Pengumuman tidak ditemukan.")
    return {"message": "Pengumuman telah dibaca."}


@api_router.get("/student/document-revisions/pending")
async def get_pending_document_revision(user: dict = Depends(require_roles("student"))):
    notification = await db.notifications.find_one(
        {
            "user_id": user["user_id"],
            "type": "document_revision_requested",
            "$or": [
                {"is_popup_seen": {"$exists": False}},
                {"is_popup_seen": False},
            ],
        },
        {"_id": 0},
        sort=[("created_at", -1)],
    )
    return {"revision": notification}


@api_router.post("/student/document-revisions/{notification_id}/seen")
async def mark_document_revision_seen(
    notification_id: str,
    user: dict = Depends(require_roles("student")),
):
    now = datetime.now(timezone.utc).isoformat()
    result = await db.notifications.update_one(
        {
            "id": notification_id,
            "user_id": user["user_id"],
            "type": "document_revision_requested",
        },
        {"$set": {"is_read": True, "is_popup_seen": True, "popup_seen_at": now}},
    )
    if not result.matched_count:
        raise HTTPException(status_code=404, detail="Notifikasi perbaikan tidak ditemukan.")
    return {"message": "Notifikasi perbaikan telah dibaca."}


@api_router.get("/student/selection-announcements/pending")
async def get_pending_selection_announcement(
    user: dict = Depends(require_roles("student")),
):
    registration = await db.registrations.find_one(
        {"user_id": user["user_id"]},
        {"_id": 0, "user_id": 1, "status": 1, "category": 1, "is_announcement_published": 1},
    )
    if (
        registration
        and registration.get("status") in SELECTION_RESULT_STATUSES
        and not registration.get("is_announcement_published")
    ):
        campaign = None
        registration_category = registration.get("category", "")
        if registration_category:
            campaign = await db.selection_announcements.find_one(
                {"status_publikasi": "Published", "category": registration_category},
                {"_id": 0},
                sort=[("published_at", -1)],
            )
        if not campaign:
            campaign = await db.selection_announcements.find_one(
                {"status_publikasi": "Published", "category": "all"},
                {"_id": 0},
                sort=[("published_at", -1)],
            )
        if campaign:
            now = datetime.now(timezone.utc).isoformat()
            claimed = await db.registrations.find_one_and_update(
                {
                    "user_id": user["user_id"],
                    "status": {"$in": SELECTION_RESULT_STATUSES},
                    "$or": [
                        {"is_announcement_published": {"$exists": False}},
                        {"is_announcement_published": False},
                    ],
                },
                {
                    "$set": {
                        "is_announcement_published": True,
                        "selection_announcement_id": campaign["id"],
                        "selection_announcement_published_at": now,
                    }
                },
                projection={"_id": 0, "user_id": 1, "status": 1, "category": 1},
                return_document=ReturnDocument.AFTER,
            )
            if claimed:
                result = (
                    "passed"
                    if claimed.get("status") in SELECTION_RESULT_PASSED_STATUSES
                    else "failed"
                )
                await db.notifications.insert_one(
                    {
                        "id": str(uuid.uuid4()),
                        "user_id": claimed["user_id"],
                        "type": "selection_result",
                        "source_id": campaign["id"],
                        "announcement_id": campaign["id"],
                        "category": claimed.get("category", ""),
                        "result": result,
                        "title": campaign["title"],
                        "message": campaign["message"],
                        "is_read": False,
                        "is_popup_seen": False,
                        "created_at": now,
                    }
                )
                await db.selection_announcements.update_one(
                    {"id": campaign["id"]},
                    {"$inc": {"late_recipient_count": 1}},
                )
    notification = await db.notifications.find_one(
        {
            "user_id": user["user_id"],
            "type": "selection_result",
            "$or": [
                {"is_popup_seen": {"$exists": False}},
                {"is_popup_seen": False},
            ],
        },
        {"_id": 0},
        sort=[("created_at", -1)],
    )
    return {"announcement": notification}


@api_router.post("/student/selection-announcements/{notification_id}/seen")
async def mark_selection_announcement_seen(
    notification_id: str,
    user: dict = Depends(require_roles("student")),
):
    now = datetime.now(timezone.utc).isoformat()
    result = await db.notifications.update_one(
        {
            "id": notification_id,
            "user_id": user["user_id"],
            "type": "selection_result",
        },
        {
            "$set": {
                "is_read": True,
                "is_popup_seen": True,
                "popup_seen_at": now,
            }
        },
    )
    if not result.matched_count:
        raise HTTPException(status_code=404, detail="Pengumuman seleksi tidak ditemukan.")
    return {"message": "Pengumuman seleksi telah dibaca."}


# ---------------------------------------------------------------------------
# Documents
# ---------------------------------------------------------------------------
@api_router.post("/documents")
async def upload_document(doc_type: str = Form(...), file: UploadFile = File(...),
                          user: dict = Depends(get_current_user)):
    await ensure_registration_documents_editable(user)
    ext = file.filename.split(".")[-1].lower() if "." in file.filename else "bin"
    content_type = file.content_type or MIME_TYPES.get(ext, "application/octet-stream")
    data, ext, content_type = compress_image(
        await file.read(),
        ext,
        content_type,
        max_w=2400,
        quality=86,
    )
    path = f"{APP_NAME}/uploads/{user['user_id']}/{uuid.uuid4()}.{ext}"
    result = put_object(path, data, content_type)
    doc_id = str(uuid.uuid4())
    # Replace any previous doc of same type (soft delete)
    await db.documents.update_many(
        {"user_id": user["user_id"], "doc_type": doc_type, "is_deleted": False},
        {"$set": {"is_deleted": True}},
    )
    record = {
        "id": doc_id, "user_id": user["user_id"], "doc_type": doc_type,
        "storage_path": result["path"], "original_filename": file.filename,
        "content_type": content_type, "size": result["size"], "is_deleted": False,
        "created_at": datetime.now(timezone.utc).isoformat(),
    }
    await db.documents.insert_one(record)
    record.pop("_id", None)
    return record


@api_router.get("/documents")
async def list_documents(user: dict = Depends(get_current_user)):
    docs = await db.documents.find({"user_id": user["user_id"], "is_deleted": False}, {"_id": 0}).to_list(200)
    return docs


@api_router.delete("/documents/{doc_id}")
async def delete_document(doc_id: str, user: dict = Depends(get_current_user)):
    await ensure_registration_documents_editable(user)
    await db.documents.update_one({"id": doc_id, "user_id": user["user_id"]}, {"$set": {"is_deleted": True}})
    return {"message": "Dokumen dihapus"}


KTP_SYSTEM_PROMPT = """Anda adalah mesin OCR ahli untuk KTP (Kartu Tanda Penduduk) Indonesia.
Ekstrak data dari gambar/PDF KTP yang diberikan dan kembalikan HANYA objek JSON valid,
tanpa penjelasan, tanpa markdown fence. Gunakan skema kunci PERSIS berikut:
{
  "namaLengkap": string,          // Nama sesuai KTP, huruf kapital seperti aslinya
  "nik": string,                  // 16 digit angka saja
  "tempatLahir": string,          // Kota/kabupaten tempat lahir
  "tanggalLahir": string,         // WAJIB format YYYY-MM-DD (konversi dari DD-MM-YYYY)
  "jenisKelamin": string,         // "L" untuk LAKI-LAKI, "P" untuk PEREMPUAN
  "agama": string,                // salah satu: islam, kristen, katolik, hindu, buddha, konghucu
  "statusPerkawinan": string,     // "belum_kawin" (BELUM KAWIN), "kawin" (KAWIN), "cerai" (CERAI HIDUP/MATI)
  "alamatLengkap": string,        // Isi baris ALAMAT
  "rt": string,                   // RT saja (angka)
  "rw": string,                   // RW saja (angka)
  "kelurahan": string,            // KEL/DESA
  "kecamatan": string,            // KECAMATAN
  "kota": string,                 // Kota/Kabupaten (dari header atau alamat)
  "provinsi": string              // Provinsi (dari header KTP)
}
Aturan: gunakan null untuk field yang tidak terbaca. Jangan mengarang data.
Jika file bukan KTP, kembalikan {"error": "bukan_ktp"}."""


KK_SYSTEM_PROMPT = """Anda adalah mesin OCR ahli untuk Kartu Keluarga (KK) Indonesia.
Ekstrak data dari gambar/PDF Kartu Keluarga yang diberikan dan kembalikan HANYA objek JSON valid,
tanpa penjelasan, tanpa markdown fence. Gunakan skema kunci PERSIS berikut:
{
  "noKK": string                  // Nomor Kartu Keluarga, 16 digit angka saja (tertera di bagian atas dokumen, "No. ...")
}
Aturan: gunakan null jika tidak terbaca. Jangan mengarang data.
Jika file bukan Kartu Keluarga, kembalikan {"error": "bukan_dokumen"}."""


KTM_SYSTEM_PROMPT = """Anda adalah mesin OCR ahli untuk Kartu Tanda Mahasiswa (KTM) Indonesia.
Ekstrak data dari gambar/PDF KTM yang diberikan dan kembalikan HANYA objek JSON valid,
tanpa penjelasan atau markdown. Gunakan skema kunci PERSIS berikut:
{
  "institusi": string,
  "nim": string,
  "jurusan": string,
  "jenjang": string
}
Aturan: jenjang hanya "D3", "D4", atau "S1" bila terbaca jelas. Gunakan null untuk field
yang tidak terbaca dan jangan mengarang data. Jika file bukan KTM, kembalikan
{"error": "bukan_ktm"}."""


ACADEMIC_RECORD_SYSTEM_PROMPT = """Anda adalah mesin OCR ahli untuk dokumen akademik Indonesia.
Dokumen dapat berupa KRS, KHS, atau Transkrip Nilai. Ekstrak data dan kembalikan HANYA
objek JSON valid, tanpa penjelasan atau markdown. Gunakan skema kunci PERSIS berikut:
{
  "institusi": string,
  "nim": string,
  "jurusan": string,
  "jenjang": string,
  "semester": string,
  "ipk": string
}
Aturan: semester berisi angka saja bila terbaca. IPK menggunakan format desimal, misalnya
"3.75". Jenjang hanya "D3", "D4", atau "S1" bila terbaca jelas. Gunakan null untuk field
yang tidak terbaca dan jangan mengarang data. Jika file bukan KRS, KHS, atau Transkrip Nilai,
kembalikan {"error": "bukan_dokumen_akademik"}."""


async def _save_scanned_doc(user: dict, doc_type: str, filename: str, ext: str, data: bytes, content_type: str):
    data, ext, content_type = compress_image(
        data,
        ext,
        content_type,
        max_w=2400,
        quality=86,
    )
    path = f"{APP_NAME}/uploads/{user['user_id']}/{uuid.uuid4()}.{ext}"
    result = put_object(path, data, content_type)
    await db.documents.update_many(
        {"user_id": user["user_id"], "doc_type": doc_type, "is_deleted": False},
        {"$set": {"is_deleted": True}},
    )
    record = {
        "id": str(uuid.uuid4()), "user_id": user["user_id"], "doc_type": doc_type,
        "storage_path": result["path"], "original_filename": filename,
        "content_type": content_type, "size": result["size"], "is_deleted": False,
        "created_at": datetime.now(timezone.utc).isoformat(),
    }
    await db.documents.insert_one(record)
    record.pop("_id", None)
    return record


async def _extract_document(file: UploadFile, user: dict, label: str, system_prompt: str, allowed: set, doc_type: str):
    ext = file.filename.split(".")[-1].lower() if "." in file.filename else "bin"
    if ext not in ("jpg", "jpeg", "png", "webp", "pdf"):
        raise HTTPException(status_code=400, detail="Format harus JPG, PNG, WEBP, atau PDF.")
    mime = MIME_TYPES.get(ext, "application/octet-stream")
    data = await file.read()
    if len(data) > 10 * 1024 * 1024:
        raise HTTPException(status_code=400, detail="Ukuran file maksimal 10MB.")
    tmp_path = None
    try:
        with tempfile.NamedTemporaryFile(delete=False, suffix=f".{ext}") as tmp:
            tmp.write(data)
            tmp_path = tmp.name
        chat = LlmChat(
            api_key=EMERGENT_KEY,
            session_id=f"ocr-{user['user_id']}-{uuid.uuid4()}",
            system_message=system_prompt,
        ).with_model("gemini", "gemini-3.1-pro-preview")
        fc = FileContentWithMimeType(mime_type=mime, file_path=tmp_path)
        chunks = []
        async for event in chat.stream_message(UserMessage(
            text=f"Ekstrak seluruh data dari {label} ini dan kembalikan HANYA JSON sesuai skema.",
            file_contents=[fc],
        )):
            if isinstance(event, TextDelta):
                chunks.append(event.content)
            elif isinstance(event, StreamDone):
                break
        raw = "".join(chunks)
    except Exception as e:
        logger.error(f"{label} extraction failed: {e}")
        raise HTTPException(status_code=502, detail=f"Gagal membaca {label}. Coba lagi dengan foto yang lebih jelas.")
    finally:
        if tmp_path and os.path.exists(tmp_path):
            os.unlink(tmp_path)

    cleaned = raw.strip()
    if cleaned.startswith("```"):
        cleaned = re.sub(r"^```(?:json)?\s*|\s*```$", "", cleaned, flags=re.DOTALL).strip()
    match = re.search(r"\{.*\}", cleaned, flags=re.DOTALL)
    if not match:
        raise HTTPException(status_code=422, detail=f"Data {label} tidak dapat dikenali. Pastikan foto jelas.")
    try:
        parsed = json.loads(match.group(0))
    except json.JSONDecodeError:
        raise HTTPException(status_code=422, detail=f"Data {label} tidak dapat dikenali. Pastikan foto jelas.")
    if parsed.get("error"):
        raise HTTPException(status_code=422, detail=f"File yang diunggah bukan {label} yang valid.")

    mapped = {k: str(v).strip() for k, v in parsed.items() if k in allowed and v not in (None, "", "null")}
    for key in ("nik", "noKK"):
        if mapped.get(key):
            mapped[key] = re.sub(r"\D", "", mapped[key])[:16]
    if mapped.get("semester"):
        mapped["semester"] = re.sub(r"\D", "", mapped["semester"])[:2]
    if mapped.get("ipk"):
        mapped["ipk"] = re.sub(r"[^0-9,.]", "", mapped["ipk"]).replace(",", ".")
    if mapped.get("jenjang"):
        jenjang = mapped["jenjang"].upper().replace(" ", "")
        mapped["jenjang"] = jenjang if jenjang in {"D3", "D4", "S1"} else ""
    document = None
    try:
        document = await _save_scanned_doc(user, doc_type, file.filename, ext, data, mime)
    except Exception as e:
        logger.error(f"Auto-save scanned {label} failed: {e}")
        raise HTTPException(
            status_code=502,
            detail="Data terbaca, tetapi dokumen belum tersimpan. Silakan unggah ulang.",
        )
    return {"data": mapped, "document": document}


@api_router.post("/profile/extract-ktp")
async def extract_ktp(file: UploadFile = File(...), user: dict = Depends(get_current_user)):
    await ensure_registration_documents_editable(user)
    allowed = {"namaLengkap", "nik", "tempatLahir", "tanggalLahir", "jenisKelamin",
               "agama", "statusPerkawinan", "alamatLengkap", "rt", "rw",
               "kelurahan", "kecamatan", "kota", "provinsi"}
    return await _extract_document(file, user, "KTP", KTP_SYSTEM_PROMPT, allowed, "KTP DKI Jakarta")


@api_router.post("/profile/extract-kk")
async def extract_kk(file: UploadFile = File(...), user: dict = Depends(get_current_user)):
    await ensure_registration_documents_editable(user)
    return await _extract_document(file, user, "Kartu Keluarga", KK_SYSTEM_PROMPT, {"noKK"}, "Kartu Keluarga (KK)")


@api_router.post("/profile/extract-ktm")
async def extract_ktm(file: UploadFile = File(...), user: dict = Depends(get_current_user)):
    await ensure_registration_documents_editable(user)
    fields = {"institusi", "nim", "jurusan", "jenjang"}
    return await _extract_document(
        file,
        user,
        "Kartu Tanda Mahasiswa",
        KTM_SYSTEM_PROMPT,
        fields,
        "Kartu Tanda Mahasiswa (KTM)",
    )


@api_router.post("/profile/extract-academic-record")
async def extract_academic_record(
    file: UploadFile = File(...),
    user: dict = Depends(get_current_user),
):
    await ensure_registration_documents_editable(user)
    fields = {"institusi", "nim", "jurusan", "jenjang", "semester", "ipk"}
    return await _extract_document(
        file,
        user,
        "KRS, KHS, atau Transkrip Nilai",
        ACADEMIC_RECORD_SYSTEM_PROMPT,
        fields,
        "KRS / KHS / Transkrip Nilai",
    )


@api_router.get("/files/{path:path}")
async def download_file(path: str, request: Request, auth: str = Query(None)):
    site_record = await db.site_files.find_one({"storage_path": path, "is_deleted": False}, {"_id": 0})
    if site_record:
        data, content_type = get_object(path)
        return StarletteResponse(content=data, media_type=site_record.get("content_type", content_type),
                                 headers={"Cache-Control": "public, max-age=31536000, immutable"})
    token = request.cookies.get("access_token") or request.cookies.get("session_token") or auth
    if not token:
        auth_header = request.headers.get("Authorization", "")
        if auth_header.startswith("Bearer "):
            token = auth_header[7:]
    if not token or not await resolve_user_from_token(token):
        raise HTTPException(status_code=401, detail="Not authenticated")
    record = await db.documents.find_one({"storage_path": path, "is_deleted": False}, {"_id": 0})
    data, content_type = get_object(path)
    return StarletteResponse(content=data, media_type=(record or {}).get("content_type", content_type),
                             headers={"Cache-Control": "private, max-age=3600"})


# ---------------------------------------------------------------------------
# Penerima Manfaat and disbursement management
# ---------------------------------------------------------------------------
PM_ACTIVE_LETTER_PROMPT = """Anda adalah mesin OCR dokumen kampus Indonesia.
Dokumen adalah surat keterangan mahasiswa aktif atau surat jawaban verifikasi yang dapat
berisi tabel banyak mahasiswa. Kembalikan HANYA JSON valid dengan skema:
{
  "campus": string,
  "students": [
    {
      "name": string,
      "nim": string,
      "study_program": string,
      "semester": string,
      "evidence": string
    }
  ]
}
Ekstrak SEMUA baris mahasiswa. Pertahankan NIM persis seperti tertulis, termasuk nol di
depan. Gunakan string kosong bila tidak terbaca dan jangan menebak data."""

PM_TRANSFER_PROOF_PROMPT = """Anda adalah mesin OCR bukti transfer Indonesia.
Bukti dapat berupa transfer massal dengan banyak transaksi. Kembalikan HANYA JSON valid:
{
  "transactions": [
    {
      "campus": string,
      "account_number": string,
      "account_holder_name": string,
      "bank_name": string,
      "status": string,
      "amount": string,
      "reference": string,
      "date": string,
      "evidence": string
    }
  ],
  "summary": {"total_amount": string, "date": string, "reference": string}
}
Gunakan HANYA nama kampus, nomor rekening tujuan, dan nama pemilik rekening sebagai dasar
informasi transaksi. JANGAN memakai atau mengekstrak nama mahasiswa maupun NIM. Prioritaskan
nomor rekening tujuan, ekstrak semua transaksi yang terlihat, jangan menebak digit rekening,
dan gunakan string kosong untuk data yang tidak terbaca."""


def pm_content_type(ext: str) -> str:
    return MIME_TYPES.get(ext, "application/octet-stream")


def pm_magic_matches(data: bytes, ext: str) -> bool:
    signatures = {
        "pdf": b"%PDF",
        "jpg": b"\xff\xd8\xff",
        "jpeg": b"\xff\xd8\xff",
        "png": b"\x89PNG\r\n\x1a\n",
    }
    return bool(data) and data.startswith(signatures[ext])


async def read_pm_upload(file: UploadFile) -> tuple:
    filename = file.filename or "dokumen"
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext not in PM_ALLOWED_EXTENSIONS:
        raise HTTPException(status_code=400, detail="Gunakan PDF, JPG, JPEG, atau PNG.")
    data = await file.read()
    if not data or len(data) > 10 * 1024 * 1024:
        raise HTTPException(status_code=400, detail="Ukuran berkas harus antara 1 byte dan 10MB.")
    if not pm_magic_matches(data, ext):
        raise HTTPException(status_code=400, detail="Isi berkas tidak sesuai dengan formatnya.")
    return filename, ext, data, pm_content_type(ext)


async def store_pm_source(
    actor: dict,
    source_type: str,
    filename: str,
    ext: str,
    data: bytes,
    content_type: str,
    stage: Optional[int] = None,
    selected_region: Optional[str] = None,
) -> dict:
    source_id = str(uuid.uuid4())
    path = f"{APP_NAME}/beneficiary-sources/{source_id}.{ext}"
    result = put_object(path, data, content_type)
    source = {
        "id": source_id,
        "source_type": source_type,
        "stage": stage,
        "selected_region": selected_region,
        "storage_path": result["path"],
        "original_filename": filename,
        "content_type": content_type,
        "size": result["size"],
        "sha256": hashlib.sha256(data).hexdigest(),
        "status": "processing",
        "uploaded_by": actor["user_id"],
        "uploaded_by_name": actor.get("name", "Admin"),
        "matched_user_ids": [],
        "regions": [],
        "created_at": datetime.now(timezone.utc).isoformat(),
    }
    await db.beneficiary_sources.insert_one(dict(source))
    await audit_beneficiary_action(actor, f"upload_{source_type}", detail={"source_id": source_id})
    return source


async def extract_pm_payload(
    data: bytes,
    content_type: str,
    actor: dict,
    prompt: str,
    source_label: str,
) -> dict:
    suffix = ".pdf" if content_type == "application/pdf" else ".img"
    tmp_path = None
    try:
        with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as temp_file:
            temp_file.write(data)
            tmp_path = temp_file.name
        chat = LlmChat(
            api_key=EMERGENT_KEY,
            session_id=f"pm-ocr-{actor['user_id']}-{uuid.uuid4()}",
            system_message=prompt,
        ).with_model("gemini", "gemini-3.1-pro-preview")
        file_content = FileContentWithMimeType(mime_type=content_type, file_path=tmp_path)
        chunks = []
        async for event in chat.stream_message(UserMessage(
            text=f"Baca {source_label} ini dan kembalikan JSON sesuai skema.",
            file_contents=[file_content],
        )):
            if isinstance(event, TextDelta):
                chunks.append(event.content)
            elif isinstance(event, StreamDone):
                break
        raw = "".join(chunks).strip()
    except Exception as error:
        logger.error("PM OCR failed: %s", error)
        raise HTTPException(status_code=502, detail="AI belum dapat membaca berkas. Coba lagi nanti.")
    finally:
        if tmp_path and os.path.exists(tmp_path):
            os.unlink(tmp_path)
    raw = re.sub(r"^```(?:json)?\s*|\s*```$", "", raw, flags=re.DOTALL).strip()
    match = re.search(r"\{.*\}", raw, flags=re.DOTALL)
    if not match:
        raise HTTPException(status_code=422, detail="Hasil OCR tidak dapat dibaca sebagai data.")
    try:
        parsed = json.loads(match.group(0))
    except json.JSONDecodeError as error:
        raise HTTPException(status_code=422, detail="Hasil OCR tidak valid.") from error
    if not isinstance(parsed, dict):
        raise HTTPException(status_code=422, detail="Hasil OCR tidak memiliki struktur yang valid.")
    return parsed


async def final_beneficiary_index(user: Optional[dict] = None) -> Dict[str, List[dict]]:
    query = {"status": "lolos"}
    if user:
        query.update(await participant_scope_query(user))
    registrations = await db.registrations.find(query, {"_id": 0}).to_list(5000)
    user_ids = [registration["user_id"] for registration in registrations]
    profiles = await db.profiles.find(
        {"user_id": {"$in": user_ids}},
        {"_id": 0, "user_id": 1, "data": 1},
    ).to_list(5000)
    profiles_by_user = {profile["user_id"]: profile.get("data", {}) for profile in profiles}
    index: Dict[str, List[dict]] = {}
    for registration in registrations:
        profile = profiles_by_user.get(registration["user_id"], {})
        name = registration.get("name") or profile.get("namaLengkap", "")
        nim = normalize_nim(profile.get("nim"))
        key = f"{normalized_name(name)}|{nim}"
        if not normalized_name(name) or not nim:
            continue
        index.setdefault(key, []).append({
            "registration": registration,
            "profile": profile,
            "region": profile_region(profile),
        })
    return index


def beneficiary_campus(profile: dict) -> str:
    return str(profile.get("institusi") or "Kampus belum diisi").strip()


def campus_disbursement_key(campus: str) -> str:
    normalized = normalized_name(campus) or "kampus-belum-diisi"
    digest = hashlib.sha256(normalized.encode("utf-8")).hexdigest()[:20]
    return f"campus_{digest}"


async def final_beneficiary_rows(user: Optional[dict] = None) -> List[dict]:
    query = {"status": "lolos"}
    if user:
        query.update(await participant_scope_query(user))
    registrations = await db.registrations.find(query, {"_id": 0}).to_list(5000)
    user_ids = [item["user_id"] for item in registrations]
    profiles = await db.profiles.find(
        {"user_id": {"$in": user_ids}},
        {"_id": 0, "user_id": 1, "data": 1},
    ).to_list(5000)
    profiles_by_user = {item["user_id"]: item.get("data", {}) for item in profiles}
    rows = []
    for registration in registrations:
        profile = profiles_by_user.get(registration["user_id"], {})
        campus = beneficiary_campus(profile)
        rows.append({
            "user_id": registration["user_id"],
            "name": registration.get("name") or profile.get("namaLengkap", "-"),
            "cpm_id": registration.get("cpm_id", "-"),
            "nim": profile.get("nim", "-"),
            "campus": campus,
            "campus_key": campus_disbursement_key(campus),
            "region": profile_region(profile) or "Belum diisi",
            "study_program": profile.get("jurusan", "-"),
            "semester": profile.get("semester", "-"),
        })
    return rows


async def campus_bank_candidates_for_region(region: str) -> List[dict]:
    beneficiaries = await final_beneficiary_rows()
    rows_by_campus = {}
    for row in beneficiaries:
        if row["region"] != region:
            continue
        rows_by_campus.setdefault(row["campus_key"], []).append(row)
    if not rows_by_campus:
        return []
    campuses = await db.campuses.find(
        {"bank_account_cipher": {"$exists": True}},
        {"_id": 0},
    ).to_list(5000)
    candidates = []
    for campus in campuses:
        campus_key = campus_disbursement_key(campus.get("name", ""))
        account_number = decrypt_account_number(campus.get("bank_account_cipher", ""))
        account_holder_name = normalized_name(campus.get("bank_account_holder_name", ""))
        if campus_key not in rows_by_campus or not account_number or not account_holder_name:
            continue
        candidates.append({
            "campus": campus,
            "recipient_rows": rows_by_campus[campus_key],
            "campus_key": campus_key,
            "account_number": account_number,
            "account_holder_name": account_holder_name,
        })
    return candidates


def match_campus_bank_transaction(transaction: dict, candidates: List[dict]) -> List[dict]:
    campus_name = normalized_name(transaction.get("campus"))
    account_number = normalize_account_number(transaction.get("account_number"))
    account_holder_name = normalized_name(transaction.get("account_holder_name"))
    if not campus_name or not account_number or not account_holder_name:
        return []
    return [
        candidate
        for candidate in candidates
        if normalized_name(candidate["campus"].get("name", "")) == campus_name
        and candidate["account_number"] == account_number
        and candidate["account_holder_name"] == account_holder_name
    ]


def default_campus_disbursement(stage: int) -> dict:
    return {
        "stage": stage,
        "status": "belum_diproses",
        "amount": None,
        "disbursed_at": None,
        "reference": "",
        "notes": "",
        "proofs": [],
        "transactions": [],
        "transfer_amount": None,
        "expected_amount": None,
        "paid_student_count": 0,
        "recipient_count": 0,
        "payment_assessment": "belum_ada_transfer",
        "payment_difference": None,
        "remarks": [],
    }


async def ensure_campus_disbursement(
    campus_key: str,
    campus: str,
    region: str,
    stage: int,
    recipient_ids: List[str],
) -> dict:
    now = datetime.now(timezone.utc).isoformat()
    defaults = default_campus_disbursement(stage)
    # Fields also set in $set below must NOT be in $setOnInsert (MongoDB
    # rejects overlapping paths).
    defaults.pop("recipient_count", None)
    defaults.update({
        "id": str(uuid.uuid4()),
        "campus_key": campus_key,
        "campus": campus,
        "region": region,
        "created_at": now,
    })
    await db.campus_disbursements.update_one(
        {"campus_key": campus_key, "region": region, "stage": stage},
        {
            "$setOnInsert": defaults,
            "$set": {
                "recipient_user_ids": recipient_ids,
                "recipient_count": len(recipient_ids),
                "synced_at": now,
            },
        },
        upsert=True,
    )
    return await db.campus_disbursements.find_one(
        {"campus_key": campus_key, "region": region, "stage": stage},
        {"_id": 0},
    )


def serialize_campus_disbursement(record: Optional[dict], stage: int) -> dict:
    data = record or default_campus_disbursement(stage)
    return {
        "stage": stage,
        "status": data.get("status", "belum_diproses"),
        "amount": data.get("amount"),
        "disbursed_at": data.get("disbursed_at"),
        "reference": data.get("reference", ""),
        "notes": data.get("notes", ""),
        "proofs": data.get("proofs", []),
        "transactions": data.get("transactions", []),
        "transfer_amount": data.get("transfer_amount"),
        "expected_amount": data.get("expected_amount"),
        "paid_student_count": data.get("paid_student_count", 0),
        "recipient_count": data.get("recipient_count", 0),
        "payment_assessment": data.get("payment_assessment", "belum_ada_transfer"),
        "payment_difference": data.get("payment_difference"),
        "remarks": build_disbursement_remarks(data),
        "updated_at": data.get("updated_at"),
    }


def build_disbursement_remarks(data: dict) -> List[str]:
    remarks = []
    assessment = data.get("payment_assessment")
    difference = abs(int(data.get("payment_difference") or 0))
    if assessment == "kurang" and difference:
        remarks.append(f"Kekurangan transfer Rp{difference:,.0f}".replace(",", "."))
    if assessment == "lebih" and difference:
        remarks.append(f"Kelebihan transfer Rp{difference:,.0f}".replace(",", "."))
    if data.get("manually_edited"):
        remarks.append("Data diedit oleh Admin")
    notes = str(data.get("notes") or "").strip()
    if notes:
        remarks.append(notes)
    return remarks


def apply_recipient_financial_defaults(stage_data: dict, recipient_count: int) -> dict:
    data = dict(stage_data)
    data["recipient_count"] = recipient_count
    if data.get("expected_amount") is None:
        data["expected_amount"] = recipient_count * TRANSFER_AMOUNT_PER_STUDENT
    return data


def combine_campus_stage(records: List[dict], stage: int) -> dict:
    items = [serialize_campus_disbursement(record, stage) for record in records]
    if not items:
        return default_campus_disbursement(stage)
    statuses = {item["status"] for item in items}
    proof_ids = {
        proof.get("source_id")
        for item in items
        for proof in item.get("proofs", [])
        if proof.get("source_id")
    }
    transfer_amount = sum(int(item.get("transfer_amount") or 0) for item in items)
    expected_amount = sum(int(item.get("expected_amount") or 0) for item in items)
    paid_student_count = sum(int(item.get("paid_student_count") or 0) for item in items)
    assessments = {item.get("payment_assessment") for item in items if item.get("transfer_amount")}
    proof_map = {
        proof.get("source_id"): proof
        for item in items
        for proof in item.get("proofs", [])
        if proof.get("source_id")
    }
    remarks = list(dict.fromkeys(
        remark
        for item in items
        for remark in item.get("remarks", [])
        if remark
    ))
    dates = {item.get("disbursed_at") for item in items if item.get("disbursed_at")}
    return {
        "stage": stage,
        "status": statuses.pop() if len(statuses) == 1 else "bervariasi",
        "amount": sum(float(item["amount"] or 0) for item in items) or None,
        "proof_count": len(proof_ids),
        "region_count": len(items),
        "transfer_amount": transfer_amount or None,
        "expected_amount": expected_amount or None,
        "paid_student_count": paid_student_count,
        "recipient_count": sum(int(item.get("recipient_count") or 0) for item in items),
        "payment_assessment": assessments.pop() if len(assessments) == 1 else "bervariasi",
        "payment_difference": transfer_amount - expected_amount if transfer_amount else None,
        "disbursed_at": dates.pop() if len(dates) == 1 else None,
        "proofs": list(proof_map.values()),
        "remarks": remarks,
    }


async def visible_campus_groups(user: dict, requested_region: Optional[str] = None) -> dict:
    scope_region = regional_admin_region(user)
    selected_region = scope_region or requested_region
    rows = await final_beneficiary_rows(user)
    grouped: Dict[str, dict] = {}
    for row in rows:
        if selected_region and selected_region != "all" and row["region"] != selected_region:
            continue
        campus = grouped.setdefault(row["campus_key"], {
            "campus_key": row["campus_key"],
            "campus": row["campus"],
            "by_region": {},
        })
        campus["by_region"].setdefault(row["region"], []).append(row)
    campus_keys = list(grouped)
    if not campus_keys:
        return grouped
    record_query = {"campus_key": {"$in": campus_keys}}
    if selected_region and selected_region != "all":
        record_query["region"] = selected_region
    records = await db.campus_disbursements.find(record_query, {"_id": 0}).to_list(5000)
    for record in records:
        campus = grouped.get(record.get("campus_key"))
        if not campus or record.get("region") not in campus["by_region"]:
            continue
        campus.setdefault("records", {}).setdefault(record["region"], {})[record["stage"]] = record
    return grouped


@api_router.get("/admin/disbursements/campuses")
async def list_campus_disbursements(
    region: Optional[str] = None,
    search: Optional[str] = None,
    user: dict = Depends(require_roles(*MANAGEMENT_ROLES)),
):
    groups = await visible_campus_groups(user, region)
    keyword = (search or "").strip().lower()
    result = []
    for group in groups.values():
        if keyword and keyword not in group["campus"].lower():
            continue
        region_records = group.get("records", {})
        stage_one = combine_campus_stage(
            [records.get(1) for records in region_records.values() if records.get(1)],
            1,
        )
        stage_two = combine_campus_stage(
            [records.get(2) for records in region_records.values() if records.get(2)],
            2,
        )
        recipient_count = sum(len(items) for items in group["by_region"].values())
        stage_one = apply_recipient_financial_defaults(stage_one, recipient_count)
        stage_two = apply_recipient_financial_defaults(stage_two, recipient_count)
        result.append({
            "campus_key": group["campus_key"],
            "campus": group["campus"],
            "recipient_count": recipient_count,
            "region_count": len(group["by_region"]),
            "regions": sorted(group["by_region"]),
            "students": sorted(
                [
                    {"name": row["name"], "region": row["region"]}
                    for students in group["by_region"].values()
                    for row in students
                ],
                key=lambda item: item["name"].lower(),
            ),
            "stage_one": stage_one,
            "stage_two": stage_two,
            "proof_count": len({
                proof.get("source_id")
                for stage in (stage_one, stage_two)
                for record in group.get("records", {}).values()
                for item in record.values()
                for proof in item.get("proofs", [])
                if proof.get("source_id")
            }),
        })
    return sorted(result, key=lambda item: item["campus"].lower())


@api_router.get("/admin/disbursements/export")
async def export_campus_disbursements(
    stage: int = 1,
    region: Optional[str] = None,
    user: dict = Depends(require_roles(*MANAGEMENT_ROLES)),
):
    if stage not in {1, 2}:
        raise HTTPException(status_code=400, detail="Tahap pencairan tidak valid.")
    campuses = await list_campus_disbursements(region=region, user=user)
    workbook = Workbook()
    recap = workbook.active
    recap.title = f"Rekap Tahap {stage}"
    headers = ["No", "Kampus", "Wilayah", "Mahasiswa", "Status", "Nominal", "Tanggal", "Bukti TF"]
    recap.merge_cells("A1:H1")
    recap["A1"] = "REKAPITULASI PENCAIRAN MASA DEPAN JAKARTA SCHOLARSHIP BAZNAS (BAZIS) PROVINSI DKI JAKARTA"
    recap["A1"].font = Font(bold=True, color="0B6B3A", size=12)
    recap.merge_cells("A2:H2")
    recap["A2"] = f"TAHAP {stage} WILAYAH {region or 'SELURUH DKI JAKARTA'}"
    recap["A2"].font = Font(bold=True, color="0B6B3A")
    recap.append(headers)
    for cell in recap[3]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0B6B3A")
    proof_sheet = workbook.create_sheet("Bukti Transfer")
    proof_sheet.append(["Kampus", "Wilayah", "Nama Berkas", "Pratinjau Bukti Transfer"])
    for cell in proof_sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0B6B3A")
    proof_row = 2
    for index, campus in enumerate(campuses, 1):
        stage_data = campus["stage_one"] if stage == 1 else campus["stage_two"]
        recap.append([
            index,
            campus["campus"],
            ", ".join(campus["regions"]),
            "\n".join(student["name"] for student in campus["students"]),
            stage_data.get("status", "belum_diproses"),
            stage_data.get("transfer_amount") or stage_data.get("expected_amount") or 0,
            stage_data.get("disbursed_at") or "-",
            len(stage_data.get("proofs", [])),
        ])
        for proof in stage_data.get("proofs", []):
            source = await db.beneficiary_sources.find_one({"id": proof.get("source_id")}, {"_id": 0})
            proof_sheet.append([campus["campus"], ", ".join(campus["regions"]), proof.get("original_filename", "Bukti TF"), ""])
            if source and str(source.get("content_type", "")).startswith("image/"):
                try:
                    data, _ = get_object(source["storage_path"])
                    image = PILImage.open(io.BytesIO(data))
                    image.thumbnail((320, 180))
                    buffer = io.BytesIO()
                    image.convert("RGB").save(buffer, format="PNG")
                    excel_image = ExcelImage(io.BytesIO(buffer.getvalue()))
                    excel_image.width = image.width
                    excel_image.height = image.height
                    proof_sheet.add_image(excel_image, f"D{proof_row}")
                    proof_sheet.row_dimensions[proof_row].height = max(80, image.height * 0.75)
                except Exception:
                    proof_sheet.cell(proof_row, 4, "Pratinjau gambar tidak tersedia")
            elif source:
                proof_sheet.cell(proof_row, 4, "Bukti PDF tersedia sebagai lampiran sumber")
            proof_row += 1
    recap.column_dimensions["B"].width = 34
    recap.column_dimensions["C"].width = 26
    recap.column_dimensions["D"].width = 34
    proof_sheet.column_dimensions["A"].width = 30
    proof_sheet.column_dimensions["B"].width = 24
    proof_sheet.column_dimensions["C"].width = 42
    proof_sheet.column_dimensions["D"].width = 48
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=rekap-pencairan-tahap-{stage}.xlsx"},
    )


@api_router.get("/admin/disbursements/campuses/{campus_key}")
async def campus_disbursement_detail(
    campus_key: str,
    user: dict = Depends(require_roles(*MANAGEMENT_ROLES)),
):
    groups = await visible_campus_groups(user)
    group = groups.get(campus_key)
    if not group:
        raise HTTPException(status_code=404, detail="Kampus atau data pencairan tidak ditemukan.")
    region_records = group.get("records", {})
    regions = []
    for region, recipients in sorted(group["by_region"].items()):
        records = region_records.get(region, {})
        recipient_count = len(recipients)
        regions.append({
            "region": region,
            "recipient_count": recipient_count,
            "recipients": recipients,
            "stage_one": apply_recipient_financial_defaults(
                serialize_campus_disbursement(records.get(1), 1),
                recipient_count,
            ),
            "stage_two": apply_recipient_financial_defaults(
                serialize_campus_disbursement(records.get(2), 2),
                recipient_count,
            ),
        })
    return {
        "campus_key": campus_key,
        "campus": group["campus"],
        "regions": regions,
    }


@api_router.get("/admin/disbursements/campuses/{campus_key}/audit")
async def campus_disbursement_audit(
    campus_key: str,
    stage: Optional[int] = None,
    user: dict = Depends(require_roles(*MANAGEMENT_ROLES)),
):
    if stage is not None and stage not in (1, 2):
        raise HTTPException(status_code=400, detail="Tahap pencairan harus 1 atau 2.")
    group = (await visible_campus_groups(user)).get(campus_key)
    if not group:
        raise HTTPException(status_code=404, detail="Kampus atau riwayat tidak ditemukan.")
    query = {"detail.campus_key": campus_key}
    if stage is not None:
        query["detail.stage"] = stage
    scope_region = regional_admin_region(user)
    if scope_region:
        query["detail.region"] = scope_region
    return await db.beneficiary_audit_events.find(query, {"_id": 0}).sort(
        "created_at",
        -1,
    ).to_list(100)


@api_router.put("/admin/disbursements/campuses/{campus_key}/regions/{region}/stages/{stage}")
async def update_campus_disbursement(
    campus_key: str,
    region: str,
    stage: int,
    payload: DisbursementInput,
    user: dict = Depends(require_roles(*PM_MANAGERS)),
):
    if stage not in (1, 2):
        raise HTTPException(status_code=400, detail="Tahap pencairan harus 1 atau 2.")
    if payload.status not in DISBURSEMENT_STATUSES:
        raise HTTPException(status_code=400, detail="Status pencairan tidak valid.")
    if payload.amount is not None and payload.amount <= 0:
        raise HTTPException(status_code=400, detail="Nominal pencairan harus lebih dari nol.")
    all_groups = await visible_campus_groups(user)
    group = all_groups.get(campus_key)
    recipients = (group or {}).get("by_region", {}).get(region, [])
    if not recipients:
        raise HTTPException(status_code=404, detail="Wilayah pencairan tidak ditemukan.")
    record = await ensure_campus_disbursement(
        campus_key,
        group["campus"],
        region,
        stage,
        [item["user_id"] for item in recipients],
    )
    await db.campus_disbursements.update_one(
        {"id": record["id"]},
        {
            "$set": {
                "status": payload.status,
                "amount": payload.amount,
                "disbursed_at": payload.disbursed_at,
                "reference": payload.reference.strip(),
                "notes": payload.notes.strip(),
                "manually_edited": True,
                "manually_edited_at": datetime.now(timezone.utc).isoformat(),
                "manually_edited_by": user["user_id"],
                "updated_at": datetime.now(timezone.utc).isoformat(),
                "updated_by": user["user_id"],
            }
        },
    )
    await audit_beneficiary_action(
        user,
        "update_campus_disbursement",
        detail={"campus_key": campus_key, "region": region, "stage": stage},
    )
    return await campus_disbursement_detail(campus_key, user)


async def attach_transfer_to_campus_disbursement(
    actor: dict,
    source: dict,
    transaction: dict,
    beneficiary: dict,
    account_match: str,
) -> None:
    registration = beneficiary["registration"]
    profile = beneficiary["profile"]
    campus = beneficiary_campus(profile)
    campus_key = campus_disbursement_key(campus)
    region = profile_region(profile) or "Belum diisi"
    stage = int(source["stage"])
    await ensure_campus_disbursement(
        campus_key,
        campus,
        region,
        stage,
        [registration["user_id"]],
    )
    proof = {
        "source_id": source["id"],
        "original_filename": source["original_filename"],
        "attached_at": datetime.now(timezone.utc).isoformat(),
    }
    transaction_entry = {
        "source_id": source["id"],
        "user_id": registration["user_id"],
        "name": registration.get("name") or profile.get("namaLengkap", ""),
        "nim": normalize_nim(profile.get("nim")),
        "amount": str(transaction.get("amount") or ""),
        "reference": str(transaction.get("reference") or ""),
        "account_number_masked": mask_account_number(transaction.get("account_number")),
        "account_match": account_match,
    }
    aggregate = await db.campus_disbursements.find_one(
        {"campus_key": campus_key, "region": region, "stage": stage},
        {"_id": 0},
    ) or {}
    existing_transactions = aggregate.get("transactions", [])
    if not any(
        item.get("source_id") == source["id"]
        and item.get("user_id") == registration["user_id"]
        for item in existing_transactions
    ):
        await db.campus_disbursements.update_one(
            {"campus_key": campus_key, "region": region, "stage": stage},
            {
                "$addToSet": {
                    "proofs": proof,
                    "recipient_user_ids": registration["user_id"],
                },
                "$push": {"transactions": transaction_entry},
                "$set": {"updated_at": datetime.now(timezone.utc).isoformat()},
            },
        )
    await audit_beneficiary_action(
        actor,
        "attach_transfer_to_campus_disbursement",
        registration["user_id"],
        {"campus_key": campus_key, "region": region, "stage": stage, "source_id": source["id"]},
    )


async def attach_transfer_to_campus_bank_disbursement(
    actor: dict,
    source: dict,
    transaction: dict,
    candidate: dict,
    match_basis: str = "campus_bank_account",
) -> None:
    campus = candidate["campus"]
    recipient_rows = candidate["recipient_rows"]
    campus_name = campus["name"]
    campus_key = candidate["campus_key"]
    region = source["selected_region"]
    stage = int(source["stage"])
    await ensure_campus_disbursement(
        campus_key,
        campus_name,
        region,
        stage,
        [row["user_id"] for row in recipient_rows],
    )
    proof = {
        "source_id": source["id"],
        "original_filename": source["original_filename"],
        "attached_at": datetime.now(timezone.utc).isoformat(),
    }
    transaction_entry = {
        "source_id": source["id"],
        "campus_id": campus["id"],
        "campus": campus_name,
        "account_number_masked": mask_account_number(transaction.get("account_number")),
        "account_holder_name": str(transaction.get("account_holder_name") or ""),
        "amount": str(transaction.get("amount") or ""),
        "reference": str(transaction.get("reference") or ""),
        "date": str(transaction.get("date") or ""),
        "transaction_status": str(transaction.get("status") or ""),
        "match_basis": match_basis,
    }
    aggregate = await db.campus_disbursements.find_one(
        {"campus_key": campus_key, "region": region, "stage": stage},
        {"_id": 0},
    ) or {}
    existing_transactions = aggregate.get("transactions", [])
    if not any(
        item.get("source_id") == source["id"]
        and item.get("campus_id") == campus["id"]
        for item in existing_transactions
    ):
        await db.campus_disbursements.update_one(
            {"campus_key": campus_key, "region": region, "stage": stage},
            {
                "$addToSet": {"proofs": proof},
                "$push": {"transactions": transaction_entry},
                "$set": {"updated_at": datetime.now(timezone.utc).isoformat()},
            },
        )
    await refresh_campus_transfer_financials(
        campus_key,
        region,
        stage,
        len(recipient_rows),
    )
    for row in recipient_rows:
        existing_notice = await db.notifications.find_one(
            {
                "user_id": row["user_id"],
                "type": "disbursement_proof",
                "source_id": source["id"],
            },
            {"_id": 0, "id": 1},
        )
        if existing_notice:
            continue
        await db.notifications.insert_one({
            "id": str(uuid.uuid4()),
            "user_id": row["user_id"],
            "type": "disbursement_proof",
            "source_id": source["id"],
            "stage": stage,
            "campus": campus_name,
            "message": (
                f"Bukti transfer Pencairan Tahap {'I' if stage == 1 else 'II'} "
                f"untuk {campus_name} telah tersedia."
            ),
            "is_read": False,
            "created_at": datetime.now(timezone.utc).isoformat(),
        })
    await audit_beneficiary_action(
        actor,
        "attach_transfer_by_campus_bank",
        detail={
            "campus_id": campus["id"],
            "campus_key": campus_key,
            "region": region,
            "stage": stage,
            "source_id": source["id"],
            "match_basis": match_basis,
        },
    )


@api_router.get("/student/disbursement-proofs")
async def list_student_disbursement_proofs(
    user: dict = Depends(require_roles("student")),
):
    records = await db.campus_disbursements.find(
        {"recipient_user_ids": user["user_id"]},
        {"_id": 0, "campus": 1, "region": 1, "stage": 1, "proofs": 1},
    ).to_list(200)
    proof_items = {}
    for record in records:
        for proof in record.get("proofs", []):
            source_id = proof.get("source_id")
            if not source_id:
                continue
            item = proof_items.setdefault(source_id, {
                "source_id": source_id,
                "filename": proof.get("original_filename", "Bukti Transfer"),
                "campuses": [],
            })
            item["campuses"].append({
                "campus": record.get("campus", "-"),
                "region": record.get("region", "-"),
                "stage": record.get("stage"),
            })
    return list(proof_items.values())


@api_router.get("/student/disbursement-proofs/{source_id}")
async def download_student_disbursement_proof(
    source_id: str,
    user: dict = Depends(require_roles("student")),
):
    access = await db.campus_disbursements.find_one(
        {"recipient_user_ids": user["user_id"], "proofs.source_id": source_id},
        {"_id": 0, "id": 1},
    )
    if not access:
        raise HTTPException(status_code=404, detail="Bukti transfer tidak ditemukan.")
    source = await db.beneficiary_sources.find_one(
        {"id": source_id, "source_type": "transfer_proof"},
        {"_id": 0},
    )
    if not source:
        raise HTTPException(status_code=404, detail="Berkas bukti transfer tidak ditemukan.")
    data, content_type = get_object(source["storage_path"])
    await audit_beneficiary_action(
        user,
        "student_download_disbursement_proof",
        user["user_id"],
        {"source_id": source_id},
    )
    return StarletteResponse(
        content=data,
        media_type=source.get("content_type", content_type),
        headers={"Content-Disposition": f'inline; filename="{source["original_filename"]}"'},
    )


async def refresh_campus_transfer_financials(
    campus_key: str,
    region: str,
    stage: int,
    recipient_count: int,
) -> None:
    aggregate = await db.campus_disbursements.find_one(
        {"campus_key": campus_key, "region": region, "stage": stage},
        {"_id": 0},
    ) or {}
    amounts = [
        amount
        for amount in (parse_transfer_amount(item.get("amount")) for item in aggregate.get("transactions", []))
        if amount is not None
    ]
    expected_amount = recipient_count * TRANSFER_AMOUNT_PER_STUDENT
    if not amounts:
        financials = {
            "transfer_amount": None,
            "expected_amount": expected_amount,
            "paid_student_count": 0,
            "recipient_count": recipient_count,
            "payment_assessment": "belum_ada_transfer",
            "payment_difference": None,
        }
    else:
        transfer_amount = sum(amounts)
        payment_difference = transfer_amount - expected_amount
        if 0 <= payment_difference <= UNIQUE_TRANSFER_TOLERANCE:
            assessment = "sesuai"
        elif payment_difference < 0:
            assessment = "kurang"
        else:
            assessment = "lebih"
        financials = {
            "amount": transfer_amount,
            "transfer_amount": transfer_amount,
            "expected_amount": expected_amount,
            "paid_student_count": min(
                transfer_amount // TRANSFER_AMOUNT_PER_STUDENT,
                recipient_count,
            ),
            "recipient_count": recipient_count,
            "payment_assessment": assessment,
            "payment_difference": payment_difference,
        }
    await db.campus_disbursements.update_one(
        {"campus_key": campus_key, "region": region, "stage": stage},
        {"$set": {**financials, "updated_at": datetime.now(timezone.utc).isoformat()}},
    )


async def get_final_beneficiary(user: dict, user_id: str) -> dict:
    await require_participant_scope(user, user_id)
    registration = await db.registrations.find_one(
        {"user_id": user_id, "status": "lolos"},
        {"_id": 0},
    )
    if not registration:
        raise HTTPException(status_code=404, detail="Penerima Manfaat tidak ditemukan.")
    profile = await db.profiles.find_one({"user_id": user_id}, {"_id": 0}) or {}
    return {"registration": registration, "profile": profile.get("data", {})}


async def ensure_disbursement(user_id: str, stage: int) -> dict:
    await db.beneficiary_disbursements.update_one(
        {"user_id": user_id, "stage": stage},
        {
            "$setOnInsert": {
                "id": str(uuid.uuid4()),
                "user_id": user_id,
                "stage": stage,
                "status": "belum_diproses",
                "amount": None,
                "disbursed_at": None,
                "reference": "",
                "notes": "",
                "proofs": [],
                "created_at": datetime.now(timezone.utc).isoformat(),
            }
        },
        upsert=True,
    )
    return await db.beneficiary_disbursements.find_one(
        {"user_id": user_id, "stage": stage},
        {"_id": 0},
    )


def serialize_disbursement(record: Optional[dict], stage: int) -> dict:
    data = record or {"stage": stage, "status": "belum_diproses", "proofs": []}
    return {
        "stage": stage,
        "status": data.get("status", "belum_diproses"),
        "amount": data.get("amount"),
        "disbursed_at": data.get("disbursed_at"),
        "reference": data.get("reference", ""),
        "notes": data.get("notes", ""),
        "proofs": data.get("proofs", []),
        "updated_at": data.get("updated_at"),
    }


async def beneficiary_view(user: dict, user_id: str) -> dict:
    beneficiary = await get_final_beneficiary(user, user_id)
    registration = beneficiary["registration"]
    profile = beneficiary["profile"]
    record = await db.beneficiary_records.find_one({"user_id": user_id}, {"_id": 0}) or {}
    stages = await db.beneficiary_disbursements.find(
        {"user_id": user_id},
        {"_id": 0},
    ).to_list(2)
    by_stage = {item.get("stage"): item for item in stages}
    return {
        "user_id": user_id,
        "name": registration.get("name") or profile.get("namaLengkap", "-"),
        "email": registration.get("email") or profile.get("email", "-"),
        "cpm_id": registration.get("cpm_id", "-"),
        "nim": profile.get("nim", "-"),
        "campus": profile.get("institusi", "-"),
        "study_program": profile.get("jurusan", "-"),
        "semester": profile.get("semester", "-"),
        "region": profile_region(profile) or "Belum diisi",
        "bank": {
            "bank_name": record.get("bank_name", ""),
            "account_number_masked": mask_account_number(
                decrypt_account_number(record.get("bank_account_cipher", ""))
            ),
            "is_recorded": bool(record.get("bank_account_cipher")),
        },
        "active_letters": record.get("active_letters", []),
        "disbursements": [
            serialize_disbursement(by_stage.get(1), 1),
            serialize_disbursement(by_stage.get(2), 2),
        ],
    }


async def create_beneficiary_review(
    actor: dict,
    source: dict,
    kind: str,
    payload: dict,
    reason: str,
) -> None:
    review = {
        "id": str(uuid.uuid4()),
        "source_id": source["id"],
        "source_type": source["source_type"],
        "stage": source.get("stage"),
        "selected_region": source.get("selected_region"),
        "kind": kind,
        "payload": payload,
        "reason": reason,
        "status": "pending",
        "created_at": datetime.now(timezone.utc).isoformat(),
    }
    await db.beneficiary_reviews.insert_one(review)
    await audit_beneficiary_action(actor, "create_pm_review", detail={"review_id": review["id"]})


@api_router.get("/admin/beneficiaries")
async def list_beneficiaries(
    region: Optional[str] = None,
    search: Optional[str] = None,
    user: dict = Depends(require_roles(*MANAGEMENT_ROLES)),
):
    scope_region = regional_admin_region(user)
    query = {"status": "lolos"}
    query.update(await participant_scope_query(user))
    registrations = await db.registrations.find(query, {"_id": 0}).sort("created_at", -1).to_list(2000)
    user_ids = [registration["user_id"] for registration in registrations]
    profiles = await db.profiles.find(
        {"user_id": {"$in": user_ids}},
        {"_id": 0, "user_id": 1, "data": 1},
    ).to_list(2000)
    records = await db.beneficiary_records.find(
        {"user_id": {"$in": user_ids}},
        {"_id": 0},
    ).to_list(2000)
    disbursements = await db.beneficiary_disbursements.find(
        {"user_id": {"$in": user_ids}},
        {"_id": 0},
    ).to_list(4000)
    profiles_by_user = {item["user_id"]: item.get("data", {}) for item in profiles}
    records_by_user = {item["user_id"]: item for item in records}
    disbursements_by_user: Dict[str, Dict[int, dict]] = {}
    for item in disbursements:
        disbursements_by_user.setdefault(item["user_id"], {})[item["stage"]] = item
    selected_region = scope_region or region
    result = []
    for registration in registrations:
        profile = profiles_by_user.get(registration["user_id"], {})
        beneficiary_region = profile_region(profile) or "Belum diisi"
        if selected_region and selected_region != "all" and beneficiary_region != selected_region:
            continue
        record = records_by_user.get(registration["user_id"], {})
        keyword = " ".join([
            str(registration.get("name", "")),
            str(registration.get("email", "")),
            str(profile.get("nim", "")),
            str(profile.get("institusi", "")),
        ]).lower()
        if search and search.strip().lower() not in keyword:
            continue
        stages = disbursements_by_user.get(registration["user_id"], {})
        result.append({
            "user_id": registration["user_id"],
            "name": registration.get("name") or profile.get("namaLengkap", "-"),
            "email": registration.get("email") or profile.get("email", "-"),
            "cpm_id": registration.get("cpm_id", "-"),
            "nim": profile.get("nim", "-"),
            "campus": profile.get("institusi", "-"),
            "study_program": profile.get("jurusan", "-"),
            "semester": profile.get("semester", "-"),
            "region": beneficiary_region,
            "active_letter_count": len(record.get("active_letters", [])),
            "bank_recorded": bool(record.get("bank_account_cipher")),
            "stage_one": serialize_disbursement(stages.get(1), 1),
            "stage_two": serialize_disbursement(stages.get(2), 2),
        })
    return result


@api_router.get("/admin/beneficiaries/{user_id}")
async def beneficiary_detail(
    user_id: str,
    user: dict = Depends(require_roles(*MANAGEMENT_ROLES)),
):
    return await beneficiary_view(user, user_id)


@api_router.put("/admin/beneficiaries/{user_id}/bank")
async def update_beneficiary_bank(
    user_id: str,
    payload: BeneficiaryBankInput,
    user: dict = Depends(require_roles(*PM_MANAGERS)),
):
    await get_final_beneficiary(user, user_id)
    account_number = normalize_account_number(payload.account_number)
    if account_number and not 6 <= len(account_number) <= 24:
        raise HTTPException(status_code=400, detail="Nomor rekening harus terdiri dari 6–24 digit.")
    now = datetime.now(timezone.utc).isoformat()
    update = {
        "bank_name": payload.bank_name.strip(),
        "updated_at": now,
        "updated_by": user["user_id"],
    }
    if account_number:
        update["bank_account_cipher"] = encrypt_account_number(account_number)
    await db.beneficiary_records.update_one(
        {"user_id": user_id},
        {"$set": update, "$setOnInsert": {"user_id": user_id, "active_letters": []}},
        upsert=True,
    )
    await audit_beneficiary_action(user, "update_beneficiary_bank", user_id)
    return await beneficiary_view(user, user_id)


@api_router.put("/admin/beneficiaries/{user_id}/disbursements/{stage}")
async def update_disbursement(
    user_id: str,
    stage: int,
    payload: DisbursementInput,
    user: dict = Depends(require_roles(*PM_MANAGERS)),
):
    if stage not in (1, 2):
        raise HTTPException(status_code=400, detail="Tahap pencairan harus 1 atau 2.")
    if payload.status not in DISBURSEMENT_STATUSES:
        raise HTTPException(status_code=400, detail="Status pencairan tidak valid.")
    if payload.amount is not None and payload.amount <= 0:
        raise HTTPException(status_code=400, detail="Nominal pencairan harus lebih dari nol.")
    await get_final_beneficiary(user, user_id)
    await ensure_disbursement(user_id, stage)
    await db.beneficiary_disbursements.update_one(
        {"user_id": user_id, "stage": stage},
        {
            "$set": {
                "status": payload.status,
                "amount": payload.amount,
                "disbursed_at": payload.disbursed_at,
                "reference": payload.reference.strip(),
                "notes": payload.notes.strip(),
                "updated_at": datetime.now(timezone.utc).isoformat(),
                "updated_by": user["user_id"],
            }
        },
    )
    await audit_beneficiary_action(
        user,
        "update_disbursement",
        user_id,
        {"stage": stage, "status": payload.status},
    )
    return await beneficiary_view(user, user_id)


@api_router.get("/admin/beneficiaries/{user_id}/audit")
async def beneficiary_audit(
    user_id: str,
    user: dict = Depends(require_roles(*MANAGEMENT_ROLES)),
):
    await get_final_beneficiary(user, user_id)
    return await db.beneficiary_audit_events.find(
        {"user_id": user_id},
        {"_id": 0},
    ).sort("created_at", -1).to_list(100)


@api_router.get("/admin/beneficiary-reviews")
async def list_beneficiary_reviews(
    stage: Optional[int] = None,
    region: Optional[str] = None,
    user: dict = Depends(require_roles(*PM_MANAGERS)),
):
    if stage is not None and stage not in (1, 2):
        raise HTTPException(status_code=400, detail="Tahap pencairan harus 1 atau 2.")
    if region is not None and region not in DKI_REGIONS:
        raise HTTPException(status_code=400, detail="Wilayah pencairan tidak valid.")
    query = {"status": "pending"}
    if stage is not None:
        query["stage"] = stage
    if region is not None:
        query["selected_region"] = region
    return await db.beneficiary_reviews.find(
        query,
        {"_id": 0},
    ).sort("created_at", -1).to_list(300)


@api_router.get("/admin/beneficiary-files/{source_id}")
async def download_beneficiary_file(
    source_id: str,
    user: dict = Depends(require_roles(*PM_MANAGERS)),
):
    source = await db.beneficiary_sources.find_one({"id": source_id}, {"_id": 0})
    if not source:
        raise HTTPException(status_code=404, detail="Berkas Penerima Manfaat tidak ditemukan.")
    data, content_type = get_object(source["storage_path"])
    headers = {"Content-Disposition": f'inline; filename="{source["original_filename"]}"'}
    await audit_beneficiary_action(user, "download_beneficiary_source", detail={"source_id": source_id})
    return StarletteResponse(
        content=data,
        media_type=source.get("content_type", content_type),
        headers=headers,
    )


async def attach_active_letter(
    actor: dict,
    source: dict,
    student: dict,
    beneficiary: dict,
) -> bool:
    registration = beneficiary["registration"]
    profile = beneficiary["profile"]
    user_id = registration["user_id"]
    letter = {
        "source_id": source["id"],
        "original_filename": source["original_filename"],
        "matched_at": datetime.now(timezone.utc).isoformat(),
        "matched_name": student.get("name", ""),
        "matched_nim": normalize_nim(student.get("nim")),
        "study_program": student.get("study_program", ""),
        "semester": student.get("semester", ""),
        "evidence": student.get("evidence", ""),
    }
    record = await db.beneficiary_records.find_one({"user_id": user_id}, {"_id": 0}) or {}
    if any(item.get("source_id") == source["id"] for item in record.get("active_letters", [])):
        return False
    await db.beneficiary_records.update_one(
        {"user_id": user_id},
        {
            "$setOnInsert": {
                "user_id": user_id,
                "active_letters": [],
            },
            "$push": {"active_letters": letter},
            "$set": {"updated_at": datetime.now(timezone.utc).isoformat()},
        },
        upsert=True,
    )
    await audit_beneficiary_action(
        actor,
        "attach_active_letter",
        user_id,
        {"source_id": source["id"], "region": profile_region(profile)},
    )
    return True


@api_router.post("/admin/beneficiaries/active-letters")
async def upload_active_letters(
    files: List[UploadFile] = File(...),
    user: dict = Depends(require_roles(*PM_MANAGERS)),
):
    if not files or len(files) > 20:
        raise HTTPException(status_code=400, detail="Unggah antara 1 hingga 20 berkas surat.")
    beneficiary_index = await final_beneficiary_index()
    summary = {"uploaded": 0, "matched": 0, "duplicate": 0, "review": 0, "failed": 0}
    sources = []
    for file in files:
        filename, ext, data, content_type = await read_pm_upload(file)
        source = await store_pm_source(
            user,
            "active_letter",
            filename,
            ext,
            data,
            content_type,
        )
        summary["uploaded"] += 1
        try:
            extracted = await extract_pm_payload(
                data,
                content_type,
                user,
                PM_ACTIVE_LETTER_PROMPT,
                "surat keterangan mahasiswa aktif",
            )
            students = extracted.get("students", [])
            if not isinstance(students, list):
                students = []
            matched_user_ids = []
            regions = set()
            for raw_student in students:
                student = raw_student if isinstance(raw_student, dict) else {}
                name = normalized_name(student.get("name"))
                nim = normalize_nim(student.get("nim"))
                if not name or not nim:
                    await create_beneficiary_review(
                        user,
                        source,
                        "active_letter_incomplete",
                        student,
                        "Nama atau NIM tidak terbaca lengkap.",
                    )
                    summary["review"] += 1
                    continue
                matches = beneficiary_index.get(f"{name}|{nim}", [])
                if len(matches) != 1:
                    reason = (
                        "Mahasiswa belum berstatus lulus tahap akhir."
                        if not matches
                        else "Data nama dan NIM cocok ke lebih dari satu Penerima Manfaat."
                    )
                    await create_beneficiary_review(
                        user,
                        source,
                        "active_letter_unmatched",
                        student,
                        reason,
                    )
                    summary["review"] += 1
                    continue
                beneficiary = matches[0]
                attached = await attach_active_letter(user, source, student, beneficiary)
                if attached:
                    matched_user_ids.append(beneficiary["registration"]["user_id"])
                    regions.add(beneficiary["region"])
                    summary["matched"] += 1
                else:
                    await create_beneficiary_review(
                        user,
                        source,
                        "active_letter_duplicate",
                        student,
                        "Surat ini sudah tertempel pada Penerima Manfaat tersebut.",
                    )
                    summary["duplicate"] += 1
            status = "processed" if students else "needs_review"
            await db.beneficiary_sources.update_one(
                {"id": source["id"]},
                {
                    "$set": {
                        "status": status,
                        "extraction": {"students": students},
                        "matched_user_ids": matched_user_ids,
                        "regions": sorted(regions),
                        "processed_at": datetime.now(timezone.utc).isoformat(),
                    }
                },
            )
        except HTTPException as error:
            await db.beneficiary_sources.update_one(
                {"id": source["id"]},
                {"$set": {"status": "failed", "error": error.detail}},
            )
            await create_beneficiary_review(
                user,
                source,
                "active_letter_ocr_failed",
                {},
                str(error.detail),
            )
            summary["failed"] += 1
        sources.append({"id": source["id"], "filename": filename})
    return {"summary": summary, "sources": sources}


async def active_letter_candidate_index(workflow: str) -> Dict[str, List[dict]]:
    statuses = ["wawancara", "verifikasi_faktual"] if workflow == "factual_verification" else ["lolos"]
    registrations = await db.registrations.find(
        {"status": {"$in": statuses}},
        {"_id": 0},
    ).to_list(5000)
    user_ids = [item["user_id"] for item in registrations]
    profiles = await db.profiles.find(
        {"user_id": {"$in": user_ids}},
        {"_id": 0, "user_id": 1, "data": 1},
    ).to_list(5000)
    profiles_by_user = {item["user_id"]: item.get("data", {}) for item in profiles}
    index: Dict[str, List[dict]] = {}
    for registration in registrations:
        profile = profiles_by_user.get(registration["user_id"], {})
        name = registration.get("name") or profile.get("namaLengkap", "")
        nim = normalize_nim(profile.get("nim"))
        if not normalized_name(name) or not nim:
            continue
        index.setdefault(f"{normalized_name(name)}|{nim}", []).append({
            "registration": registration,
            "profile": profile,
            "region": profile_region(profile),
        })
    return index


async def create_active_letter_recommendation(
    source: dict,
    workflow: str,
    student: dict,
    match: Optional[dict],
    recommendation: str,
    reason: str,
) -> None:
    document = {
        "id": str(uuid.uuid4()),
        "source_id": source["id"],
        "workflow": workflow,
        "campus": source.get("extracted_campus", ""),
        "student": student,
        "user_id": match["registration"]["user_id"] if match else None,
        "region": match.get("region") if match else None,
        "recommendation": recommendation,
        "reason": reason,
        "status": "pending",
        "created_at": datetime.now(timezone.utc).isoformat(),
    }
    await db.active_letter_approvals.insert_one(document)


@api_router.post("/admin/active-letter-approvals/upload")
async def upload_active_letter_for_approval(
    workflow: str = Form(...),
    files: List[UploadFile] = File(...),
    user: dict = Depends(require_roles(*PM_MANAGERS)),
):
    if workflow not in {"factual_verification", "stage_ii_disbursement"}:
        raise HTTPException(status_code=400, detail="Jenis proses surat aktif tidak valid.")
    if not files or len(files) > 20:
        raise HTTPException(status_code=400, detail="Unggah antara 1 hingga 20 surat aktif.")
    candidate_index = await active_letter_candidate_index(workflow)
    summary = {"uploaded": 0, "recommended": 0, "quarantined": 0, "failed": 0}
    sources = []
    for file in files:
        filename, ext, data, content_type = await read_pm_upload(file)
        source = await store_pm_source(
            user,
            "active_letter_approval",
            filename,
            ext,
            data,
            content_type,
        )
        summary["uploaded"] += 1
        try:
            extracted = await extract_pm_payload(
                data,
                content_type,
                user,
                PM_ACTIVE_LETTER_PROMPT,
                "surat keterangan mahasiswa aktif untuk persetujuan",
            )
            campus = str(extracted.get("campus") or "").strip()
            source["extracted_campus"] = campus
            students = extracted.get("students", [])
            if not isinstance(students, list):
                students = []
            matched_user_ids = []
            for raw_student in students:
                student = raw_student if isinstance(raw_student, dict) else {}
                name = normalized_name(student.get("name"))
                nim = normalize_nim(student.get("nim"))
                if not campus or not name or not nim:
                    await create_active_letter_recommendation(
                        source,
                        workflow,
                        student,
                        None,
                        "quarantined",
                        "Nama kampus, nama mahasiswa, atau NIM tidak terbaca lengkap.",
                    )
                    summary["quarantined"] += 1
                    continue
                matches = [
                    item
                    for item in candidate_index.get(f"{name}|{nim}", [])
                    if normalized_name(beneficiary_campus(item["profile"])) == normalized_name(campus)
                ]
                if len(matches) != 1:
                    reason = (
                        "Mahasiswa tidak ditemukan pada proses atau kampus yang sesuai."
                        if not matches
                        else "Mahasiswa cocok ke lebih dari satu data penerima."
                    )
                    await create_active_letter_recommendation(
                        source,
                        workflow,
                        student,
                        None,
                        "quarantined",
                        reason,
                    )
                    summary["quarantined"] += 1
                    continue
                match = matches[0]
                await create_active_letter_recommendation(
                    source,
                    workflow,
                    student,
                    match,
                    "recommended",
                    "Cocok berdasarkan kampus, nama, dan NIM pada surat aktif.",
                )
                matched_user_ids.append(match["registration"]["user_id"])
                summary["recommended"] += 1
            await db.beneficiary_sources.update_one(
                {"id": source["id"]},
                {
                    "$set": {
                        "workflow": workflow,
                        "status": "processed" if students else "needs_review",
                        "extraction": {"campus": campus, "students": students},
                        "extracted_campus": campus,
                        "matched_user_ids": matched_user_ids,
                        "processed_at": datetime.now(timezone.utc).isoformat(),
                    }
                },
            )
        except HTTPException as error:
            await db.beneficiary_sources.update_one(
                {"id": source["id"]},
                {"$set": {"status": "failed", "error": error.detail, "workflow": workflow}},
            )
            summary["failed"] += 1
        sources.append({"id": source["id"], "filename": filename})
    await audit_beneficiary_action(user, "upload_active_letter_for_approval", detail={"workflow": workflow})
    return {"summary": summary, "sources": sources}


@api_router.get("/admin/active-letter-approvals")
async def list_active_letter_approvals(
    workflow: Optional[str] = None,
    user: dict = Depends(require_roles(*PM_MANAGERS)),
):
    query = {"status": "pending"}
    if workflow:
        query["workflow"] = workflow
    return await db.active_letter_approvals.find(query, {"_id": 0}).sort(
        "created_at",
        -1,
    ).to_list(5000)


@api_router.post("/admin/active-letter-approvals/decision")
async def decide_active_letter_approvals(
    payload: ActiveLetterDecisionInput,
    user: dict = Depends(require_roles(*PM_MANAGERS)),
):
    if payload.action not in {"approve", "reject"}:
        raise HTTPException(status_code=400, detail="Keputusan surat aktif tidak valid.")
    if payload.workflow not in {"factual_verification", "stage_ii_disbursement"}:
        raise HTTPException(status_code=400, detail="Jenis proses surat aktif tidak valid.")
    query = {
        "status": "pending",
        "recommendation": "recommended",
        "workflow": payload.workflow,
    }
    if payload.apply_all:
        if not payload.source_ids:
            raise HTTPException(
                status_code=400,
                detail="Persetujuan semua harus dibatasi ke berkas sumber yang dipilih.",
            )
        query["source_id"] = {"$in": payload.source_ids}
    else:
        if not payload.recommendation_ids:
            raise HTTPException(status_code=400, detail="Pilih rekomendasi untuk diproses.")
        query["id"] = {"$in": payload.recommendation_ids}
    recommendations = await db.active_letter_approvals.find(query, {"_id": 0}).to_list(5000)
    if not recommendations:
        raise HTTPException(status_code=404, detail="Tidak ada rekomendasi yang dapat diproses.")
    now = datetime.now(timezone.utc).isoformat()
    processed = 0
    skipped = 0
    for recommendation in recommendations:
        user_id = recommendation.get("user_id")
        if payload.action == "approve" and user_id:
            if recommendation["workflow"] == "factual_verification":
                update_result = await db.registrations.update_one(
                    {
                        "user_id": user_id,
                        "status": {"$in": ["wawancara", "verifikasi_faktual"]},
                    },
                    {
                        "$set": {
                            "status": "lolos",
                            "factual_verification_status": "approved",
                            "factual_verification_source_id": recommendation["source_id"],
                            "updated_at": now,
                        }
                    },
                )
                if not update_result.modified_count:
                    await db.active_letter_approvals.update_one(
                        {"id": recommendation["id"], "status": "pending"},
                        {
                            "$set": {
                                "status": "stale",
                                "stale_at": now,
                                "stale_reason": "Status peserta telah berubah sebelum persetujuan.",
                            }
                        },
                    )
                    skipped += 1
                    continue
            else:
                await db.stage_ii_eligibilities.update_one(
                    {"user_id": user_id},
                    {
                        "$set": {
                            "status": "approved",
                            "source_id": recommendation["source_id"],
                            "approved_at": now,
                            "approved_by": user["user_id"],
                        }
                    },
                    upsert=True,
                )
            message = (
                "Surat aktif kampus Anda disetujui untuk Verifikasi Faktual."
                if recommendation["workflow"] == "factual_verification"
                else "Surat aktif kampus Anda disetujui untuk kelayakan Pencairan Tahap II."
            )
            await db.notifications.insert_one({
                "id": str(uuid.uuid4()),
                "user_id": user_id,
                "message": message,
                "is_read": False,
                "created_at": now,
            })
        await db.active_letter_approvals.update_one(
            {"id": recommendation["id"], "status": "pending"},
            {
                "$set": {
                    "status": "approved" if payload.action == "approve" else "rejected",
                    "decided_by": user["user_id"],
                    "decided_at": now,
                }
            },
        )
        await audit_beneficiary_action(
            user,
            f"{payload.action}_active_letter_recommendation",
            user_id,
            {"recommendation_id": recommendation["id"], "workflow": recommendation["workflow"]},
        )
        processed += 1
    return {"processed": processed, "skipped": skipped, "action": payload.action}


async def attach_transfer_transaction(
    actor: dict,
    source: dict,
    transaction: dict,
    beneficiary: dict,
) -> None:
    user_id = beneficiary["registration"]["user_id"]
    stage = int(source["stage"])
    record = await db.beneficiary_records.find_one({"user_id": user_id}, {"_id": 0}) or {}
    saved_account = decrypt_account_number(record.get("bank_account_cipher", ""))
    extracted_account = normalize_account_number(transaction.get("account_number"))
    if extracted_account and saved_account:
        account_match = "sesuai" if extracted_account == saved_account else "tidak_sesuai"
    else:
        account_match = "perlu_tinjau"
    await ensure_disbursement(user_id, stage)
    proof = {
        "source_id": source["id"],
        "original_filename": source["original_filename"],
        "account_number_masked": mask_account_number(extracted_account),
        "bank_name": str(transaction.get("bank_name") or ""),
        "transaction_status": str(transaction.get("status") or ""),
        "amount": str(transaction.get("amount") or ""),
        "reference": str(transaction.get("reference") or ""),
        "date": str(transaction.get("date") or ""),
        "account_match": account_match,
        "evidence": str(transaction.get("evidence") or ""),
        "attached_at": datetime.now(timezone.utc).isoformat(),
    }
    disbursement = await db.beneficiary_disbursements.find_one(
        {"user_id": user_id, "stage": stage},
        {"_id": 0},
    ) or {}
    if not any(item.get("source_id") == source["id"] for item in disbursement.get("proofs", [])):
        await db.beneficiary_disbursements.update_one(
            {"user_id": user_id, "stage": stage},
            {
                "$push": {"proofs": proof},
                "$set": {"updated_at": datetime.now(timezone.utc).isoformat()},
            },
        )
    await audit_beneficiary_action(
        actor,
        "attach_transfer_proof",
        user_id,
        {"source_id": source["id"], "stage": stage, "account_match": account_match},
    )
    await attach_transfer_to_campus_disbursement(
        actor,
        source,
        transaction,
        beneficiary,
        account_match,
    )


@api_router.post("/admin/disbursements/transfer-proofs")
@api_router.post("/admin/beneficiaries/disbursement-proofs")
async def upload_disbursement_proofs(
    stage: int = Form(...),
    region: Optional[str] = Form(None),
    files: List[UploadFile] = File(...),
    user: dict = Depends(require_roles(*PM_MANAGERS)),
):
    if stage not in (1, 2):
        raise HTTPException(status_code=400, detail="Tahap pencairan harus 1 atau 2.")
    region = (region or "").strip()
    if region not in DKI_REGIONS:
        raise HTTPException(status_code=400, detail="Pilih wilayah pencairan yang valid.")
    if not files or len(files) > 20:
        raise HTTPException(status_code=400, detail="Unggah antara 1 hingga 20 bukti transfer.")
    campus_candidates = await campus_bank_candidates_for_region(region)
    summary = {"uploaded": 0, "matched": 0, "review": 0, "failed": 0}
    sources = []
    for file in files:
        filename, ext, data, content_type = await read_pm_upload(file)
        source = await store_pm_source(
            user,
            "transfer_proof",
            filename,
            ext,
            data,
            content_type,
            stage,
            region,
        )
        summary["uploaded"] += 1
        try:
            extracted = await extract_pm_payload(
                data,
                content_type,
                user,
                f"{PM_TRANSFER_PROOF_PROMPT}\n\nKONTEKS WILAYAH UNGGAHAN: {region}. "
                "Prioritaskan pembacaan transaksi untuk kampus di wilayah ini. "
                "Jangan menyimpulkan wilayah lain sebagai kecocokan otomatis.",
                "bukti transfer",
            )
            transactions = extracted.get("transactions", [])
            if not isinstance(transactions, list):
                transactions = []
            matched_user_ids = []
            matched_campus_ids = []
            regions = set()
            for raw_transaction in transactions:
                transaction = raw_transaction if isinstance(raw_transaction, dict) else {}
                campus_name = normalized_name(transaction.get("campus"))
                account_number = normalize_account_number(transaction.get("account_number"))
                account_holder_name = normalized_name(transaction.get("account_holder_name"))
                if not campus_name or not account_number or not account_holder_name:
                    await create_beneficiary_review(
                        user,
                        source,
                        "transfer_campus_incomplete",
                        transaction,
                        "Nama kampus, nomor rekening, atau nama pemilik rekening belum lengkap.",
                    )
                    summary["review"] += 1
                    continue
                matches = match_campus_bank_transaction(transaction, campus_candidates)
                if len(matches) != 1:
                    reason = (
                        f"Data kampus tidak cocok dengan rekening master wilayah unggahan {region}."
                        if not matches
                        else "Data rekening cocok ke lebih dari satu master kampus."
                    )
                    await create_beneficiary_review(
                        user,
                        source,
                        "transfer_campus_unmatched",
                        transaction,
                        reason,
                    )
                    summary["review"] += 1
                    continue
                candidate = matches[0]
                await attach_transfer_to_campus_bank_disbursement(
                    user,
                    source,
                    transaction,
                    candidate,
                )
                matched_user_ids.extend(row["user_id"] for row in candidate["recipient_rows"])
                matched_campus_ids.append(candidate["campus"]["id"])
                regions.add(region)
                summary["matched"] += 1
            status = "processed" if transactions else "needs_review"
            await db.beneficiary_sources.update_one(
                {"id": source["id"]},
                {
                    "$set": {
                        "status": status,
                        "extraction": {
                            "transactions": transactions,
                            "summary": extracted.get("summary", {}),
                        },
                        "matched_user_ids": matched_user_ids,
                        "matched_campus_ids": matched_campus_ids,
                        "regions": sorted(regions),
                        "processed_at": datetime.now(timezone.utc).isoformat(),
                    }
                },
            )
        except HTTPException as error:
            await db.beneficiary_sources.update_one(
                {"id": source["id"]},
                {"$set": {"status": "failed", "error": error.detail}},
            )
            await create_beneficiary_review(
                user,
                source,
                "transfer_ocr_failed",
                {},
                str(error.detail),
            )
            summary["failed"] += 1
        sources.append({"id": source["id"], "filename": filename})
    return {"summary": summary, "sources": sources}


@api_router.post("/admin/beneficiary-reviews/{review_id}/resolve")
async def resolve_beneficiary_review(
    review_id: str,
    payload: BeneficiaryReviewResolveInput,
    user: dict = Depends(require_roles(*PM_MANAGERS)),
):
    review = await db.beneficiary_reviews.find_one(
        {"id": review_id, "status": "pending"},
        {"_id": 0},
    )
    if not review:
        raise HTTPException(status_code=404, detail="Item tinjauan tidak ditemukan.")
    source = await db.beneficiary_sources.find_one({"id": review["source_id"]}, {"_id": 0})
    if not source:
        raise HTTPException(status_code=404, detail="Berkas sumber tidak ditemukan.")
    selected_region = source.get("selected_region")
    if source["source_type"] == "active_letter":
        if not payload.user_id:
            raise HTTPException(status_code=400, detail="Pilih mahasiswa untuk surat aktif.")
        beneficiary = await get_final_beneficiary(user, payload.user_id)
        if selected_region and profile_region(beneficiary["profile"]) != selected_region:
            raise HTTPException(
                status_code=400,
                detail="Penerima harus berasal dari wilayah yang dipilih saat unggah.",
            )
        await attach_active_letter(user, source, review.get("payload", {}), beneficiary)
        resolved_user_id = payload.user_id
        resolved_campus_id = None
    elif source["source_type"] == "transfer_proof":
        if not payload.campus_id:
            raise HTTPException(status_code=400, detail="Pilih kampus untuk bukti transfer.")
        if not selected_region:
            raise HTTPException(status_code=400, detail="Wilayah sumber bukti transfer tidak tersedia.")
        candidates = await campus_bank_candidates_for_region(selected_region)
        candidate = next(
            (item for item in candidates if item["campus"].get("id") == payload.campus_id),
            None,
        )
        if not candidate:
            raise HTTPException(
                status_code=400,
                detail="Kampus tidak memiliki Penerima Manfaat pada wilayah unggahan.",
            )
        await attach_transfer_to_campus_bank_disbursement(
            user,
            source,
            review.get("payload", {}),
            candidate,
            "manual_campus_confirmation",
        )
        resolved_user_id = None
        resolved_campus_id = payload.campus_id
    else:
        raise HTTPException(status_code=400, detail="Jenis item tinjauan tidak didukung.")
    await db.beneficiary_reviews.update_one(
        {"id": review_id},
        {
            "$set": {
                "status": "resolved",
                "resolved_user_id": resolved_user_id,
                "resolved_campus_id": resolved_campus_id,
                "resolved_by": user["user_id"],
                "resolved_at": datetime.now(timezone.utc).isoformat(),
            }
        },
    )
    await audit_beneficiary_action(
        user,
        "resolve_pm_review",
        resolved_user_id,
        {
            "review_id": review_id,
            "source_id": source["id"],
            "campus_id": resolved_campus_id,
        },
    )
    return {"message": "Item tinjauan berhasil dipetakan ke Penerima Manfaat."}


# ---------------------------------------------------------------------------
# Site content (public + super admin)
# ---------------------------------------------------------------------------
def clean_campus_value(value: Any) -> str:
    return str(value or "").strip()


def serialize_campus(campus: dict, include_bank_data: bool = False) -> dict:
    data = dict(campus)
    data.pop("_id", None)
    account_number = decrypt_account_number(data.pop("bank_account_cipher", ""))
    if include_bank_data:
        data["bank_account_number_masked"] = mask_account_number(account_number)
        data["bank_account_holder_name"] = data.get("bank_account_holder_name", "")
        data["has_bank_account"] = bool(account_number)
    else:
        data.pop("bank_account_holder_name", None)
    return data


def campus_bank_input(payload: CampusInput, existing: Optional[dict] = None) -> dict:
    existing = existing or {}
    account_number = normalize_account_number(payload.bank_account_number)
    holder_name = clean_campus_value(payload.bank_account_holder_name)
    current_holder = clean_campus_value(existing.get("bank_account_holder_name"))
    update = {}
    if account_number:
        effective_holder = holder_name or current_holder
        if not effective_holder:
            raise HTTPException(
                status_code=400,
                detail="Nama pemilik rekening wajib diisi bersama nomor rekening kampus.",
            )
        if not 6 <= len(account_number) <= 24:
            raise HTTPException(status_code=400, detail="Nomor rekening kampus harus 6–24 digit.")
        update["bank_account_cipher"] = encrypt_account_number(account_number)
    if holder_name:
        update["bank_account_holder_name"] = holder_name
    return update


async def create_student_campus_code() -> str:
    counter = await db.counters.find_one_and_update(
        {"key": "student-campus-code"},
        {"$inc": {"sequence": 1}},
        upsert=True,
        return_document=ReturnDocument.AFTER,
    )
    return f"MDJ-KMP-{int(counter['sequence']):06d}"


async def register_student_campus(value: Any) -> Optional[dict]:
    name = clean_campus_value(value)
    if len(name) < 2:
        return None
    existing = await db.campuses.find_one(
        {"name": {"$regex": f"^{re.escape(name)}$", "$options": "i"}},
        {"_id": 0},
    )
    if existing:
        return None
    campus = {
        "id": str(uuid.uuid4()),
        "name": name,
        "code": await create_student_campus_code(),
        "source": "student_submission",
        "created_at": datetime.now(timezone.utc).isoformat(),
        "updated_at": datetime.now(timezone.utc).isoformat(),
    }
    await db.campuses.insert_one(dict(campus))
    return campus


@api_router.get("/campuses")
async def list_campuses(user: dict = Depends(get_current_user)):
    campuses = await db.campuses.find({}, {"_id": 0}).sort("name", 1).to_list(5000)
    include_bank_data = user.get("role") in {"admin", "super_admin"}
    return [serialize_campus(campus, include_bank_data) for campus in campuses]


@api_router.post("/campuses")
async def create_campus(
    payload: CampusInput,
    user: dict = Depends(require_roles("admin", "super_admin")),
):
    name = clean_campus_value(payload.name)
    code = clean_campus_value(payload.code)
    if len(name) < 2:
        raise HTTPException(status_code=400, detail="Nama kampus minimal 2 karakter.")
    if code and await db.campuses.find_one({"code": code}, {"_id": 0, "id": 1}):
        raise HTTPException(status_code=400, detail="Kode kampus sudah terdaftar.")
    if await db.campuses.find_one({"name": name}, {"_id": 0, "id": 1}):
        raise HTTPException(status_code=400, detail="Nama kampus sudah terdaftar.")
    campus = {
        "id": str(uuid.uuid4()),
        "name": name,
        "created_at": datetime.now(timezone.utc).isoformat(),
        "updated_at": datetime.now(timezone.utc).isoformat(),
    }
    if code:
        campus["code"] = code
    campus.update(campus_bank_input(payload))
    await db.campuses.insert_one(dict(campus))
    await audit_beneficiary_action(user, "create_campus_bank_master", detail={"campus_id": campus["id"]})
    return serialize_campus(campus, True)


@api_router.put("/campuses/{campus_id}")
async def update_campus(
    campus_id: str,
    payload: CampusInput,
    user: dict = Depends(require_roles("admin", "super_admin")),
):
    name = clean_campus_value(payload.name)
    code = clean_campus_value(payload.code)
    if len(name) < 2:
        raise HTTPException(status_code=400, detail="Nama kampus minimal 2 karakter.")
    existing = await db.campuses.find_one({"id": campus_id}, {"_id": 0})
    if not existing:
        raise HTTPException(status_code=404, detail="Data kampus tidak ditemukan.")
    duplicate = await db.campuses.find_one(
        {"id": {"$ne": campus_id}, "$or": [{"name": name}, {"code": code}]},
        {"_id": 0, "id": 1},
    )
    if duplicate and (duplicate.get("name") == name or (code and duplicate.get("code") == code)):
        raise HTTPException(status_code=400, detail="Nama atau kode kampus sudah terdaftar.")
    update = {
        "name": name,
        "updated_at": datetime.now(timezone.utc).isoformat(),
    }
    update.update(campus_bank_input(payload, existing))
    operations = {"$set": update}
    if code:
        update["code"] = code
    else:
        operations["$unset"] = {"code": ""}
    await db.campuses.update_one({"id": campus_id}, operations)
    updated = await db.campuses.find_one({"id": campus_id}, {"_id": 0})
    await audit_beneficiary_action(user, "update_campus_bank_master", detail={"campus_id": campus_id})
    return serialize_campus(updated or {"id": campus_id, **update}, True)


@api_router.delete("/campuses/{campus_id}")
async def delete_campus(
    campus_id: str,
    user: dict = Depends(require_roles("admin", "super_admin")),
):
    result = await db.campuses.delete_one({"id": campus_id})
    if not result.deleted_count:
        raise HTTPException(status_code=404, detail="Data kampus tidak ditemukan.")
    return {"message": "Data kampus dihapus."}


@api_router.post("/campuses/import")
async def import_campuses(
    file: UploadFile = File(...),
    user: dict = Depends(require_roles("admin", "super_admin")),
):
    if not file.filename.lower().endswith(".xlsx"):
        raise HTTPException(status_code=400, detail="File import harus berformat XLSX.")
    data = await file.read()
    if len(data) > 5 * 1024 * 1024:
        raise HTTPException(status_code=400, detail="Ukuran file XLSX maksimal 5MB.")
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        worksheet = workbook.active
        rows = worksheet.iter_rows(values_only=True)
        name_index = None
        code_index = None
        account_number_index = None
        account_holder_index = None
        for header_row in rows:
            headers = [clean_campus_value(value).lower() for value in header_row]
            if "nama kampus" in headers and "kode kampus" in headers:
                name_index = headers.index("nama kampus")
                code_index = headers.index("kode kampus")
                for label in ("nomor rekening", "no rekening", "rekening kampus"):
                    if label in headers:
                        account_number_index = headers.index(label)
                        break
                for label in ("atas nama rekening", "nama pemilik rekening", "pemilik rekening"):
                    if label in headers:
                        account_holder_index = headers.index(label)
                        break
                break
        if name_index is None or code_index is None:
            raise ValueError("Kolom kampus tidak ditemukan")
    except (StopIteration, ValueError, OSError) as error:
        raise HTTPException(
            status_code=400,
            detail="XLSX harus memiliki kolom 'Nama Kampus' dan 'Kode Kampus'.",
        ) from error

    inserted = 0
    updated = 0
    now = datetime.now(timezone.utc).isoformat()
    for row in rows:
        name = clean_campus_value(row[name_index] if len(row) > name_index else "")
        code = clean_campus_value(row[code_index] if len(row) > code_index else "")
        if not name:
            continue
        query = {"code": code} if code else {"name": name}
        existing = await db.campuses.find_one(query, {"_id": 0})
        campus_id = existing["id"] if existing else str(uuid.uuid4())
        account_number = clean_campus_value(
            row[account_number_index] if account_number_index is not None and len(row) > account_number_index else ""
        )
        account_holder_name = clean_campus_value(
            row[account_holder_index] if account_holder_index is not None and len(row) > account_holder_index else ""
        )
        bank_update = campus_bank_input(
            CampusInput(
                name=name,
                code=code or None,
                bank_account_number=account_number or None,
                bank_account_holder_name=account_holder_name or None,
            ),
            existing,
        )
        operations = {
            "$set": {"name": name, "updated_at": now, **bank_update},
            "$setOnInsert": {"id": campus_id, "created_at": now},
        }
        if code:
            operations["$set"]["code"] = code
        else:
            operations["$unset"] = {"code": ""}
        await db.campuses.update_one(
            query,
            operations,
            upsert=True,
        )
        if existing:
            updated += 1
        else:
            inserted += 1
    workbook.close()
    return {"inserted": inserted, "updated": updated, "total": inserted + updated}


DEFAULT_REGISTRATION_FLOW = [
    {"title": "Pembuatan Akun", "desc": ""},
    {"title": "Pendaftaran Online", "desc": ""},
    {"title": "Seleksi Administrasi", "desc": ""},
    {"title": "Pengumuman Kelulusan Seleksi Administrasi", "desc": ""},
    {"title": "Wawancara Assessment", "desc": ""},
    {"title": "Verifikasi Faktual", "desc": ""},
    {"title": "Pengumuman Kelulusan Penerima Manfaat", "desc": ""},
    {"title": "Pengukuhan / Inagurasi", "desc": ""},
    {"title": "Pencairan Bantuan Tahap I", "desc": ""},
    {"title": "Pembinaan", "desc": ""},
    {"title": "Pencairan Bantuan Tahap II", "desc": ""},
]


def default_registration_flow() -> List[dict]:
    return [dict(item) for item in DEFAULT_REGISTRATION_FLOW]


def normalize_registration_flow(value: Any) -> List[dict]:
    if not isinstance(value, list) or len(value) != len(DEFAULT_REGISTRATION_FLOW):
        raise HTTPException(
            status_code=400,
            detail="Alur pendaftaran harus berisi tepat 11 tahapan.",
        )
    return [
        {
            "title": DEFAULT_REGISTRATION_FLOW[index]["title"],
            "desc": str(item.get("desc") or "").strip()[:800]
            if isinstance(item, dict)
            else "",
        }
        for index, item in enumerate(value)
    ]


DEFAULT_SITE_CONTENT = {
    "settings": {
        "registration_open": True,
        "period_start": "8 September 2026",
        "period_end": "30 September 2026",
        "announcement_date": "12 Januari 2027",
        "year": "2026",
    },
    "banner_url": None,
    "hero": {
        "eyebrow": "Program Beasiswa Pendidikan DKI Jakarta",
        "title": "Wujudkan Pendidikan dan Masa Depan Terbaikmu",
        "subtitle": "Program Masa Depan Jakarta (MDJ) hadir sebagai wujud nyata dukungan BAZNAS (BAZIS) Provinsi DKI Jakarta. Muda dengan Zakat, Bahagia dengan Manfaat.",
    },
    "stats": [
        {"label": "Total Pendaftar (22-25)", "value": "16.935"},
        {"label": "Penerima Manfaat 2025", "value": "3.049"},
        {"label": "Wilayah Terjangkau", "value": "6"},
        {"label": "Mitra Kampus (UPZ)", "value": "37+"},
    ],
    "timeline": [
        {"title": "Pembuatan Akun", "desc": "Calon pendaftar membuat akun resmi pada portal pendaftaran."},
        {"title": "Pendaftaran Online", "desc": "Pengisian formulir dan pengunggahan berkas persyaratan."},
        {"title": "Seleksi Administrasi", "desc": "Verifikasi dan validasi kesesuaian dokumen pendaftar."},
        {"title": "Pengumuman Seleksi Administrasi", "desc": "Pengumuman peserta yang memenuhi syarat dokumen."},
        {"title": "Wawancara Assessment", "desc": "Sesi penggalian potensi, motivasi, dan komitmen peserta."},
        {"title": "Verifikasi Faktual", "desc": "Survei dan validasi lapangan ke kediaman serta kampus."},
        {"title": "Pengumuman Penerima Manfaat", "desc": "Penetapan akhir peserta terpilih sebagai penerima beasiswa."},
        {"title": "Pengukuhan / Inagurasi", "desc": "Acara peresmian dan penyambutan resmi penerima manfaat baru."},
    ],
    "registration_flow": default_registration_flow(),
    "announcements": [
        {"date": "12 Jan 2026", "category": "Informasi", "title": "Pengukuhan Kader Akademia MDJ", "summary": "Peserta yang lolos seleksi akhir wajib mengikuti kegiatan pengukuhan dan orientasi program pembinaan."},
        {"date": "22 Des 2025", "category": "Pengumuman", "title": "Pengumuman Awardee MDJ Scholarship", "summary": "Daftar nama 3.049 penerima manfaat yang lolos tahap wawancara dan verifikasi faktual telah diterbitkan."},
        {"date": "18 Okt 2025", "category": "Jadwal", "title": "Pelaksanaan Verifikasi Faktual", "summary": "Tim BAZNAS (BAZIS) DKI Jakarta akan melakukan survey ke domisili dan audiensi ke kampus mahasiswa."},
    ],
    "requirements": [
        "Surat Permohonan yang bersangkutan.",
        "KTP Provinsi DKI Jakarta.",
        "Pas Foto berwarna ukuran 3x4.",
        "Kartu Keluarga (KK).",
        "Kartu Tanda Mahasiswa (KTM).",
        "Surat Keterangan Mahasiswa Aktif (Masih Kuliah).",
        "SKTM atau Surat Rekomendasi Masjid/Lembaga Keagamaan Resmi.",
        "Surat Persetujuan Orang Tua/Wali.",
        "Surat Keterangan tidak sedang menerima beasiswa lain.",
        "Pakta Integritas Kader Akademia MDJ (Sesuai format BAZNAS).",
    ],
    "eligibility_requirements": [
        "Warga DKI Jakarta (dibuktikan dengan KTP & KK DKI Jakarta).",
        "Berstatus sebagai Mahasiswa Aktif di perguruan tinggi.",
        "Berasal dari keluarga kurang mampu atau mendapat rekomendasi dari lembaga keagamaan/ormas Islam resmi.",
        "Tidak sedang menerima beasiswa dari pihak/lembaga lain.",
        "Mendapatkan izin dan persetujuan dari Orang Tua/Wali.",
        "Bersedia menyetujui dan mematuhi aturan/integritas yang ditetapkan oleh BAZNAS DKI Jakarta.",
    ],
    "required_documents": [
        "Surat Permohonan.",
        "KTP DKI Jakarta.",
        "Kartu Keluarga (KK).",
        "Kartu Tanda Mahasiswa (KTM).",
        "Pas Foto Ukuran 3x4.",
        "Surat Keterangan Tidak Mampu (SKTM) atau Surat Rekomendasi dari Masjid, Majelis Taklim, Lembaga Keagamaan, atau Ormas Islam Resmi.",
        "Surat Keterangan Mahasiswa Aktif.",
        "Transkrip Nilai Terakhir.",
        "Surat Persetujuan Orang Tua/Wali.",
        "Surat Keterangan Tidak Menerima Beasiswa Lain.",
        "Pakta Integritas (format disesuaikan dengan ketentuan BAZNAS DKI Jakarta).",
    ],
    "faqs": [
        {"q": "Apa itu Program Masa Depan Jakarta?", "a": "Program Masa Depan Jakarta Scholarship merupakan Bantuan Biaya Pendidikan Mahasiswa DKI Jakarta yang diselenggarakan oleh BAZNAS (BAZIS) Provinsi DKI Jakarta."},
        {"q": "Siapa saja yang dapat mendaftar?", "a": "Mahasiswa (Diploma/Sarjana) aktif yang berdomisili di DKI Jakarta dan memenuhi kriteria keluarga prasejahtera."},
        {"q": "Apakah pendaftaran dikenakan biaya?", "a": "Tidak. Seluruh proses pendaftaran, seleksi, hingga penyaluran dana TIDAK DIPUNGUT BIAYA (Gratis)."},
        {"q": "Bagaimana cara mengetahui status pendaftaran?", "a": "Login ke akun Anda di portal ini. Status akan diperbarui secara berkala pada menu dashboard pendaftar."},
    ],
}


@api_router.get("/site/content")
async def get_site_content():
    content = await db.site_content.find_one({"key": "main"}, {"_id": 0})
    if not content:
        content = {"key": "main", **DEFAULT_SITE_CONTENT}
        await db.site_content.insert_one(dict(content))
        content.pop("_id", None)
    if "registration_flow" not in content:
        content["registration_flow"] = default_registration_flow()
        await db.site_content.update_one(
            {"key": "main"},
            {"$set": {"registration_flow": content["registration_flow"]}},
            upsert=True,
        )
    missing_lists = {
        key: value
        for key, value in {
            "eligibility_requirements": DEFAULT_SITE_CONTENT["eligibility_requirements"],
            "required_documents": DEFAULT_SITE_CONTENT["required_documents"],
        }.items()
        if key not in content
    }
    if missing_lists:
        content.update(missing_lists)
        await db.site_content.update_one({"key": "main"}, {"$set": missing_lists}, upsert=True)
    content.pop("key", None)
    return content


@api_router.put("/site/content")
async def update_site_content(payload: SiteContentInput, user: dict = Depends(require_roles("super_admin"))):
    current = await db.site_content.find_one({"key": "main"}, {"_id": 0}) or {}
    content_update = dict(payload.content)
    if "registration_flow" in content_update:
        content_update["registration_flow"] = normalize_registration_flow(
            content_update["registration_flow"]
        )
    incoming_announcements = content_update.get("announcements")
    new_announcements = []
    if isinstance(incoming_announcements, list):
        existing_signatures = {
            announcement_signature(announcement)
            for announcement in current.get("announcements", [])
        }
        new_announcements = [
            announcement
            for announcement in incoming_announcements
            if announcement_signature(announcement) not in existing_signatures
        ]
    await db.site_content.update_one(
        {"key": "main"},
        {"$set": {**content_update, "updated_at": datetime.now(timezone.utc).isoformat()}},
        upsert=True,
    )
    await notify_new_announcements(new_announcements)
    content = await db.site_content.find_one({"key": "main"}, {"_id": 0})
    content.pop("key", None)
    return content


@api_router.post("/site/banner")
async def upload_banner(file: UploadFile = File(...), user: dict = Depends(require_roles("super_admin"))):
    ext = file.filename.split(".")[-1].lower() if "." in file.filename else "png"
    content_type = file.content_type or MIME_TYPES.get(ext, "image/png")
    data, ext, content_type = compress_image(
        await file.read(),
        ext,
        content_type,
        max_w=1920,
        quality=80,
    )
    path = f"{APP_NAME}/site/banner_{uuid.uuid4().hex[:8]}.{ext}"
    result = put_object(path, data, content_type)
    await db.site_files.insert_one({
        "storage_path": result["path"], "content_type": content_type,
        "is_deleted": False, "created_at": datetime.now(timezone.utc).isoformat(),
    })
    banner_url = f"/api/files/{result['path']}"
    await db.site_content.update_one({"key": "main"}, {"$set": {"banner_url": banner_url}}, upsert=True)
    return {"banner_url": banner_url, "storage_path": result["path"]}


@api_router.post("/site/about-image")
async def upload_about_image(
    file: UploadFile = File(...),
    user: dict = Depends(require_roles("super_admin")),
):
    ext = file.filename.split(".")[-1].lower() if "." in file.filename else "png"
    content_type = file.content_type or MIME_TYPES.get(ext, "image/png")
    data, ext, content_type = compress_image(
        await file.read(),
        ext,
        content_type,
        max_w=1600,
        quality=84,
    )
    path = f"{APP_NAME}/site/about_{uuid.uuid4().hex[:8]}.{ext}"
    result = put_object(path, data, content_type)
    await db.site_files.insert_one({
        "storage_path": result["path"],
        "content_type": content_type,
        "is_deleted": False,
        "created_at": datetime.now(timezone.utc).isoformat(),
    })
    about_url = f"/api/files/{result['path']}"
    await db.site_content.update_one(
        {"key": "main"},
        {"$set": {"about_url": about_url}},
        upsert=True,
    )
    return {"about_url": about_url, "storage_path": result["path"]}


@api_router.post("/site/logo")
async def upload_site_logo(
    file: UploadFile = File(...),
    user: dict = Depends(require_roles("super_admin")),
):
    ext = file.filename.split(".")[-1].lower() if "." in file.filename else "bin"
    if ext not in ("jpg", "jpeg", "png", "webp"):
        raise HTTPException(status_code=400, detail="Logo harus berupa JPG, PNG, atau WEBP.")
    content_type = file.content_type or MIME_TYPES.get(ext, "image/png")
    data, ext, content_type = compress_image(
        await file.read(),
        ext,
        content_type,
        max_w=512,
        quality=92,
    )
    path = f"{APP_NAME}/site/logo_{uuid.uuid4().hex[:8]}.{ext}"
    result = put_object(path, data, content_type)
    await db.site_files.insert_one({
        "storage_path": result["path"],
        "content_type": content_type,
        "is_deleted": False,
        "created_at": datetime.now(timezone.utc).isoformat(),
    })
    logo_url = f"/api/files/{result['path']}"
    await db.site_content.update_one(
        {"key": "main"},
        {"$set": {"logo_url": logo_url}},
        upsert=True,
    )
    return {"logo_url": logo_url, "storage_path": result["path"]}


# ---------------------------------------------------------------------------
# Admin: participants management
# ---------------------------------------------------------------------------
IMPORT_STAGE_CONFIG = {
    "wawancara_lolos": {
        "label": "Hasil Lolos Wawancara",
        "target_status": "verifikasi_faktual",
    },
    "kelulusan_akhir": {
        "label": "Hasil Kelulusan Akhir",
        "target_status": "lolos",
    },
}

REQUIRED_VERIFICATION_DOCUMENTS = [
    "KTP DKI Jakarta",
    "Kartu Keluarga (KK)",
    "Pas Foto 3x4",
    "Kartu Tanda Mahasiswa (KTM)",
    "KRS / KHS / Transkrip Nilai",
    "Surat Keterangan Mahasiswa Aktif",
    "SKTM / Surat Rekomendasi",
    "Surat Persetujuan Orang Tua",
    "Surat Keterangan Tidak Menerima Beasiswa Lain",
]

REQUIRED_VERIFICATION_FIELDS = {
    "namaLengkap": "Nama lengkap",
    "email": "Email",
    "nik": "NIK",
    "noTelp": "Nomor telepon",
    "alamatLengkap": "Alamat lengkap",
    "kota": "Wilayah/kota domisili",
    "provinsi": "Provinsi",
    "institusi": "Perguruan tinggi",
    "nim": "NIM",
    "jenjang": "Jenjang pendidikan",
    "jurusan": "Program studi",
    "semester": "Semester",
    "ipk": "IPK",
    "biayaPendidikanSemester": "Biaya pendidikan per semester",
}

IMPORT_FIELD_ALIASES = {
    "cpm_id": {"idcpm", "cpmid", "nomorcpm", "idcalonpenerimamanfaat"},
    "nik": {"nik", "nomorindukkependudukan"},
    "email": {"email", "emailmahasiswa", "alamatemail"},
    "name": {"nama", "namamahasiswa", "namalengkap", "peserta"},
    "campus": {"kampus", "namakampus", "perguruantinggi", "universitas"},
    "region": {"wilayah", "kota", "kotakabupaten", "domisili", "kabupaten"},
}

STATUS_RANK = {
    "draft": 0,
    "submitted": 1,
    "verifikasi": 2,
    "lolos_administrasi": 3,
    "wawancara": 4,
    "verifikasi_faktual": 5,
    "lolos": 6,
    "ditolak": -1,
}


def normalized_import_value(value: Any) -> str:
    return re.sub(r"[^a-z0-9]", "", str(value or "").lower())


def normalized_name(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip().lower())


def normalized_cpm(value: Any) -> str:
    match = re.search(r"CPM-MDJ/\d{4}/\d{6}", str(value or "").upper())
    return match.group(0) if match else ""


def find_import_header_row(rows: List[tuple]) -> int:
    best_index = 0
    best_score = -1
    for index, row in enumerate(rows[:10]):
        score = sum(
            normalized_import_value(value) in aliases
            for value in row
            for aliases in IMPORT_FIELD_ALIASES.values()
        )
        if score > best_score:
            best_index = index
            best_score = score
    return best_index


def deterministic_import_mapping(headers: List[str]) -> Dict[str, Optional[str]]:
    mapping = {}
    for field, aliases in IMPORT_FIELD_ALIASES.items():
        mapping[field] = next(
            (header for header in headers if normalized_import_value(header) in aliases),
            None,
        )
    return mapping


async def ai_import_mapping(headers: List[str], sample_rows: List[dict]) -> Dict[str, Optional[str]]:
    if not EMERGENT_KEY:
        return {}
    prompt = (
        "Anda membantu import hasil seleksi beasiswa Indonesia. Petakan nama kolom sumber ke field "
        "berikut: cpm_id, nik, email, name, campus, region. Gunakan hanya nama kolom yang persis "
        "ada pada daftar headers atau null jika tidak ada. Kembalikan HANYA JSON object.\n"
        f"Headers: {json.dumps(headers, ensure_ascii=False)}\n"
        f"Contoh baris: {json.dumps(sample_rows[:5], ensure_ascii=False)}"
    )
    try:
        chat = LlmChat(
            api_key=EMERGENT_KEY,
            session_id=f"selection-import-map-{uuid.uuid4()}",
            system_message="Anda adalah asisten pemetaan kolom data yang teliti.",
        ).with_model("gemini", "gemini-3.1-pro-preview")
        chunks = []
        async for event in chat.stream_message(UserMessage(text=prompt)):
            if isinstance(event, TextDelta):
                chunks.append(event.content)
            elif isinstance(event, StreamDone):
                break
        match = re.search(r"\{.*\}", "".join(chunks), flags=re.DOTALL)
        parsed = json.loads(match.group(0)) if match else {}
        return {
            field: value
            for field, value in parsed.items()
            if field in IMPORT_FIELD_ALIASES and value in headers
        }
    except Exception as error:
        logger.warning(f"AI import mapping skipped: {error}")
        return {}


async def ai_review_notes(items: List[dict]) -> Dict[int, str]:
    if not EMERGENT_KEY or not items:
        return {}
    prompt_items = [
        {
            "row_number": item["row_number"],
            "kind": item["kind"],
            "reason": item["reason"],
            "data": item["imported_data"],
        }
        for item in items[:25]
    ]
    prompt = (
        "Berikan catatan singkat Bahasa Indonesia untuk tiap data import beasiswa yang perlu ditinjau. "
        "Jangan mengubah atau menyimpulkan kecocokan peserta. Kembalikan HANYA JSON object dengan "
        "format {\"notes\":[{\"row_number\":2,\"note\":\"...\"}]}. Data: "
        f"{json.dumps(prompt_items, ensure_ascii=False)}"
    )
    try:
        chat = LlmChat(
            api_key=EMERGENT_KEY,
            session_id=f"selection-import-review-{uuid.uuid4()}",
            system_message="Anda adalah asisten pemeriksaan data seleksi beasiswa.",
        ).with_model("gemini", "gemini-3.1-pro-preview")
        chunks = []
        async for event in chat.stream_message(UserMessage(text=prompt)):
            if isinstance(event, TextDelta):
                chunks.append(event.content)
            elif isinstance(event, StreamDone):
                break
        match = re.search(r"\{.*\}", "".join(chunks), flags=re.DOTALL)
        notes = json.loads(match.group(0)).get("notes", []) if match else []
        return {int(note["row_number"]): str(note["note"]) for note in notes if note.get("note")}
    except Exception as error:
        logger.warning(f"AI import review skipped: {error}")
        return {}


async def build_participant_match_indexes(
    user: Optional[dict] = None,
) -> Dict[str, Dict[str, List[dict]]]:
    query = await participant_scope_query(user) if user else {}
    registrations = await db.registrations.find(query, {"_id": 0}).to_list(5000)
    user_ids = [registration["user_id"] for registration in registrations]
    profiles = await db.profiles.find(
        {"user_id": {"$in": user_ids}},
        {"_id": 0, "user_id": 1, "data": 1},
    ).to_list(5000)
    profile_data = {profile["user_id"]: profile.get("data", {}) for profile in profiles}
    indexes = {"cpm_id": {}, "nik": {}, "email": {}, "name_campus_region": {}}
    for registration in registrations:
        data = profile_data.get(registration["user_id"], {})
        candidate = {"registration": registration, "profile": data}
        keys = {
            "cpm_id": normalized_cpm(registration.get("cpm_id")),
            "nik": re.sub(r"\D", "", str(data.get("nik", ""))),
            "email": normalized_name(registration.get("email") or data.get("email")),
            "name_campus_region": "|".join([
                normalized_name(registration.get("name") or data.get("namaLengkap")),
                normalized_name(data.get("institusi")),
                normalized_name(data.get("kota") or data.get("provinsi")),
            ]),
        }
        for index_name, key in keys.items():
            if key:
                indexes[index_name].setdefault(key, []).append(candidate)
    return indexes


def match_imported_participant(data: dict, indexes: Dict[str, Dict[str, List[dict]]]) -> tuple:
    checks = [
        ("cpm_id", normalized_cpm(data.get("cpm_id"))),
        ("nik", re.sub(r"\D", "", str(data.get("nik", "")))),
        ("email", normalized_name(data.get("email"))),
        ("name_campus_region", "|".join([
            normalized_name(data.get("name")),
            normalized_name(data.get("campus")),
            normalized_name(data.get("region")),
        ])),
    ]
    for index_name, key in checks:
        if not key or (index_name == "name_campus_region" and key.count("|") != 2):
            continue
        candidates = indexes[index_name].get(key, [])
        if len(candidates) == 1:
            return candidates[0], index_name, None
        if len(candidates) > 1:
            return None, index_name, "Data identitas cocok ke lebih dari satu peserta."
    return None, None, None


def imported_differences(imported: dict, candidate: dict) -> List[str]:
    registration = candidate["registration"]
    profile = candidate["profile"]
    comparisons = [
        ("Nama", imported.get("name"), registration.get("name") or profile.get("namaLengkap")),
        ("Kampus", imported.get("campus"), profile.get("institusi")),
        ("Wilayah", imported.get("region"), profile.get("kota") or profile.get("provinsi")),
    ]
    return [
        label
        for label, source, saved in comparisons
        if source and saved and normalized_name(source) != normalized_name(saved)
    ]


def verification_issues(profile: dict, document_types: set, campus_names: set) -> List[str]:
    issues = [
        label
        for key, label in REQUIRED_VERIFICATION_FIELDS.items()
        if not str(profile.get(key, "")).strip()
    ]
    nik = re.sub(r"\D", "", str(profile.get("nik", "")))
    if nik and len(nik) != 16:
        issues.append("Format NIK tidak 16 digit")
    try:
        ipk = float(str(profile.get("ipk", "")).replace(",", "."))
        if not 0 <= ipk <= 4:
            issues.append("IPK di luar rentang 0 sampai 4")
    except (TypeError, ValueError):
        if profile.get("ipk"):
            issues.append("Format IPK tidak valid")
    biaya = re.sub(r"\D", "", str(profile.get("biayaPendidikanSemester", "")))
    if profile.get("biayaPendidikanSemester") and not biaya:
        issues.append("Format biaya pendidikan tidak valid")
    institution = normalized_name(profile.get("institusi"))
    if institution and institution not in campus_names:
        issues.append("Kampus belum terdaftar pada data kampus")
    missing_documents = [
        document
        for document in REQUIRED_VERIFICATION_DOCUMENTS
        if document not in document_types
    ]
    if missing_documents:
        issues.append(f"Berkas belum lengkap: {', '.join(missing_documents)}")
    return issues


async def ai_administration_recommendations(candidates: List[dict]) -> Dict[str, dict]:
    if not EMERGENT_KEY or not candidates:
        return {}
    results = {}
    for start in range(0, len(candidates), 20):
        batch = candidates[start:start + 20]
        prompt = (
            "Anda membantu verifikasi administrasi beasiswa. Data berikut sudah lulus pemeriksaan "
            "kelengkapan aturan sistem. Nilai kewajaran data administratif secara konservatif tanpa "
            "mengubah data. Kembalikan HANYA JSON {\"results\":[{\"user_id\":\"...\","
            "\"recommend\":\"recommended|review\",\"summary\":\"...\",\"issues\":[\"...\"]}]}. "
            "Gunakan recommended jika tidak ada konflik eksplisit di data yang diberikan. Jangan menebak "
            "bahwa data fiktif dari pola nama, domain email, nomor, atau alamat. Gunakan review hanya jika "
            "ada kontradiksi eksplisit. Catatan: Anda tidak memeriksa keaslian visual dokumen. Data: "
            f"{json.dumps(batch, ensure_ascii=False)}"
        )
        try:
            chat = LlmChat(
                api_key=EMERGENT_KEY,
                session_id=f"admin-verification-{uuid.uuid4()}",
                system_message="Anda adalah asisten verifikasi administrasi beasiswa yang teliti.",
            ).with_model("gemini", "gemini-3.1-pro-preview")
            chunks = []
            async for event in chat.stream_message(UserMessage(text=prompt)):
                if isinstance(event, TextDelta):
                    chunks.append(event.content)
                elif isinstance(event, StreamDone):
                    break
            match = re.search(r"\{.*\}", "".join(chunks), flags=re.DOTALL)
            parsed = json.loads(match.group(0)).get("results", []) if match else []
            for item in parsed:
                user_id = item.get("user_id")
                if user_id:
                    results[user_id] = item
        except Exception as error:
            logger.warning(f"AI administrative verification skipped: {error}")
    return results


@api_router.post("/admin/verification-recommendations/generate")
async def generate_verification_recommendations(
    user: dict = Depends(require_roles(*MANAGEMENT_ROLES)),
):
    registration_query = await participant_scope_query(user)
    registration_query["status"] = {"$in": ["submitted", "verifikasi"]}
    registrations = await db.registrations.find(
        registration_query,
        {"_id": 0},
    ).to_list(5000)
    user_ids = [registration["user_id"] for registration in registrations]
    profiles = await db.profiles.find(
        {"user_id": {"$in": user_ids}},
        {"_id": 0, "user_id": 1, "data": 1},
    ).to_list(5000)
    documents = await db.documents.find(
        {"user_id": {"$in": user_ids}, "is_deleted": {"$ne": True}},
        {"_id": 0, "user_id": 1, "doc_type": 1},
    ).to_list(100000)
    campuses = await db.campuses.find({}, {"_id": 0, "name": 1}).to_list(5000)
    profiles_by_user = {profile["user_id"]: profile.get("data", {}) for profile in profiles}
    documents_by_user = {}
    for document in documents:
        documents_by_user.setdefault(document["user_id"], set()).add(document["doc_type"])
    campus_names = {normalized_name(campus.get("name")) for campus in campuses}

    candidate_checks = []
    ai_candidates = []
    for registration in registrations:
        profile = profiles_by_user.get(registration["user_id"], {})
        document_types = documents_by_user.get(registration["user_id"], set())
        issues = verification_issues(profile, document_types, campus_names)
        candidate = {
            "user_id": registration["user_id"],
            "name": registration.get("name") or profile.get("namaLengkap", "-"),
            "cpm_id": registration.get("cpm_id", "-"),
            "profile": profile,
            "documents": sorted(document_types),
            "issues": issues,
        }
        candidate_checks.append(candidate)
        if not issues:
            ai_candidates.append(candidate)

    ai_results = await ai_administration_recommendations(ai_candidates)
    now = datetime.now(timezone.utc).isoformat()
    recommendations = []
    for candidate in candidate_checks:
        ai_result = ai_results.get(candidate["user_id"], {})
        issues = candidate["issues"] + ai_result.get("issues", [])
        is_recommended = not candidate["issues"] and ai_result.get("recommend") == "recommended"
        recommendation = {
            "id": str(uuid.uuid4()),
            "user_id": candidate["user_id"],
            "name": candidate["name"],
            "cpm_id": candidate["cpm_id"],
            "recommendation": "recommended" if is_recommended else "review",
            "approval_status": "pending",
            "score": 100 if is_recommended else max(0, 100 - len(issues) * 8),
            "issues": issues,
            "ai_summary": ai_result.get("summary") or (
                "Pemeriksaan aturan menemukan data yang perlu dilengkapi."
                if candidate["issues"]
                else "AI belum dapat memberikan rekomendasi otomatis."
            ),
            "document_count": len(candidate["documents"]),
            "region": profile_region(candidate["profile"]),
            "analyzed_at": now,
            "analyzed_by": user["user_id"],
        }
        existing = await db.verification_recommendations.find_one(
            {"user_id": candidate["user_id"]},
            {"_id": 0, "id": 1},
        )
        if existing:
            recommendation["id"] = existing["id"]
        await db.verification_recommendations.update_one(
            {"user_id": candidate["user_id"]},
            {"$set": recommendation},
            upsert=True,
        )
        recommendations.append(recommendation)

    return {
        "total": len(recommendations),
        "recommended": sum(item["recommendation"] == "recommended" for item in recommendations),
        "review": sum(item["recommendation"] == "review" for item in recommendations),
    }


@api_router.get("/admin/verification-recommendations")
async def list_verification_recommendations(
    user: dict = Depends(require_roles(*MANAGEMENT_ROLES)),
):
    query = {"approval_status": "pending"}
    scope_query = await participant_scope_query(user)
    if scope_query:
        query.update(scope_query)
    return await db.verification_recommendations.find(
        query,
        {"_id": 0},
    ).sort("analyzed_at", -1).to_list(5000)


@api_router.post("/admin/verification-recommendations/approve")
async def approve_verification_recommendations(
    payload: VerificationApprovalInput,
    user: dict = Depends(require_roles(*MANAGEMENT_ROLES)),
):
    if not payload.approve_all and not payload.recommendation_ids:
        raise HTTPException(status_code=400, detail="Pilih minimal satu rekomendasi untuk disetujui.")
    query = {"recommendation": "recommended", "approval_status": "pending"}
    scope_query = await participant_scope_query(user)
    if scope_query:
        query.update(scope_query)
    if not payload.approve_all:
        query["id"] = {"$in": payload.recommendation_ids}
    recommendations = await db.verification_recommendations.find(query, {"_id": 0}).to_list(5000)
    now = datetime.now(timezone.utc).isoformat()
    approved = 0
    for recommendation in recommendations:
        result = await db.registrations.update_one(
            {
                "user_id": recommendation["user_id"],
                "status": {"$in": ["submitted", "verifikasi"]},
            },
            {
                "$set": {"status": "lolos_administrasi", "updated_at": now},
                "$push": {
                    "history": {
                        "status": "lolos_administrasi",
                        "note": "Lolos Administrasi disetujui Admin dari rekomendasi AI.",
                        "at": now,
                    },
                },
            },
        )
        if not result.modified_count:
            continue
        await db.verification_recommendations.update_one(
            {"id": recommendation["id"]},
            {"$set": {"approval_status": "approved", "approved_at": now, "approved_by": user["user_id"]}},
        )
        await create_notification(
            recommendation["user_id"],
            "selection_progress",
            "Perkembangan seleksi diperbarui",
            "Anda dinyatakan Lolos Administrasi setelah verifikasi oleh Admin.",
        )
        approved += 1
    return {"message": "Rekomendasi berhasil disetujui.", "approved": approved}


@api_router.post("/admin/selection-imports")
async def import_selection_results(
    stage: str = Form(...),
    file: UploadFile = File(...),
    user: dict = Depends(require_roles(*MANAGEMENT_ROLES)),
):
    if stage not in IMPORT_STAGE_CONFIG:
        raise HTTPException(status_code=400, detail="Jenis hasil import tidak valid.")
    if not file.filename.lower().endswith(".xlsx"):
        raise HTTPException(status_code=400, detail="File hasil seleksi harus berformat XLSX.")
    file_data = await file.read()
    if len(file_data) > 10 * 1024 * 1024:
        raise HTTPException(status_code=400, detail="Ukuran file XLSX maksimal 10MB.")
    try:
        workbook = load_workbook(io.BytesIO(file_data), read_only=True, data_only=True)
        worksheet = workbook.active
        rows = list(worksheet.iter_rows(values_only=True))
        workbook.close()
    except Exception as error:
        raise HTTPException(status_code=400, detail="File XLSX tidak dapat dibaca.") from error
    if len(rows) < 2:
        raise HTTPException(status_code=400, detail="File XLSX belum memiliki data peserta.")

    header_index = find_import_header_row(rows)
    headers = [str(value or "").strip() for value in rows[header_index]]
    if not any(headers):
        raise HTTPException(status_code=400, detail="Header kolom XLSX tidak ditemukan.")
    raw_rows = rows[header_index + 1:header_index + 1001]
    sample_rows = [
        {headers[index]: str(value or "").strip() for index, value in enumerate(row) if index < len(headers)}
        for row in raw_rows[:5]
    ]
    mapping = deterministic_import_mapping(headers)
    ai_mapping = await ai_import_mapping(headers, sample_rows)
    for field, header in ai_mapping.items():
        if not mapping.get(field):
            mapping[field] = header

    import_id = str(uuid.uuid4())
    path = f"{APP_NAME}/selection-imports/{import_id}.xlsx"
    try:
        storage_result = put_object(path, file_data, MIME_TYPES["xlsx"])
    except Exception as error:
        logger.error(f"Selection import storage failed: {error}")
        raise HTTPException(status_code=502, detail="File import belum dapat disimpan untuk audit.")

    header_positions = {header: index for index, header in enumerate(headers)}
    indexes = await build_participant_match_indexes(user)
    region_scope = regional_admin_region(user)
    target_status = IMPORT_STAGE_CONFIG[stage]["target_status"]
    summary = {"matched": 0, "updated": 0, "unchanged": 0, "review": 0, "new": 0, "invalid": 0}
    review_items = []
    now = datetime.now(timezone.utc).isoformat()

    for offset, row in enumerate(raw_rows, start=header_index + 2):
        raw_data = {
            header: str(row[index] or "").strip()
            for header, index in header_positions.items()
            if index < len(row) and row[index] not in (None, "")
        }
        if not raw_data:
            continue
        imported = {
            field: raw_data.get(header, "")
            for field, header in mapping.items()
            if header
        }
        candidate, matched_by, matching_error = match_imported_participant(imported, indexes)
        review_base = {
            "id": str(uuid.uuid4()),
            "import_id": import_id,
            "stage": stage,
            "status": "pending",
            "row_number": offset,
            "imported_data": imported,
            "raw_data": raw_data,
            "region_scope": region_scope,
            "created_at": now,
        }
        if matching_error:
            review_items.append({
                **review_base,
                "kind": "mismatch",
                "reason": matching_error,
                "differences": [],
            })
            summary["review"] += 1
            continue
        if not candidate:
            identity_values = [imported.get(key) for key in ("cpm_id", "nik", "email", "name")]
            kind = "new_participant" if any(identity_values) else "invalid"
            reason = (
                "Peserta belum ditemukan di data pendaftaran."
                if kind == "new_participant"
                else "Tidak ada data identitas yang dapat dipakai untuk mencocokkan peserta."
            )
            review_items.append({
                **review_base,
                "kind": kind,
                "reason": reason,
                "differences": [],
            })
            summary["new" if kind == "new_participant" else "invalid"] += 1
            continue

        differences = imported_differences(imported, candidate)
        if differences:
            review_items.append({
                **review_base,
                "kind": "mismatch",
                "reason": "Data import tidak sesuai dengan data pendaftaran.",
                "differences": differences,
                "matched_user_id": candidate["registration"]["user_id"],
                "matched_by": matched_by,
            })
            summary["review"] += 1
            continue

        registration = candidate["registration"]
        summary["matched"] += 1
        if STATUS_RANK.get(registration.get("status"), 0) >= STATUS_RANK[target_status]:
            summary["unchanged"] += 1
            continue
        entry = {
            "status": target_status,
            "note": f"{IMPORT_STAGE_CONFIG[stage]['label']} diimpor oleh Admin.",
            "at": now,
        }
        await db.registrations.update_one(
            {"user_id": registration["user_id"]},
            {
                "$set": {"status": target_status, "updated_at": now},
                "$push": {"history": entry},
            },
        )
        await create_notification(
            registration["user_id"],
            "selection_progress",
            "Perkembangan seleksi diperbarui",
            f"{IMPORT_STAGE_CONFIG[stage]['label']} telah diperbarui dari hasil import Admin.",
        )
        summary["updated"] += 1

    ai_notes = await ai_review_notes(review_items)
    for item in review_items:
        item["ai_note"] = ai_notes.get(item["row_number"], "")
    if review_items:
        await db.selection_import_review.insert_many([dict(item) for item in review_items])

    import_record = {
        "id": import_id,
        "stage": stage,
        "stage_label": IMPORT_STAGE_CONFIG[stage]["label"],
        "original_filename": file.filename,
        "storage_path": storage_result["path"],
        "content_type": MIME_TYPES["xlsx"],
        "size": storage_result["size"],
        "mapping": mapping,
        "summary": summary,
        "review_count": len(review_items),
        "imported_by": user["user_id"],
        "region_scope": region_scope,
        "created_at": now,
    }
    await db.selection_imports.insert_one(dict(import_record))
    return {
        "import_id": import_id,
        "stage": stage,
        "mapping": mapping,
        "summary": summary,
        "review_count": len(review_items),
    }


@api_router.get("/admin/selection-imports")
async def list_selection_imports(
    user: dict = Depends(require_roles(*MANAGEMENT_ROLES)),
):
    query = {}
    region_scope = regional_admin_region(user)
    if region_scope:
        query["region_scope"] = region_scope
    return await db.selection_imports.find(query, {"_id": 0}).sort("created_at", -1).to_list(50)


@api_router.get("/admin/selection-import-review")
async def list_selection_import_review(
    kind: Optional[str] = None,
    user: dict = Depends(require_roles(*MANAGEMENT_ROLES)),
):
    query = {"status": "pending"}
    region_scope = regional_admin_region(user)
    if region_scope:
        query["region_scope"] = region_scope
    if kind == "new":
        query["kind"] = "new_participant"
    elif kind == "review":
        query["kind"] = {"$in": ["mismatch", "invalid"]}
    return await db.selection_import_review.find(query, {"_id": 0}).sort("created_at", -1).to_list(300)


@api_router.delete("/admin/selection-import-review/{item_id}")
async def dismiss_selection_import_review(
    item_id: str,
    user: dict = Depends(require_roles(*MANAGEMENT_ROLES)),
):
    query = {"id": item_id, "status": "pending"}
    region_scope = regional_admin_region(user)
    if region_scope:
        query["region_scope"] = region_scope
    result = await db.selection_import_review.update_one(
        query,
        {"$set": {"status": "dismissed", "dismissed_at": datetime.now(timezone.utc).isoformat()}},
    )
    if not result.matched_count:
        raise HTTPException(status_code=404, detail="Data tinjauan tidak ditemukan.")
    return {"message": "Data dihapus dari daftar tinjauan."}


@api_router.get("/admin/selection-imports/{import_id}/file")
async def download_selection_import_file(
    import_id: str,
    user: dict = Depends(require_roles(*MANAGEMENT_ROLES)),
):
    record = await db.selection_imports.find_one({"id": import_id}, {"_id": 0})
    region_scope = regional_admin_region(user)
    if not record or (region_scope and record.get("region_scope") != region_scope):
        raise HTTPException(status_code=404, detail="Riwayat import tidak ditemukan.")
    data, content_type = get_object(record["storage_path"])
    headers = {"Content-Disposition": f'attachment; filename="{record["original_filename"]}"'}
    return StarletteResponse(content=data, media_type=content_type, headers=headers)


@api_router.get("/admin/stats")
async def admin_stats(
    region: Optional[str] = None,
    user: dict = Depends(require_roles(*MANAGEMENT_ROLES)),
):
    scope_region = regional_admin_region(user)
    scope_query = await participant_scope_query(user)
    all_regs = await db.registrations.find(scope_query, {"_id": 0}).to_list(5000)
    user_ids = [registration["user_id"] for registration in all_regs]
    profiles = await db.profiles.find(
        {"user_id": {"$in": user_ids}},
        {"_id": 0, "user_id": 1, "data": 1},
    ).to_list(5000)
    profile_data = {
        profile["user_id"]: profile.get("data", {})
        for profile in profiles
    }

    def region_name(registration: dict) -> str:
        data = profile_data.get(registration["user_id"], {})
        return str(data.get("kota") or data.get("provinsi") or "Belum diisi").strip()

    available_regions = sorted({
        region_name(registration)
        for registration in all_regs
        if region_name(registration) != "Belum diisi"
    })
    if scope_region:
        available_regions = [scope_region]
    regs = all_regs
    selected_region = scope_region or region or "all"
    if selected_region != "all":
        regs = [
            registration
            for registration in all_regs
            if region_name(registration) == selected_region
        ]
    total = len(regs)
    by_status = {}
    for r in regs:
        by_status[r.get("status", "draft")] = by_status.get(r.get("status", "draft"), 0) + 1
    if scope_region:
        total_students = len({registration["user_id"] for registration in regs})
        total_admins = await db.users.count_documents(
            {"role": REGIONAL_ADMIN_ROLE, "region": scope_region}
        )
    else:
        total_students = await db.users.count_documents({"role": "student"})
        total_admins = await db.users.count_documents(
            {"role": {"$in": ["admin", "super_admin", REGIONAL_ADMIN_ROLE]}}
        )
    # trend by month (submitted_at)
    trend = {}
    for r in regs:
        ts = r.get("submitted_at") or r.get("created_at")
        if ts:
            month = ts[:7]
            trend[month] = trend.get(month, 0) + 1
    trend_list = [{"month": k, "count": v} for k, v in sorted(trend.items())]
    passed_by_region = {}
    for registration in regs:
        if registration.get("status") != "lolos":
            continue
        passed_region = region_name(registration)
        passed_by_region[passed_region] = passed_by_region.get(passed_region, 0) + 1
    passed_region_data = [
        {"region": region, "count": count}
        for region, count in sorted(passed_by_region.items(), key=lambda item: (-item[1], item[0]))
    ]
    site_content = await db.site_content.find_one({"key": "main"}, {"_id": 0}) or {}
    return {
        "total_registrations": total, "total_students": total_students, "total_admins": total_admins,
        "by_status": by_status, "trend": trend_list,
        "verified": by_status.get("lolos", 0),
        "pending": by_status.get("submitted", 0) + by_status.get("verifikasi", 0),
        "passed_by_region": passed_region_data,
        "available_regions": available_regions,
        "announcement_count": len(site_content.get("announcements", [])),
        "selected_region": selected_region,
    }


@api_router.get("/admin/participants")
async def list_participants(
    status: Optional[str] = None,
    search: Optional[str] = None,
    region: Optional[str] = None,
    user: dict = Depends(require_roles(*MANAGEMENT_ROLES)),
):
    scope_region = regional_admin_region(user)
    query = await participant_scope_query(user)
    query["is_archived"] = {"$ne": True}
    if status and status != "all":
        query["status"] = status
    regs = await db.registrations.find(query, {"_id": 0}).sort("created_at", -1).to_list(2000)
    result = []
    for r in regs:
        profile = await db.profiles.find_one({"user_id": r["user_id"]}, {"_id": 0})
        pdata = (profile or {}).get("data", {})
        doc_count = await db.documents.count_documents({"user_id": r["user_id"], "is_deleted": False})
        item = {
            **r,
            "institusi": pdata.get("institusi", "-"),
            "jenjang": pdata.get("jenjang", "-"),
            "nim": pdata.get("nim", "-"),
            "jurusan": pdata.get("jurusan", "-"),
            "semester": pdata.get("semester", "-"),
            "ipk": pdata.get("ipk", "-"),
            "biaya_pendidikan_semester": pdata.get("biayaPendidikanSemester", "-"),
            "wilayah": pdata.get("kota") or pdata.get("provinsi") or "-",
            "provinsi": pdata.get("provinsi", "-"),
            "phone": pdata.get("noTelp", "-"),
            "doc_count": doc_count,
        }
        selected_region = scope_region or region
        if selected_region and item["wilayah"] != selected_region:
            continue
        if search:
            s = search.lower()
            if s not in (r.get("name", "").lower() + r.get("email", "").lower() + pdata.get("institusi", "").lower()):
                continue
        result.append(item)
    return result


@api_router.get("/admin/participants/export.xlsx")
async def export_participants_excel(
    status: Optional[str] = None,
    search: Optional[str] = None,
    region: Optional[str] = None,
    user: dict = Depends(require_roles(*MANAGEMENT_ROLES)),
):
    participants = await list_participants(status, search, region, user)
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Data Peserta MDJ"
    headers = [
        "No.",
        "Nama",
        "Email",
        "ID CPM",
        "Kampus",
        "NIM",
        "Jenjang",
        "Jurusan",
        "Semester",
        "IPK",
        "Biaya Pendidikan / Semester (Rp)",
        "Wilayah",
        "Provinsi",
        "Nomor Telepon",
        "Dokumen Diunggah",
        "Status Pendaftaran",
        "Tanggal Daftar",
    ]
    worksheet.append(headers)
    header_fill = PatternFill("solid", fgColor="0B6B3A")
    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill

    for index, participant in enumerate(participants, start=1):
        worksheet.append([
            index,
            participant.get("name", "-"),
            participant.get("email", "-"),
            participant.get("cpm_id", "-"),
            participant.get("institusi", "-"),
            participant.get("nim", "-"),
            participant.get("jenjang", "-"),
            participant.get("jurusan", "-"),
            participant.get("semester", "-"),
            participant.get("ipk", "-"),
            participant.get("biaya_pendidikan_semester", "-"),
            participant.get("wilayah", "-"),
            participant.get("provinsi", "-"),
            participant.get("phone", "-"),
            participant.get("doc_count", 0),
            STATUS_LABELS.get(participant.get("status"), participant.get("status", "-")),
            participant.get("created_at", "-"),
        ])

    widths = [6, 28, 34, 24, 34, 18, 12, 24, 12, 10, 30, 20, 16, 20, 18, 28, 24]
    for index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[chr(64 + index)].width = width
    worksheet.freeze_panes = "A2"

    stream = io.BytesIO()
    workbook.save(stream)
    stream.seek(0)
    filename = f"data-peserta-mdj-{datetime.now(timezone.utc).strftime('%Y%m%d-%H%M%S')}.xlsx"
    headers_response = {"Content-Disposition": f'attachment; filename="{filename}"'}
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=headers_response,
    )


@api_router.post("/admin/demo-students/seed")
async def seed_demo_students(
    user: dict = Depends(require_roles("super_admin")),
):
    seed_batch = "mdj-student-sample-2026"
    campuses = await db.campuses.find({}, {"_id": 0, "id": 1, "name": 1}).to_list(5000)
    if not campuses:
        raise HTTPException(status_code=400, detail="Impor data kampus sebelum membuat data uji.")

    statuses = (
        ["lolos"] * 12
        + ["submitted"] * 8
        + ["verifikasi"] * 8
        + ["lolos_administrasi"] * 7
        + ["wawancara"] * 6
        + ["verifikasi_faktual"] * 5
        + ["ditolak"] * 4
    )
    regions = [
        "Jakarta Pusat",
        "Jakarta Utara",
        "Jakarta Barat",
        "Jakarta Selatan",
        "Jakarta Timur",
        "Kepulauan Seribu",
    ]
    majors = [
        "Ilmu Komputer",
        "Akuntansi",
        "Manajemen",
        "Ilmu Komunikasi",
        "Kesehatan Masyarakat",
        "Teknik Informatika",
    ]
    first_names = [
        "Alya", "Bagas", "Citra", "Dimas", "Erika", "Fajar", "Gita", "Hafiz", "Indah", "Jihan",
    ]
    last_names = ["Pratama", "Lestari", "Saputra", "Permata", "Nugroho"]
    generator = random.Random(20260912)
    generator.shuffle(statuses)
    generator.shuffle(campuses)
    for region_index in range(len(regions)):
        if statuses[region_index] == "lolos":
            continue
        passed_index = statuses.index("lolos", len(regions))
        statuses[region_index], statuses[passed_index] = statuses[passed_index], statuses[region_index]

    created = 0
    skipped = 0
    by_status = {}
    for index in range(50):
        user_id = f"student_uji_mdj_{index + 1:03d}"
        existing = await db.registrations.find_one({"user_id": user_id}, {"_id": 0, "id": 1})
        if existing:
            await db.registrations.update_one(
                {"user_id": user_id},
                {
                    "$set": {
                        "status": statuses[index],
                        "updated_at": datetime.now(timezone.utc).isoformat(),
                    },
                },
            )
            by_status[statuses[index]] = by_status.get(statuses[index], 0) + 1
            skipped += 1
            continue

        status = statuses[index]
        region = regions[index % len(regions)]
        campus = campuses[index % len(campuses)]
        name = f"{first_names[index % len(first_names)]} {last_names[index % len(last_names)]}"
        email = f"uji.mdj.{index + 1:03d}@contoh.invalid"
        created_at = datetime(2026, 8, 1, tzinfo=timezone.utc) + timedelta(days=index)
        timestamp = created_at.isoformat()
        user_record = await db.users.find_one({"user_id": user_id}, {"_id": 0, "user_id": 1})
        if not user_record:
            await db.users.insert_one({
                "user_id": user_id,
                "email": email,
                "password_hash": hash_password(secrets.token_urlsafe(32)),
                "name": name,
                "role": "student",
                "auth_provider": "demo_seed",
                "is_active": True,
                "is_test_data": True,
                "seed_batch": seed_batch,
                "created_at": timestamp,
            })

        profile = {
            "namaLengkap": name,
            "email": email,
            "noTelp": f"+62812{index + 1000000:07d}",
            "institusi": campus["name"],
            "jurusan": majors[index % len(majors)],
            "jenjang": "S1",
            "semester": str((index % 8) + 1),
            "ipk": f"{3.10 + ((index * 7) % 80) / 100:.2f}",
            "biayaPendidikanSemester": str(2500000 + (index % 7) * 500000),
            "kota": region,
            "provinsi": "DKI Jakarta",
        }
        await db.profiles.update_one(
            {"user_id": user_id},
            {"$set": {"data": profile, "updated_at": timestamp, "is_test_data": True}},
            upsert=True,
        )
        cpm_id = await create_cpm_id(timestamp)
        registration = {
            "id": str(uuid.uuid4()),
            "user_id": user_id,
            "name": name,
            "email": email,
            "category": "Mahasiswa Sarjana (S1)",
            "status": status,
            "cpm_id": cpm_id,
            "history": [{"status": status, "note": "Data uji dibuat", "at": timestamp}],
            "created_at": timestamp,
            "submitted_at": timestamp,
            "updated_at": timestamp,
            "is_test_data": True,
            "seed_batch": seed_batch,
        }
        await db.registrations.insert_one(registration)
        created += 1
        by_status[status] = by_status.get(status, 0) + 1

    return {
        "message": "Data uji mahasiswa berhasil disiapkan.",
        "created": created,
        "skipped": skipped,
        "total": created + skipped,
        "by_status": by_status,
        "regions": regions,
    }


@api_router.get("/admin/participants/{user_id}")
async def participant_detail(
    user_id: str,
    user: dict = Depends(require_roles(*MANAGEMENT_ROLES)),
):
    await require_participant_scope(user, user_id)
    reg = await db.registrations.find_one({"user_id": user_id}, {"_id": 0})
    if not reg:
        raise HTTPException(status_code=404, detail="Pendaftar tidak ditemukan")
    profile = await db.profiles.find_one({"user_id": user_id}, {"_id": 0})
    docs = await db.documents.find({"user_id": user_id, "is_deleted": False}, {"_id": 0}).to_list(200)
    account = await db.users.find_one({"user_id": user_id}, {"_id": 0, "password_hash": 0})
    for d in docs:
        d["url"] = f"/api/files/{d['storage_path']}"
    return {"registration": reg or {}, "profile": (profile or {}).get("data", {}),
            "documents": docs, "account": account or {}}


@api_router.put("/admin/participants/{user_id}/status")
async def update_participant_status(user_id: str, payload: StatusUpdateInput,
                                    user: dict = Depends(require_roles(*MANAGEMENT_ROLES))):
    if payload.status not in REG_STATUSES:
        raise HTTPException(status_code=400, detail="Status tidak valid")
    await require_participant_scope(user, user_id)
    reg = await db.registrations.find_one({"user_id": user_id})
    if not reg:
        raise HTTPException(status_code=404, detail="Pendaftar tidak ditemukan")
    now = datetime.now(timezone.utc).isoformat()
    entry = {"status": payload.status, "note": payload.note or "", "at": now, "by": user.get("name")}
    await db.registrations.update_one(
        {"user_id": user_id},
        {"$set": {"status": payload.status, "updated_at": now}, "$push": {"history": entry}},
    )
    if payload.status != reg.get("status"):
        label = STATUS_LABELS.get(payload.status, payload.status)
        await create_notification(
            user_id,
            "selection_progress",
            "Perkembangan seleksi diperbarui",
            f"Status seleksi Anda diperbarui menjadi {label}.",
        )
    updated = await db.registrations.find_one({"user_id": user_id}, {"_id": 0})
    return updated


@api_router.put("/admin/participants/{user_id}/document-editing-permission")
async def update_document_editing_permission(
    user_id: str,
    payload: DocumentEditingPermissionInput,
    user: dict = Depends(require_roles("admin", "super_admin")),
):
    if payload.allowed:
        raise HTTPException(
            status_code=400,
            detail="Gunakan pengembalian berkas dengan catatan perbaikan.",
        )
    registration = await db.registrations.find_one({"user_id": user_id}, {"_id": 0})
    if not registration:
        raise HTTPException(status_code=404, detail="Pendaftar tidak ditemukan")
    now = datetime.now(timezone.utc).isoformat()
    permission_note = (
        "Panitia mengizinkan perbaikan dokumen."
        if payload.allowed
        else "Panitia mengunci kembali perbaikan dokumen."
    )
    await db.registrations.update_one(
        {"user_id": user_id},
        {
            "$set": {
                "documents_editing_allowed": payload.allowed,
                "documents_editing_permission_updated_at": now,
                "documents_editing_permission_updated_by": user.get("name", "Panitia"),
                "updated_at": now,
            },
            "$push": {
                "history": {
                    "status": registration.get("status", "submitted"),
                    "note": permission_note,
                    "at": now,
                    "by": user.get("name", "Panitia"),
                }
            },
        },
    )
    await create_notification(
        user_id,
        "document_edit_permission",
        "Izin edit dokumen diperbarui",
        permission_note,
    )
    updated = await db.registrations.find_one({"user_id": user_id}, {"_id": 0})
    return updated


@api_router.post("/admin/participants/{user_id}/return-documents")
async def return_documents_for_revision(
    user_id: str,
    user: dict = Depends(require_roles("admin", "super_admin")),
):
    note = DOCUMENT_REVISION_DEFAULT_NOTE
    registration = await db.registrations.find_one({"user_id": user_id}, {"_id": 0})
    if not registration:
        raise HTTPException(status_code=404, detail="Pendaftar tidak ditemukan")
    if registration.get("status") not in {"submitted", "verifikasi", "perlu_perbaikan"}:
        raise HTTPException(
            status_code=400,
            detail="Berkas hanya dapat dikembalikan setelah pendaftaran dikirim.",
        )
    now = datetime.now(timezone.utc).isoformat()
    history_entry = {
        "status": "perlu_perbaikan",
        "note": note,
        "at": now,
        "by": user.get("name", "Panitia"),
    }
    await db.registrations.update_one(
        {"user_id": user_id},
        {
            "$set": {
                "status": "perlu_perbaikan",
                "documents_editing_allowed": True,
                "revision_requested": True,
                "revision_note": note,
                "revision_requested_at": now,
                "revision_requested_by": user.get("name", "Panitia"),
                "revision_requested_by_user_id": user["user_id"],
                "updated_at": now,
            },
            "$push": {"history": history_entry},
        },
    )
    await create_notification(
        user_id,
        "document_revision_requested",
        "Berkas Perlu Dilengkapi",
        f"Panitia mengembalikan berkas Anda untuk diperbaiki. Catatan: {note}",
        {"revision_note": note},
    )
    updated = await db.registrations.find_one({"user_id": user_id}, {"_id": 0})
    return updated


@api_router.get("/admin/selection-announcements/summary")
async def get_selection_announcement_summary(
    category: str = "all",
    user: dict = Depends(require_roles("admin", "super_admin")),
):
    return await selection_announcement_summary(category)


@api_router.post("/admin/selection-announcements/publish")
async def publish_selection_announcement(
    payload: SelectionAnnouncementPublishInput,
    user: dict = Depends(require_roles("admin", "super_admin")),
):
    summary = await selection_announcement_summary(payload.category)
    if not summary["recipient_count"]:
        raise HTTPException(status_code=400, detail="Tidak ada hasil seleksi baru yang dapat diumumkan.")

    selected_category = summary["category"]
    category_query = {}
    if selected_category != "all":
        category_query["category"] = selected_category
    recipient_query = {
        **category_query,
        "status": {"$in": SELECTION_RESULT_STATUSES},
        "$or": [
            {"is_announcement_published": {"$exists": False}},
            {"is_announcement_published": False},
        ],
    }
    registrations = await db.registrations.find(
        recipient_query,
        {"_id": 0, "user_id": 1, "status": 1, "category": 1},
    ).to_list(5000)
    if not registrations:
        raise HTTPException(status_code=400, detail="Hasil seleksi sudah tidak tersedia untuk diumumkan.")

    now = datetime.now(timezone.utc).isoformat()
    announcement_id = str(uuid.uuid4())
    announcement = {
        "id": announcement_id,
        "category": selected_category,
        "status_publikasi": "Published",
        "title": "Pengumuman Hasil Seleksi Berkas",
        "message": (
            "Pengumuman hasil seleksi tahap berkas telah diterbitkan oleh Admin Provinsi. "
            "Silakan cek status Anda."
        ),
        "passed_count": summary["passed_count"],
        "failed_count": summary["failed_count"],
        "recipient_count": len(registrations),
        "published_by": user["user_id"],
        "published_by_name": user.get("name", "Admin Provinsi"),
        "published_at": now,
        "created_at": now,
    }
    await db.selection_announcements.insert_one(dict(announcement))

    notifications = []
    for registration in registrations:
        result = (
            "passed"
            if registration.get("status") in SELECTION_RESULT_PASSED_STATUSES
            else "failed"
        )
        notifications.append({
            "id": str(uuid.uuid4()),
            "user_id": registration["user_id"],
            "type": "selection_result",
            "source_id": announcement_id,
            "announcement_id": announcement_id,
            "category": registration.get("category", ""),
            "result": result,
            "title": announcement["title"],
            "message": announcement["message"],
            "is_read": False,
            "is_popup_seen": False,
            "created_at": now,
        })
    await db.notifications.insert_many(notifications)
    await db.registrations.update_many(
        {"user_id": {"$in": [item["user_id"] for item in registrations]}},
        {
            "$set": {
                "is_announcement_published": True,
                "selection_announcement_id": announcement_id,
                "selection_announcement_published_at": now,
            }
        },
    )
    return {
        "message": "Pengumuman hasil seleksi berhasil dikirim.",
        "announcement_id": announcement_id,
        "status_publikasi": "Published",
        "passed_count": summary["passed_count"],
        "failed_count": summary["failed_count"],
        "recipient_count": len(registrations),
    }


# ---------------------------------------------------------------------------
# Super admin: user/admin management
# ---------------------------------------------------------------------------
@api_router.get("/admin/users")
async def list_users(
    role: Optional[str] = None,
    user: dict = Depends(require_roles("admin", "super_admin")),
):
    query = {}
    if user.get("role") == "admin":
        query["role"] = REGIONAL_ADMIN_ROLE
    elif role and role != "all":
        query["role"] = role
    users = await db.users.find(query, {"_id": 0, "password_hash": 0}).sort("created_at", -1).to_list(2000)
    return users


@api_router.post("/admin/users")
async def create_admin(
    payload: CreateAdminInput,
    user: dict = Depends(require_roles("admin", "super_admin")),
):
    is_provincial_admin = user.get("role") == "admin"
    allowed_roles = [REGIONAL_ADMIN_ROLE] if is_provincial_admin else [
        "admin",
        "super_admin",
        REGIONAL_ADMIN_ROLE,
    ]
    if payload.role not in allowed_roles:
        raise HTTPException(status_code=400, detail="Role tidak valid")
    region = str(payload.region or "").strip()
    if payload.role == REGIONAL_ADMIN_ROLE and region not in DKI_REGIONS:
        raise HTTPException(status_code=400, detail="Wilayah Admin Wilayah wajib dipilih.")
    if payload.role != REGIONAL_ADMIN_ROLE:
        region = None
    email = payload.email.lower().strip()
    if await db.users.find_one({"email": email}):
        raise HTTPException(status_code=400, detail="Email sudah terdaftar")
    user_id = f"user_{uuid.uuid4().hex[:12]}"
    doc = {
        "user_id": user_id, "email": email, "password_hash": hash_password(payload.password),
        "name": payload.name, "role": payload.role, "auth_provider": "password",
        "nik": None, "phone": None, "picture": None, "is_active": True,
        "region": region,
        "created_at": datetime.now(timezone.utc).isoformat(),
    }
    await db.users.insert_one(doc)
    return clean_user(doc)


@api_router.put("/admin/users/{user_id}")
async def toggle_user(
    user_id: str,
    payload: UserToggleInput,
    user: dict = Depends(require_roles("admin", "super_admin")),
):
    if user_id == user["user_id"]:
        raise HTTPException(status_code=400, detail="Tidak dapat menonaktifkan akun sendiri")
    target = await db.users.find_one({"user_id": user_id}, {"_id": 0})
    if not target or not can_manage_admin_account(user, target):
        raise HTTPException(status_code=404, detail="Akun admin tidak ditemukan")
    await db.users.update_one({"user_id": user_id}, {"$set": {"is_active": payload.is_active}})
    return {"message": "Status akun diperbarui"}


@api_router.delete("/admin/users/{user_id}")
async def delete_user(
    user_id: str,
    user: dict = Depends(require_roles("admin", "super_admin")),
):
    target = await db.users.find_one({"user_id": user_id})
    if not target or not can_manage_admin_account(user, target):
        raise HTTPException(status_code=404, detail="Akun admin tidak ditemukan")
    if target.get("email") == ADMIN_EMAIL:
        raise HTTPException(status_code=400, detail="Tidak dapat menghapus super admin utama")
    await db.users.delete_one({"user_id": user_id})
    return {"message": "User dihapus"}


@api_router.get("/super-admin/applicant-archive/summary")
async def get_applicant_archive_summary(
    user: dict = Depends(require_roles("super_admin")),
):
    active_count = await db.users.count_documents(
        {"role": "student", "is_archived": {"$ne": True}}
    )
    archived_count = await db.users.count_documents(
        {"role": "student", "is_archived": True}
    )
    restorable_batch = await db.applicant_archive_batches.find_one(
        {"status": "archived"},
        {"_id": 0},
        sort=[("archived_at", -1)],
    )
    return {
        "active_applicant_count": active_count,
        "archived_applicant_count": archived_count,
        "restorable_batch": restorable_batch,
    }


@api_router.post("/super-admin/applicant-archive")
async def archive_all_applicants(user: dict = Depends(require_roles("super_admin"))):
    student_cursor = db.users.find(
        {"role": "student", "is_archived": {"$ne": True}},
        {"_id": 0, "user_id": 1},
    )
    user_ids = [student["user_id"] async for student in student_cursor]
    if not user_ids:
        raise HTTPException(status_code=400, detail="Tidak ada data pendaftar aktif untuk diarsipkan.")

    batch_id = str(uuid.uuid4())
    now = datetime.now(timezone.utc).isoformat()
    scope = {"user_id": {"$in": user_ids}}
    counts = {
        "applicant_count": len(user_ids),
        "registration_count": await db.registrations.count_documents(scope),
        "profile_count": await db.profiles.count_documents(scope),
        "document_count": await db.documents.count_documents(scope),
        "notification_count": await db.notifications.count_documents(scope),
    }
    archive_fields = {
        "is_archived": True,
        "archive_batch_id": batch_id,
        "archived_at": now,
        "archived_by": user.get("user_id"),
    }
    await db.users.update_many(
        {**scope, "role": "student"},
        {"$set": archive_fields, "$inc": {"auth_version": 1}},
    )
    for collection in [db.registrations, db.profiles, db.documents, db.notifications]:
        await collection.update_many(scope, {"$set": archive_fields})
    await db.user_sessions.delete_many(scope)
    batch = {
        "id": batch_id,
        "status": "archived",
        **counts,
        "archived_at": now,
        "archived_by": user.get("user_id"),
        "archived_by_name": user.get("name", "Super Admin"),
    }
    await db.applicant_archive_batches.insert_one(dict(batch))
    return {"message": "Seluruh data pendaftar telah diarsipkan.", "batch": batch}


@api_router.post("/super-admin/applicant-archive/{batch_id}/restore")
async def restore_applicant_archive(
    batch_id: str,
    user: dict = Depends(require_roles("super_admin")),
):
    batch = await db.applicant_archive_batches.find_one(
        {"id": batch_id, "status": "archived"},
        {"_id": 0},
    )
    if not batch:
        raise HTTPException(status_code=404, detail="Batch arsip aktif tidak ditemukan.")
    now = datetime.now(timezone.utc).isoformat()
    restore_fields = {"is_archived": False, "restored_at": now, "restored_by": user.get("user_id")}
    archive_scope = {"archive_batch_id": batch_id, "is_archived": True}
    for collection in [db.users, db.registrations, db.profiles, db.documents, db.notifications]:
        await collection.update_many(
            archive_scope,
            {"$set": restore_fields, "$unset": {"archive_batch_id": ""}},
        )
    await db.applicant_archive_batches.update_one(
        {"id": batch_id},
        {
            "$set": {
                "status": "restored",
                "restored_at": now,
                "restored_by": user.get("user_id"),
                "restored_by_name": user.get("name", "Super Admin"),
            }
        },
    )
    return {"message": "Data pendaftar berhasil dipulihkan.", "batch_id": batch_id}


@api_router.post("/super-admin/student-data/purge")
async def purge_non_demo_student_data(
    payload: StudentDataPurgeInput,
    user: dict = Depends(require_roles("super_admin")),
):
    expected_confirmation = (
        "HAPUS SEMUA DATA PENDAFTAR"
        if payload.include_orphaned_registrations
        else "HAPUS SEMUA DATA MAHASISWA"
    )
    if payload.confirmation.strip() != expected_confirmation:
        raise HTTPException(
            status_code=400,
            detail="Konfirmasi penghapusan permanen tidak sesuai.",
        )

    demo_email = DEMO_STUDENT_EMAIL.lower()
    demo_user = await db.users.find_one(
        {"email": demo_email, "role": "student"},
        {"_id": 0, "user_id": 1},
    )
    if not demo_user:
        raise HTTPException(status_code=404, detail="Akun demo mahasiswa tidak ditemukan.")

    demo_user_id = demo_user["user_id"]
    student_cursor = db.users.find(
        {"role": "student", "email": {"$ne": demo_email}},
        {"_id": 0, "user_id": 1, "email": 1},
    )
    students = [student async for student in student_cursor]
    registration_scope = {"user_id": {"$ne": demo_user_id}}
    if not payload.include_orphaned_registrations:
        registration_scope["is_test_data"] = True
    test_registration_cursor = db.registrations.find(
        registration_scope,
        {"_id": 0, "user_id": 1},
    )
    test_registration_user_ids = [
        registration["user_id"]
        async for registration in test_registration_cursor
    ]
    user_ids = list({
        *(student["user_id"] for student in students),
        *test_registration_user_ids,
    })
    emails = [student["email"] for student in students]
    if not user_ids and not payload.include_orphaned_registrations:
        await db.users.update_one(
            {"user_id": demo_user_id},
            {
                "$set": {"is_active": True, "is_archived": False},
                "$unset": {
                    "archive_batch_id": "",
                    "archived_at": "",
                    "archived_by": "",
                },
            },
        )
        return {
            "message": "Tidak ada data mahasiswa non-demo yang perlu dihapus.",
            "deleted_student_count": 0,
            "deleted_test_registration_count": 0,
            "deleted_records": {},
        }

    user_scope = {"user_id": {"$in": user_ids}}
    deleted_records = {}
    direct_user_collections = [
        "profiles",
        "registrations",
        "documents",
        "notifications",
        "user_sessions",
        "verification_recommendations",
        "beneficiary_records",
        "beneficiary_disbursements",
        "active_letter_approvals",
        "stage_ii_eligibilities",
        "beneficiary_audit_events",
    ]
    for collection_name in direct_user_collections:
        deleted_records[collection_name] = 0
    if user_ids:
        for collection_name in direct_user_collections:
            result = await db[collection_name].delete_many(user_scope)
            deleted_records[collection_name] = result.deleted_count

    if payload.include_orphaned_registrations:
        known_user_ids = [
            account["user_id"]
            async for account in db.users.find({}, {"_id": 0, "user_id": 1})
        ]
        orphan_scope = {
            "user_id": {"$exists": True, "$nin": known_user_ids},
        }
        for collection_name in direct_user_collections:
            result = await db[collection_name].delete_many(orphan_scope)
            deleted_records[collection_name] += result.deleted_count

    for collection_name in ["password_reset_tokens", "email_change_tokens"]:
        deleted_records[collection_name] = 0
        if emails:
            result = await db[collection_name].delete_many({"email": {"$in": emails}})
            deleted_records[collection_name] = result.deleted_count

    deleted_source_ids = []
    source_cursor = db.beneficiary_sources.find(
        {"matched_user_ids": {"$in": user_ids}},
        {"_id": 0, "id": 1, "matched_user_ids": 1},
    )
    async for source in source_cursor:
        remaining_user_ids = [
            source_user_id
            for source_user_id in source.get("matched_user_ids", [])
            if source_user_id not in user_ids
        ]
        if remaining_user_ids:
            await db.beneficiary_sources.update_one(
                {"id": source["id"]},
                {"$set": {"matched_user_ids": remaining_user_ids}},
            )
            continue
        await db.beneficiary_sources.delete_one({"id": source["id"]})
        deleted_source_ids.append(source["id"])
    deleted_records["beneficiary_sources"] = len(deleted_source_ids)

    if deleted_source_ids:
        result = await db.beneficiary_reviews.delete_many(
            {"source_id": {"$in": deleted_source_ids}}
        )
        deleted_records["beneficiary_reviews"] = result.deleted_count
    else:
        deleted_records["beneficiary_reviews"] = 0

    deleted_disbursement_aggregates = 0
    disbursement_cursor = db.campus_disbursements.find(
        {"recipient_user_ids": {"$in": user_ids}},
        {"_id": 0, "id": 1, "recipient_user_ids": 1},
    )
    async for disbursement in disbursement_cursor:
        remaining_user_ids = [
            recipient_user_id
            for recipient_user_id in disbursement.get("recipient_user_ids", [])
            if recipient_user_id not in user_ids
        ]
        if remaining_user_ids:
            await db.campus_disbursements.update_one(
                {"id": disbursement["id"]},
                {"$set": {"recipient_user_ids": remaining_user_ids}},
            )
            continue
        result = await db.campus_disbursements.delete_one({"id": disbursement["id"]})
        deleted_disbursement_aggregates += result.deleted_count
    deleted_records["campus_disbursements"] = deleted_disbursement_aggregates

    result = await db.users.delete_many(
        {"role": "student", "email": {"$ne": demo_email}},
    )
    deleted_records["users"] = result.deleted_count
    await db.users.update_one(
        {"user_id": demo_user_id},
        {
            "$set": {"is_active": True, "is_archived": False},
            "$unset": {
                "archive_batch_id": "",
                "archived_at": "",
                "archived_by": "",
            },
        },
    )
    await db.beneficiary_audit_events.insert_one({
        "id": str(uuid.uuid4()),
        "actor_id": user["user_id"],
        "actor_name": user.get("name", "Super Admin"),
        "action": "purge_non_demo_student_data",
        "detail": {
            "deleted_student_count": result.deleted_count,
            "deleted_test_registration_count": deleted_records["registrations"],
            "included_orphaned_registrations": payload.include_orphaned_registrations,
            "deleted_records": deleted_records,
        },
        "created_at": datetime.now(timezone.utc).isoformat(),
    })
    return {
        "message": "Seluruh data mahasiswa non-demo berhasil dihapus permanen.",
        "deleted_student_count": result.deleted_count,
        "deleted_test_registration_count": deleted_records["registrations"],
        "included_orphaned_registrations": payload.include_orphaned_registrations,
        "deleted_records": deleted_records,
        "preserved_demo_email": demo_email,
    }


@api_router.get("/admin/live-chat/messages")
async def list_admin_live_chat_messages(
    session_id: str = "admin-team",
    user: dict = Depends(require_roles(*MANAGEMENT_ROLES)),
):
    if session_id != "admin-team":
        raise HTTPException(status_code=404, detail="Sesi Live Chat tidak ditemukan.")
    messages = await db.admin_live_chat_messages.find(
        {"session_id": session_id},
        {"_id": 0},
    ).sort("created_at", -1).to_list(100)
    messages.reverse()
    return {"session_id": session_id, "messages": [serialize_admin_chat_message(item) for item in messages]}


@api_router.post("/admin/live-chat/messages")
async def create_admin_live_chat_message(
    payload: AdminLiveChatMessageInput,
    user: dict = Depends(require_roles(*MANAGEMENT_ROLES)),
):
    if payload.session_id != "admin-team":
        raise HTTPException(status_code=404, detail="Sesi Live Chat tidak ditemukan.")
    message = payload.message.strip()
    if not message:
        raise HTTPException(status_code=400, detail="Pesan Live Chat tidak boleh kosong.")
    record = {
        "id": str(uuid.uuid4()),
        "session_id": payload.session_id,
        "sender_id": user["user_id"],
        "sender_name": user.get("name", "Admin"),
        "sender_role": user.get("role"),
        "message": message,
        "created_at": datetime.now(timezone.utc).isoformat(),
    }
    await db.admin_live_chat_messages.insert_one(dict(record))
    return {"message": record}


def serialize_admin_chat_message(record: dict) -> dict:
    item = dict(record)
    item.pop("_id", None)
    if item.get("attachment"):
        attachment = item["attachment"]
        item["attachment"] = {
            "id": attachment["id"],
            "original_filename": attachment["original_filename"],
            "content_type": attachment["content_type"],
            "size": attachment["size"],
        }
    return item


@api_router.post("/admin/live-chat/messages/attachment")
async def create_admin_live_chat_attachment(
    message: str = Form(""),
    attachment: UploadFile = File(...),
    user: dict = Depends(require_roles(*MANAGEMENT_ROLES)),
):
    filename = attachment.filename or "lampiran"
    extension = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if extension not in ADMIN_CHAT_MIME_TYPES:
        raise HTTPException(status_code=400, detail="Format lampiran belum didukung.")
    data = await attachment.read()
    if not data or len(data) > 15 * 1024 * 1024:
        raise HTTPException(status_code=400, detail="Ukuran lampiran harus antara 1 byte dan 15MB.")
    content_type = ADMIN_CHAT_MIME_TYPES[extension]
    data, extension, content_type = compress_image(data, extension, content_type)
    attachment_id = str(uuid.uuid4())
    stored = put_object(
        f"{APP_NAME}/admin-live-chat/{attachment_id}.{extension}",
        data,
        content_type,
    )
    record = {
        "id": str(uuid.uuid4()),
        "session_id": "admin-team",
        "sender_id": user["user_id"],
        "sender_name": user.get("name", "Admin"),
        "sender_role": user.get("role"),
        "message": message.strip(),
        "attachment": {
            "id": attachment_id,
            "storage_path": stored["path"],
            "original_filename": filename,
            "content_type": content_type,
            "size": stored["size"],
        },
        "created_at": datetime.now(timezone.utc).isoformat(),
    }
    await db.admin_live_chat_messages.insert_one(dict(record))
    return {"message": serialize_admin_chat_message(record)}


@api_router.get("/admin/live-chat/attachments/{attachment_id}")
async def download_admin_live_chat_attachment(
    attachment_id: str,
    request: Request,
    auth: str = Query(None),
):
    token = request.cookies.get("access_token") or request.cookies.get("session_token") or auth
    if not token:
        authorization = request.headers.get("Authorization", "")
        token = authorization[7:] if authorization.startswith("Bearer ") else None
    user = await resolve_user_from_token(token) if token else None
    if not user or user.get("role") not in MANAGEMENT_ROLES:
        raise HTTPException(status_code=401, detail="Not authenticated")
    record = await db.admin_live_chat_messages.find_one(
        {"attachment.id": attachment_id},
        {"_id": 0},
    )
    if not record or not record.get("attachment"):
        raise HTTPException(status_code=404, detail="Lampiran Live Chat tidak ditemukan.")
    attachment = record["attachment"]
    data, fallback_type = get_object(attachment["storage_path"])
    return StarletteResponse(content=data, media_type=attachment.get("content_type", fallback_type))


async def resolve_student_chat_target(user: dict, student_id: Optional[str]) -> dict:
    role = user.get("role")
    if role == "student":
        if student_id and student_id != user["user_id"]:
            raise HTTPException(status_code=404, detail="Percakapan mahasiswa tidak ditemukan.")
        return user
    if role not in MANAGEMENT_ROLES:
        raise HTTPException(status_code=403, detail="Akses Live Chat mahasiswa ditolak.")
    if not student_id:
        raise HTTPException(status_code=400, detail="Pilih mahasiswa untuk membuka percakapan.")

    student = await db.users.find_one(
        {"user_id": student_id, "role": "student", "is_archived": {"$ne": True}},
        {"_id": 0},
    )
    if not student:
        raise HTTPException(status_code=404, detail="Percakapan mahasiswa tidak ditemukan.")

    region_scope = regional_admin_region(user)
    if region_scope:
        profile = await db.profiles.find_one({"user_id": student_id}, {"_id": 0, "data": 1})
        student_region = profile_region((profile or {}).get("data", {}))
        if student_region != region_scope:
            raise HTTPException(status_code=404, detail="Percakapan mahasiswa tidak ditemukan.")
    return student


def student_chat_attachment_matches(data: bytes, extension: str) -> bool:
    signatures = {
        "pdf": data.startswith(b"%PDF"),
        "jpg": data.startswith(b"\xff\xd8\xff"),
        "jpeg": data.startswith(b"\xff\xd8\xff"),
        "png": data.startswith(b"\x89PNG\r\n\x1a\n"),
        "webp": data.startswith(b"RIFF") and data[8:12] == b"WEBP",
    }
    return bool(data) and signatures.get(extension, False)


async def create_student_chat_attachment(file: UploadFile, student_id: str) -> dict:
    filename = file.filename or "lampiran"
    extension = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if extension not in STUDENT_CHAT_ALLOWED_EXTENSIONS:
        raise HTTPException(status_code=400, detail="Gunakan JPG, PNG, WEBP, atau PDF.")
    data = await file.read()
    if not data or len(data) > STUDENT_CHAT_MAX_FILE_SIZE:
        raise HTTPException(status_code=400, detail="Ukuran lampiran harus antara 1 byte dan 10MB.")
    if not student_chat_attachment_matches(data, extension):
        raise HTTPException(status_code=400, detail="Isi lampiran tidak sesuai dengan format berkas.")

    content_type = MIME_TYPES[extension]
    data, extension, content_type = compress_image(data, extension, content_type)
    attachment_id = str(uuid.uuid4())
    path = f"{APP_NAME}/student-live-chat/{student_id}/{attachment_id}.{extension}"
    result = put_object(path, data, content_type)
    return {
        "id": attachment_id,
        "storage_path": result["path"],
        "original_filename": filename,
        "content_type": content_type,
        "size": result["size"],
    }


def serialize_student_chat_message(record: dict) -> dict:
    item = dict(record)
    item.pop("_id", None)
    attachment = item.get("attachment")
    if attachment:
        item["attachment"] = {
            "id": attachment["id"],
            "original_filename": attachment["original_filename"],
            "content_type": attachment["content_type"],
            "size": attachment["size"],
        }
    return item


@api_router.get("/student-live-chat/conversations")
async def list_student_live_chat_conversations(
    user: dict = Depends(require_roles(*MANAGEMENT_ROLES)),
):
    student_query = {"role": "student", "is_archived": {"$ne": True}}
    region_scope = regional_admin_region(user)
    if region_scope:
        profiles = await db.profiles.find(
            {"$or": [{"data.kota": region_scope}, {"data.provinsi": region_scope}]},
            {"_id": 0, "user_id": 1},
        ).to_list(5000)
        student_query["user_id"] = {"$in": [profile["user_id"] for profile in profiles]}
    students = await db.users.find(
        student_query,
        {"_id": 0, "user_id": 1, "name": 1, "email": 1},
    ).to_list(5000)
    student_ids = [student["user_id"] for student in students]
    if not student_ids:
        return {"conversations": []}

    profiles = await db.profiles.find(
        {"user_id": {"$in": student_ids}},
        {"_id": 0, "user_id": 1, "data.kota": 1, "data.provinsi": 1},
    ).to_list(5000)
    region_by_student = {
        profile["user_id"]: profile_region(profile.get("data", {}))
        for profile in profiles
    }
    messages = await db.student_live_chat_messages.find(
        {"student_id": {"$in": student_ids}},
        {"_id": 0},
    ).sort("created_at", -1).to_list(1000)
    latest_by_student = {}
    for message in messages:
        latest_by_student.setdefault(message["student_id"], message)

    conversations = []
    for student in students:
        latest = latest_by_student.get(student["user_id"])
        if not latest:
            continue
        conversations.append({
            "student_id": student["user_id"],
            "student_name": student.get("name", "Mahasiswa"),
            "student_email": student.get("email", ""),
            "region": region_by_student.get(student["user_id"], ""),
            "last_message": latest.get("message", "") or "Lampiran dikirim.",
            "last_message_at": latest.get("created_at"),
        })
    conversations.sort(key=lambda item: item.get("last_message_at") or "", reverse=True)
    return {"conversations": conversations}


@api_router.get("/student-live-chat/messages")
async def list_student_live_chat_messages(
    student_id: Optional[str] = None,
    user: dict = Depends(get_current_user),
):
    student = await resolve_student_chat_target(user, student_id)
    messages = await db.student_live_chat_messages.find(
        {"student_id": student["user_id"]},
        {"_id": 0},
    ).sort("created_at", 1).to_list(500)
    return {
        "student": {
            "user_id": student["user_id"],
            "name": student.get("name", "Mahasiswa"),
        },
        "messages": [serialize_student_chat_message(message) for message in messages],
    }


@api_router.post("/student-live-chat/messages")
async def create_student_live_chat_message(
    message: str = Form(""),
    student_id: Optional[str] = Form(None),
    attachment: Optional[UploadFile] = File(None),
    user: dict = Depends(get_current_user),
):
    if user.get("role") not in {*MANAGEMENT_ROLES, "student"}:
        raise HTTPException(status_code=403, detail="Akses Live Chat mahasiswa ditolak.")
    student = await resolve_student_chat_target(user, student_id)
    cleaned_message = message.strip()
    if len(cleaned_message) > 2000:
        raise HTTPException(status_code=400, detail="Pesan maksimal terdiri dari 2.000 karakter.")
    if not cleaned_message and not attachment:
        raise HTTPException(status_code=400, detail="Tulis pesan atau lampirkan berkas terlebih dahulu.")

    chat_attachment = None
    if attachment:
        chat_attachment = await create_student_chat_attachment(attachment, student["user_id"])
    record = {
        "id": str(uuid.uuid4()),
        "student_id": student["user_id"],
        "sender_id": user["user_id"],
        "sender_name": user.get("name", "Pengguna MDJ"),
        "sender_role": user.get("role"),
        "message": cleaned_message,
        "attachment": chat_attachment,
        "created_at": datetime.now(timezone.utc).isoformat(),
    }
    await db.student_live_chat_messages.insert_one(dict(record))
    return {"message": serialize_student_chat_message(record)}


@api_router.get("/student-live-chat/attachments/{attachment_id}")
async def download_student_live_chat_attachment(
    attachment_id: str,
    request: Request,
    auth: str = Query(None),
):
    token = request.cookies.get("access_token") or request.cookies.get("session_token") or auth
    if not token:
        authorization = request.headers.get("Authorization", "")
        token = authorization[7:] if authorization.startswith("Bearer ") else None
    viewer = await resolve_user_from_token(token) if token else None
    if not viewer:
        raise HTTPException(status_code=401, detail="Not authenticated")
    record = await db.student_live_chat_messages.find_one(
        {"attachment.id": attachment_id},
        {"_id": 0},
    )
    if not record or not record.get("attachment"):
        raise HTTPException(status_code=404, detail="Lampiran Live Chat tidak ditemukan.")
    await resolve_student_chat_target(viewer, record["student_id"])
    attachment = record["attachment"]
    data, content_type = get_object(attachment["storage_path"])
    return StarletteResponse(
        content=data,
        media_type=attachment.get("content_type", content_type),
        headers={"Cache-Control": "private, max-age=3600"},
    )


AI_REFERENCE_MIME_TYPES = {
    "pdf": "application/pdf",
    "doc": "application/msword",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "xls": "application/vnd.ms-excel",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "jpg": "image/jpeg",
    "jpeg": "image/jpeg",
    "png": "image/png",
    "mp3": "audio/mpeg",
    "wav": "audio/wav",
    "ogg": "audio/ogg",
    "aac": "audio/aac",
    "flac": "audio/flac",
}


async def retrieve_ai_reference_chunks(question: str) -> list[dict]:
    terms = reference_tokens(question)
    if not terms:
        return []
    chunks = await db.ai_reference_chunks.find(
        {},
        {"_id": 0, "reference_name": 1, "text": 1},
    ).to_list(5000)
    ranked = []
    for chunk in chunks:
        score = len(terms.intersection(reference_tokens(chunk["text"])))
        if score:
            ranked.append((score, chunk))
    ranked.sort(key=lambda item: item[0], reverse=True)
    return [item[1] for item in ranked[:5]]


@api_router.get("/super-admin/ai-references")
async def list_ai_references(user: dict = Depends(require_roles("super_admin"))):
    references = await db.ai_references.find({}, {"_id": 0, "storage_path": 0}).sort(
        "created_at", -1
    ).to_list(200)
    return {"references": references}


@api_router.post("/super-admin/ai-references")
async def upload_ai_reference(
    file: UploadFile = File(...),
    user: dict = Depends(require_roles("super_admin")),
):
    filename = file.filename or "referensi"
    extension = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if extension not in AI_REFERENCE_MIME_TYPES:
        raise HTTPException(status_code=400, detail="Format referensi belum didukung.")
    data = await file.read()
    if not data or len(data) > 20 * 1024 * 1024:
        raise HTTPException(status_code=400, detail="Ukuran referensi harus antara 1 byte dan 20MB.")
    reference_id = str(uuid.uuid4())
    mime_type = AI_REFERENCE_MIME_TYPES[extension]
    try:
        extracted = await extract_reference_text(
            data,
            filename,
            extension,
            mime_type,
            GEMINI_API_KEY,
        )
    except Exception as error:
        raise HTTPException(status_code=422, detail=f"Referensi tidak dapat diproses: {str(error)[:160]}")
    chunks = split_reference_text(extracted)
    if not chunks:
        raise HTTPException(status_code=422, detail="Tidak ada informasi yang dapat diindeks.")
    stored = put_object(f"{APP_NAME}/ai-references/{reference_id}.{extension}", data, mime_type)
    reference = {
        "id": reference_id,
        "name": filename,
        "content_type": mime_type,
        "size": stored["size"],
        "storage_path": stored["path"],
        "chunk_count": len(chunks),
        "created_by": user["user_id"],
        "created_at": datetime.now(timezone.utc).isoformat(),
    }
    await db.ai_references.insert_one(dict(reference))
    await db.ai_reference_chunks.insert_many([
        {
            "id": str(uuid.uuid4()),
            "reference_id": reference_id,
            "reference_name": filename,
            "text": chunk,
        }
        for chunk in chunks
    ])
    reference.pop("storage_path", None)
    return {"reference": reference}


@api_router.delete("/super-admin/ai-references/{reference_id}")
async def delete_ai_reference(
    reference_id: str,
    user: dict = Depends(require_roles("super_admin")),
):
    result = await db.ai_references.delete_one({"id": reference_id})
    if not result.deleted_count:
        raise HTTPException(status_code=404, detail="Referensi AI tidak ditemukan.")
    await db.ai_reference_chunks.delete_many({"reference_id": reference_id})
    return {"message": "Referensi AI dan indeksnya telah dihapus."}


@api_router.get("/student/ai-chat/messages")
async def list_student_ai_chat_messages(user: dict = Depends(require_roles("student"))):
    session_id = f"student-rag-{user['user_id']}"
    messages = await db.student_ai_chat_messages.find(
        {"session_id": session_id},
        {"_id": 0},
    ).sort("created_at", 1).to_list(40)
    return {"session_id": session_id, "messages": messages}


@api_router.post("/student/ai-chat/stream")
async def stream_student_ai_chat(
    payload: StudentAiChatInput,
    user: dict = Depends(require_roles("student")),
):
    question = payload.question.strip()
    session_id = f"student-rag-{user['user_id']}"
    chunks = await retrieve_ai_reference_chunks(question)
    sources = list({chunk["reference_name"] for chunk in chunks})

    async def event_stream():
        answer = ""
        try:
            if not chunks:
                answer = OUT_OF_SCOPE_MESSAGE
                yield f"event: delta\ndata: {json.dumps({'text': answer})}\n\n"
            else:
                context = "\n\n".join(
                    f"[{chunk['reference_name']}]\n{chunk['text']}" for chunk in chunks
                )
                chat = LlmChat(
                    api_key=GEMINI_API_KEY,
                    session_id=session_id,
                    system_message=(
                        "Jawab hanya berdasarkan referensi. Gunakan Bahasa Indonesia. "
                        f"Jika tidak cukup, jawab persis: {OUT_OF_SCOPE_MESSAGE}"
                    ),
                ).with_model("gemini", GEMINI_MODEL)
                async for event in chat.stream_message(UserMessage(text=f"REFERENSI:\n{context}\n\nPERTANYAAN: {question}")):
                    if isinstance(event, TextDelta):
                        answer += event.content
                        yield f"event: delta\ndata: {json.dumps({'text': event.content})}\n\n"
                    elif isinstance(event, StreamDone):
                        break
                if not answer.strip():
                    answer = OUT_OF_SCOPE_MESSAGE
                    yield f"event: delta\ndata: {json.dumps({'text': answer})}\n\n"
        except Exception:
            answer = OUT_OF_SCOPE_MESSAGE
            yield f"event: delta\ndata: {json.dumps({'text': answer})}\n\n"
        finally:
            now = datetime.now(timezone.utc).isoformat()
            await db.student_ai_chat_messages.insert_many([
                {"id": str(uuid.uuid4()), "session_id": session_id, "role": "user", "text": question, "created_at": now},
                {"id": str(uuid.uuid4()), "session_id": session_id, "role": "assistant", "text": answer, "sources": sources, "created_at": datetime.now(timezone.utc).isoformat()},
            ])
            yield f"event: done\ndata: {json.dumps({'sources': sources})}\n\n"

    return StreamingResponse(
        event_stream(),
        media_type="text/event-stream",
        headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"},
    )


@api_router.get("/")
async def root():
    return {"message": "MDJ Scholarship API"}


app.include_router(api_router)

app.add_middleware(
    CORSMiddleware,
    allow_credentials=True,
    allow_origins=os.environ.get('CORS_ORIGINS', '*').split(','),
    allow_methods=["*"],
    allow_headers=["*"],
)


@app.on_event("startup")
async def startup():
    await db.users.create_index("email", unique=True)
    await db.users.create_index("user_id")
    await db.users.create_index([("role", 1), ("is_archived", 1)])
    await db.user_sessions.create_index("session_token")
    await db.password_reset_tokens.create_index("token_hash", unique=True)
    await db.password_reset_tokens.create_index("expires_at", expireAfterSeconds=0)
    await db.password_reset_rate_limits.create_index("key", unique=True)
    await db.password_reset_rate_limits.create_index("expires_at", expireAfterSeconds=0)
    await db.selection_announcements.create_index([("status_publikasi", 1), ("published_at", -1)])
    await db.selection_announcements.create_index([("category", 1), ("published_at", -1)])
    await db.notifications.create_index([("user_id", 1), ("type", 1), ("is_popup_seen", 1)])
    await db.registrations.create_index(
        [("status", 1), ("category", 1), ("is_announcement_published", 1)]
    )
    await db.registrations.create_index("user_id")
    await db.registrations.create_index([("is_archived", 1), ("archive_batch_id", 1)])
    await db.profiles.create_index([("is_archived", 1), ("archive_batch_id", 1)])
    await db.documents.create_index([("is_archived", 1), ("archive_batch_id", 1)])
    await db.notifications.create_index([("is_archived", 1), ("archive_batch_id", 1)])
    await db.applicant_archive_batches.create_index([("status", 1), ("archived_at", -1)])
    await db.registrations.create_index("cpm_id", unique=True, sparse=True)
    await db.documents.create_index("user_id")
    await db.campuses.create_index("id", unique=True)
    await db.campuses.create_index("code", unique=True, sparse=True)
    await db.selection_imports.create_index("id", unique=True)
    await db.selection_import_review.create_index([("status", 1), ("created_at", -1)])
    await db.verification_recommendations.create_index("user_id", unique=True)
    await db.verification_recommendations.create_index([("approval_status", 1), ("analyzed_at", -1)])
    await db.beneficiary_records.create_index("user_id", unique=True)
    await db.beneficiary_disbursements.create_index([("user_id", 1), ("stage", 1)], unique=True)
    await db.beneficiary_sources.create_index("id", unique=True)
    await db.beneficiary_sources.create_index([("source_type", 1), ("created_at", -1)])
    await db.beneficiary_reviews.create_index([("status", 1), ("created_at", -1)])
    await db.beneficiary_audit_events.create_index([("user_id", 1), ("created_at", -1)])
    await db.campus_disbursements.create_index(
        [("campus_key", 1), ("region", 1), ("stage", 1)],
        unique=True,
    )
    await db.campus_disbursements.create_index([("region", 1), ("stage", 1)])
    await db.campus_disbursements.create_index("recipient_user_ids")
    await db.notifications.create_index([("user_id", 1), ("type", 1), ("source_id", 1)])
    await db.active_letter_approvals.create_index([("workflow", 1), ("status", 1), ("created_at", -1)])
    await db.active_letter_approvals.create_index([("user_id", 1), ("source_id", 1)])
    await db.stage_ii_eligibilities.create_index("user_id", unique=True)
    await db.admin_live_chat_messages.create_index([("session_id", 1), ("created_at", 1)])
    await db.student_live_chat_messages.create_index([("student_id", 1), ("created_at", 1)])
    await db.student_live_chat_messages.create_index("attachment.id", sparse=True)
    await db.ai_reference_chunks.create_index("reference_id")
    await db.student_ai_chat_messages.create_index([("session_id", 1), ("created_at", 1)])
    # Seed super admin
    existing = await db.users.find_one({"email": ADMIN_EMAIL})
    if not existing:
        await db.users.insert_one({
            "user_id": f"user_{uuid.uuid4().hex[:12]}", "email": ADMIN_EMAIL,
            "password_hash": hash_password(ADMIN_PASSWORD), "name": "Super Admin MDJ",
            "role": "super_admin", "auth_provider": "password", "nik": None, "phone": None,
            "picture": None, "is_active": True, "created_at": datetime.now(timezone.utc).isoformat(),
        })
        logger.info("Seeded super admin")
    elif not verify_password(ADMIN_PASSWORD, existing.get("password_hash") or ""):
        await db.users.update_one(
            {"email": ADMIN_EMAIL},
            {
                "$set": {
                    "password_hash": hash_password(ADMIN_PASSWORD),
                    "role": "super_admin",
                }
            },
        )

    demo_accounts = [
        {
            "email": DEMO_STUDENT_EMAIL,
            "password": DEMO_STUDENT_PASSWORD,
            "name": "Mahasiswa Demo MDJ",
            "role": "student",
        },
        {
            "email": DEMO_ADMIN_EMAIL,
            "password": DEMO_ADMIN_PASSWORD,
            "name": "Admin Provinsi MDJ",
            "role": "admin",
        },
    ]
    for account in demo_accounts:
        exists = await db.users.find_one({"email": account["email"]}, {"_id": 0})
        if not exists:
            await db.users.insert_one(
                {
                    "user_id": f"user_{uuid.uuid4().hex[:12]}",
                    "email": account["email"],
                    "password_hash": hash_password(account["password"]),
                    "name": account["name"],
                    "role": account["role"],
                    "auth_provider": "password",
                    "nik": None,
                    "phone": None,
                    "picture": None,
                    "is_active": True,
                    "created_at": datetime.now(timezone.utc).isoformat(),
                }
            )
            logger.info("Seeded demo %s account", account["role"])
            continue

        updates = {}
        if exists.get("role") != account["role"]:
            updates["role"] = account["role"]
        if exists.get("name") != account["name"]:
            updates["name"] = account["name"]
        if not exists.get("is_active", True):
            updates["is_active"] = True
        if not verify_password(account["password"], exists.get("password_hash") or ""):
            updates["password_hash"] = hash_password(account["password"])
        if updates:
            await db.users.update_one({"email": account["email"]}, {"$set": updates})
            logger.info("Refreshed demo %s account", account["role"])
    # Seed site content
    if not await db.site_content.find_one({"key": "main"}):
        await db.site_content.insert_one({"key": "main", **DEFAULT_SITE_CONTENT})
    try:
        init_storage()
        logger.info("Storage initialized")
    except Exception as e:
        logger.error(f"Storage init failed: {e}")


@app.on_event("shutdown")
async def shutdown_db_client():
    client.close()
