"""Tests for POST /api/super-admin/legal-information/import (PDF/DOCX/DOC parsing)."""
import io
import os
import copy
import pytest
import requests
from docx import Document
from fpdf import FPDF

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL", "").rstrip("/")

SUPER_ADMIN = {"email": "supri@baznasbazisdki.id", "password": "MdjSuper2026!"}
STUDENT = {"email": "mahasiswa.mdj@baznasbazisdki.id", "password": "MahasiswaMDJ2026!"}

DOCX_TEXT = "Kebijakan Privasi MDJ testing DOCX kalimat unik AZ12345."
PDF_TEXT = "Syarat dan Ketentuan MDJ testing PDF kalimat unik BX67890."


def _login(payload):
    r = requests.post(f"{BASE_URL}/api/auth/login", json=payload, timeout=30)
    assert r.status_code == 200, f"login failed for {payload['email']}: {r.status_code} {r.text}"
    return r.json()["token"]


def _make_docx(text: str) -> bytes:
    buf = io.BytesIO()
    doc = Document()
    doc.add_paragraph(text)
    doc.save(buf)
    return buf.getvalue()


def _make_pdf(text: str) -> bytes:
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    pdf.cell(0, 10, text)
    return bytes(pdf.output(dest="S"))


@pytest.fixture(scope="module")
def super_token():
    return _login(SUPER_ADMIN)


@pytest.fixture(scope="module")
def student_token():
    return _login(STUDENT)


@pytest.fixture(scope="module")
def original_legal(super_token):
    r = requests.get(f"{BASE_URL}/api/site/content", timeout=30)
    assert r.status_code == 200
    data = r.json()
    original = {
        "privacy": copy.deepcopy(data.get("privacy") or {"title": "Kebijakan Privasi", "content": ""}),
        "terms": copy.deepcopy(data.get("terms") or {"title": "Syarat dan Ketentuan", "content": ""}),
        "guide": copy.deepcopy(data.get("guide") or {"title": "Pedoman Program", "content": ""}),
    }
    yield original
    # Restore
    requests.put(
        f"{BASE_URL}/api/site/content",
        json={"content": original},
        headers={"Authorization": f"Bearer {super_token}"},
        timeout=30,
    )


def _import(token, filename, data, mime):
    files = {"file": (filename, data, mime)}
    return requests.post(
        f"{BASE_URL}/api/super-admin/legal-information/import",
        files=files,
        headers={"Authorization": f"Bearer {token}"},
        timeout=60,
    )


# ---- Extraction tests ----

def test_import_docx_returns_extracted_text(super_token):
    r = _import(super_token, "privacy.docx", _make_docx(DOCX_TEXT),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    assert r.status_code == 200, r.text
    content = r.json().get("content", "")
    assert "AZ12345" in content, f"Extracted docx missing marker: {content!r}"


def test_import_pdf_returns_extracted_text(super_token):
    r = _import(super_token, "terms.pdf", _make_pdf(PDF_TEXT), "application/pdf")
    assert r.status_code == 200, r.text
    content = r.json().get("content", "")
    assert "BX67890" in content, f"Extracted pdf missing marker: {content!r}"


# ---- Rejection tests ----

def test_import_unsupported_type_rejected(super_token):
    r = _import(super_token, "img.jpg", b"\xff\xd8\xff\xe0garbage", "image/jpeg")
    assert r.status_code == 400
    assert "PDF" in r.json().get("detail", "")


def test_import_unauthenticated_rejected():
    files = {"file": ("x.pdf", _make_pdf("hi"), "application/pdf")}
    r = requests.post(f"{BASE_URL}/api/super-admin/legal-information/import", files=files, timeout=30)
    assert r.status_code in (401, 403)


def test_import_student_rejected(student_token):
    r = _import(student_token, "x.pdf", _make_pdf("hi"), "application/pdf")
    assert r.status_code == 403


# ---- End-to-end: save + public reflection ----

def test_save_and_public_reflects_content(super_token, original_legal):
    # extract via docx
    r = _import(super_token, "guide.docx", _make_docx(DOCX_TEXT),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    assert r.status_code == 200
    extracted = r.json()["content"]

    new_legal = copy.deepcopy(original_legal)
    new_legal["guide"] = {"title": "Pedoman Program", "content": extracted}
    put = requests.put(
        f"{BASE_URL}/api/site/content",
        json={"content": new_legal},
        headers={"Authorization": f"Bearer {super_token}"},
        timeout=30,
    )
    assert put.status_code == 200, put.text

    get = requests.get(f"{BASE_URL}/api/site/content", timeout=30)
    assert get.status_code == 200
    assert "AZ12345" in (get.json().get("guide") or {}).get("content", "")
